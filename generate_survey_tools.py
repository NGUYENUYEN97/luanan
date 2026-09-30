import os
import sys
import subprocess

def install_and_import(package):
    try:
        __import__(package)
    except ImportError:
        subprocess.check_call([sys.executable, "-m", "pip", "install", package])
        __import__(package)

install_and_import('pandas')
install_and_import('openpyxl')
install_and_import('docx')

import pandas as pd
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

OUTPUT_DIR = r"g:\My Drive\Luan an uyen 26\LuanAn_TS_SHTT_KinhTeSo\04_Survey_Interview_Design"
os.makedirs(OUTPUT_DIR, exist_ok=True)

# Data for Survey Measurement Matrix
survey_data = [
    # LF
    ["Khung pháp lý (LF)", "LF1", "Các quy định pháp luật sở hữu trí tuệ hiện hành bảo vệ hiệu quả các tài sản trí tuệ mới (như nội dung số, phần mềm, AI).", "WIPO (2020), Học giả A (2021)", "RQ1, RQ2", "H1", "Chương 2, Chương 3", "Bài báo 1, Bài báo 3", "Cán bộ quản lý, Chuyên gia pháp lý, Doanh nghiệp", "Likert 5 điểm", "Kiểm tra xem từ 'tài sản trí tuệ mới' có rõ nghĩa không."],
    ["Khung pháp lý (LF)", "LF2", "Chế tài xử phạt các hành vi vi phạm sở hữu trí tuệ trên môi trường mạng là đủ mạnh và có tính răn đe.", "Luật SHTT 2022, Học giả B (2022)", "RQ1, RQ2", "H1", "Chương 2, Chương 3", "Bài báo 1, Bài báo 3", "Cán bộ quản lý, Chuyên gia pháp lý, Doanh nghiệp", "Likert 5 điểm", ""],
    ["Khung pháp lý (LF)", "LF3", "Quy định về trách nhiệm của các nền tảng trung gian (như Facebook, YouTube) là rõ ràng và dễ thực thi.", "Học giả C (2019)", "RQ1, RQ2", "H1", "Chương 2, Chương 3", "Bài báo 1, Bài báo 3", "Cán bộ quản lý, Chuyên gia pháp lý, Doanh nghiệp", "Likert 5 điểm", "Đảm bảo đối tượng khảo sát hiểu 'nền tảng trung gian'."],
    ["Khung pháp lý (LF)", "LF4", "Thủ tục hành chính để đăng ký và xử lý vi phạm SHTT trực tuyến diễn ra thuận lợi, nhanh chóng.", "Học giả D (2021)", "RQ1, RQ2", "H1", "Chương 2, Chương 3", "Bài báo 1, Bài báo 3", "Doanh nghiệp", "Likert 5 điểm", ""],
    
    # OS
    ["Tổ chức bộ máy (OS)", "OS1", "Có sự phân công chức năng, nhiệm vụ rõ ràng giữa các cơ quan quản lý nhà nước về SHTT.", "Bộ KHCN (2021)", "RQ1, RQ2", "H2", "Chương 3", "Bài báo 3", "Cán bộ quản lý", "Likert 5 điểm", ""],
    ["Tổ chức bộ máy (OS)", "OS2", "Sự phối hợp liên ngành (Khoa học, Thông tin & Truyền thông, Công an) trong xử lý vi phạm trên Internet đạt hiệu quả cao.", "Học giả E (2020)", "RQ1, RQ2", "H2", "Chương 3", "Bài báo 3", "Cán bộ quản lý, Doanh nghiệp", "Likert 5 điểm", ""],
    ["Tổ chức bộ máy (OS)", "OS3", "Quy mô và cơ cấu tổ chức của bộ máy quản lý nhà nước hiện tại đáp ứng tốt khối lượng công việc ngày càng tăng.", "Học giả F (2022)", "RQ1, RQ2", "H2", "Chương 3", "Bài báo 3", "Cán bộ quản lý", "Likert 5 điểm", ""],

    # HR
    ["Nguồn nhân lực (HR)", "HR1", "Cán bộ quản lý nhà nước được trang bị đầy đủ kiến thức về các công nghệ số mới.", "Học giả G (2021)", "RQ1, RQ2", "H3", "Chương 3", "Bài báo 2, Bài báo 3", "Cán bộ quản lý", "Likert 5 điểm", ""],
    ["Nguồn nhân lực (HR)", "HR2", "Đội ngũ xử lý vi phạm có năng lực chuyên môn tốt trong việc giải quyết các tranh chấp phức tạp trên không gian mạng.", "Học giả H (2022)", "RQ1, RQ2", "H3", "Chương 3", "Bài báo 2, Bài báo 3", "Cán bộ quản lý", "Likert 5 điểm", ""],
    ["Nguồn nhân lực (HR)", "HR3", "Cán bộ thường xuyên được tập huấn, cập nhật về các thủ đoạn vi phạm SHTT công nghệ cao.", "Học giả I (2020)", "RQ1, RQ2", "H3", "Chương 3", "Bài báo 2, Bài báo 3", "Cán bộ quản lý", "Likert 5 điểm", ""],

    # TI
    ["Hạ tầng công nghệ (TI)", "TI1", "Hệ thống cơ sở dữ liệu quốc gia về sở hữu trí tuệ đã hoàn thiện và rất dễ tra cứu.", "Học giả J (2019)", "RQ1, RQ2", "H4", "Chương 3", "Bài báo 2, Bài báo 3", "Tất cả", "Likert 5 điểm", ""],
    ["Hạ tầng công nghệ (TI)", "TI2", "Việc ứng dụng công nghệ (như AI, Big Data) trong phát hiện vi phạm SHTT trực tuyến được triển khai hiệu quả.", "Học giả K (2023)", "RQ1, RQ2", "H4", "Chương 3", "Bài báo 2, Bài báo 3", "Cán bộ quản lý", "Likert 5 điểm", ""],
    ["Hạ tầng công nghệ (TI)", "TI3", "Hệ thống lưu trữ và bảo mật dữ liệu SHTT của cơ quan nhà nước được đảm bảo an toàn tuyệt đối.", "Học giả L (2021)", "RQ1, RQ2", "H4", "Chương 3", "Bài báo 2, Bài báo 3", "Cán bộ quản lý", "Likert 5 điểm", ""],

    # AC
    ["Nhận thức & Hợp tác (AC)", "AC1", "Các doanh nghiệp và cộng đồng mạng có ý thức cao trong việc tôn trọng quyền SHTT.", "Học giả M (2020)", "RQ1, RQ2", "H5", "Chương 3", "Bài báo 3", "Doanh nghiệp", "Likert 5 điểm", ""],
    ["Nhận thức & Hợp tác (AC)", "AC2", "Chủ sở hữu SHTT luôn chủ động và phối hợp tích cực với cơ quan chức năng khi phát hiện vi phạm.", "Học giả N (2022)", "RQ1, RQ2", "H5", "Chương 3", "Bài báo 3", "Doanh nghiệp", "Likert 5 điểm", ""],
    ["Nhận thức & Hợp tác (AC)", "AC3", "Các nền tảng số xuyên biên giới thể hiện sự tuân thủ tốt các quy định về bảo vệ SHTT của Việt Nam.", "Học giả O (2021)", "RQ1, RQ2", "H5", "Chương 3", "Bài báo 3", "Cán bộ quản lý", "Likert 5 điểm", ""],

    # ME
    ["Hiệu quả QLNN (ME)", "ME1", "Tình trạng vi phạm sở hữu trí tuệ trên môi trường số đã giảm đáng kể trong những năm gần đây.", "Học giả P (2023)", "RQ1", "Biến phụ thuộc", "Chương 3, Chương 4", "Bài báo 1, 2, 3", "Tất cả", "Likert 5 điểm", ""],
    ["Hiệu quả QLNN (ME)", "ME2", "Môi trường kinh doanh số tại Việt Nam ngày càng trở nên minh bạch và an toàn hơn cho các chủ thể sáng tạo.", "Học giả Q (2022)", "RQ1", "Biến phụ thuộc", "Chương 3, Chương 4", "Bài báo 1, 2, 3", "Tất cả", "Likert 5 điểm", ""],
    ["Hiệu quả QLNN (ME)", "ME3", "Quyền lợi hợp pháp của chủ sở hữu SHTT trên môi trường mạng luôn được bảo vệ kịp thời và thỏa đáng.", "Học giả R (2021)", "RQ1", "Biến phụ thuộc", "Chương 3, Chương 4", "Bài báo 1, 2, 3", "Tất cả", "Likert 5 điểm", ""],
]

columns = [
    "Construct/Nhóm nhân tố", "Variable code", "Survey item", "Theoretical source", 
    "Research question", "Hypothesis if any", "Thesis chapter/section", "Product use", 
    "Respondent group", "Scale", "Note for pilot test"
]

df_survey = pd.DataFrame(survey_data, columns=columns)
df_survey.to_excel(os.path.join(OUTPUT_DIR, "SURVEY_MEASUREMENT_MATRIX.xlsx"), index=False)

# Mapping Matrices
df_rq_map = df_survey[["Construct/Nhóm nhân tố", "Variable code", "Research question"]]
df_rq_map.to_excel(os.path.join(OUTPUT_DIR, "VARIABLE_TO_RESEARCH_QUESTION_MAP.xlsx"), index=False)

df_chapter_map = df_survey[["Construct/Nhóm nhân tố", "Variable code", "Thesis chapter/section"]]
df_chapter_map.to_excel(os.path.join(OUTPUT_DIR, "VARIABLE_TO_THESIS_CHAPTER_MAP.xlsx"), index=False)

df_product_map = df_survey[["Construct/Nhóm nhân tố", "Variable code", "Product use"]]
df_product_map.to_excel(os.path.join(OUTPUT_DIR, "VARIABLE_TO_PUBLICATION_PRODUCT_MAP.xlsx"), index=False)

# Create docx for Survey Questionnaire
def create_survey_docx():
    doc = Document()
    doc.add_heading('PHIẾU KHẢO SÁT CHUYÊN GIA/ DOANH NGHIỆP', 0)
    doc.add_paragraph('Đề tài: Quản lý nhà nước về quyền sở hữu trí tuệ trong nền kinh tế số ở Việt Nam\n', style='Normal')
    
    doc.add_heading('PHẦN A. THÔNG TIN NGƯỜI TRẢ LỜI', level=1)
    doc.add_paragraph('Xin vui lòng điền hoặc đánh dấu (X) vào ô phù hợp với đặc điểm của Ông/Bà:')
    doc.add_paragraph('1. Vị trí công tác hiện tại:\n[ ] Cán bộ quản lý nhà nước\n[ ] Chuyên gia pháp lý/ Học giả\n[ ] Đại diện doanh nghiệp/ Chủ sở hữu SHTT\n[ ] Khác (Ghi rõ): ..................................')
    doc.add_paragraph('2. Số năm kinh nghiệm công tác:\n[ ] Dưới 5 năm\n[ ] Từ 5 đến 10 năm\n[ ] Trên 10 năm')
    
    doc.add_heading('PHẦN B. NỘI DUNG ĐÁNH GIÁ', level=1)
    doc.add_paragraph('Xin vui lòng đánh giá mức độ đồng ý của Ông/Bà đối với các phát biểu sau theo thang điểm từ 1 đến 5:')
    doc.add_paragraph('(1: Hoàn toàn không đồng ý; 2: Không đồng ý; 3: Bình thường/Không ý kiến; 4: Đồng ý; 5: Hoàn toàn đồng ý)')
    
    table = doc.add_table(rows=1, cols=3)
    table.style = 'Table Grid'
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Mã'
    hdr_cells[1].text = 'Nội dung phát biểu'
    hdr_cells[2].text = 'Đánh giá (1-5)'
    
    for row in survey_data:
        row_cells = table.add_row().cells
        row_cells[0].text = row[1]
        row_cells[1].text = row[2]
        row_cells[2].text = ''

    doc.add_heading('PHẦN C. CÂU HỎI MỞ', level=1)
    doc.add_paragraph('OM1: Theo Ông/Bà, đâu là rào cản hoặc khó khăn lớn nhất hiện nay trong quản lý nhà nước về SHTT trên môi trường số tại Việt Nam?')
    doc.add_paragraph('\n\n\n\n')
    doc.add_paragraph('OM2: Ông/Bà có đề xuất hay kiến nghị giải pháp gì nhằm nâng cao hiệu quả QLNN về SHTT trong nền kinh tế số không?')
    doc.add_paragraph('\n\n\n\n')
    
    doc.add_paragraph('Xin chân thành cảm ơn sự tham gia và hỗ trợ quý báu của Ông/Bà!')
    doc.save(os.path.join(OUTPUT_DIR, "SURVEY_QUESTIONNAIRE.docx"))

create_survey_docx()

# Create docx for Interview Protocol
def create_interview_docx():
    doc = Document()
    doc.add_heading('KỊCH BẢN PHỎNG VẤN CHUYÊN GIA SÂU', 0)
    doc.add_paragraph('Đề tài: Quản lý nhà nước về quyền sở hữu trí tuệ trong nền kinh tế số ở Việt Nam', style='Normal')
    doc.add_paragraph('Số lượng dự kiến: 10 chuyên gia\nĐối tượng: Lãnh đạo cơ quan QLNN, Thẩm phán chuyên trách SHTT, Chuyên gia hàng đầu về Luật SHTT, Đại diện doanh nghiệp công nghệ lớn.\n')
    
    doc.add_heading('1. Nhóm câu hỏi về Khung pháp lý', level=2)
    doc.add_paragraph('- Thưa Ông/Bà, mức độ tương thích của pháp luật SHTT Việt Nam hiện nay so với các chuẩn mực quốc tế (như CPTPP, EVFTA, WIPO Internet Treaties) trong bối cảnh kinh tế số được đánh giá như thế nào?')
    doc.add_paragraph('- Có những khoảng trống hay rào cản pháp lý nào khiến việc bảo vệ tài sản trí tuệ (đặc biệt là nội dung số) gặp khó khăn?')
    
    doc.add_heading('2. Nhóm câu hỏi về Tổ chức bộ máy quản lý nhà nước', level=2)
    doc.add_paragraph('- Việc phân cấp, phân quyền giữa Bộ KH&CN, Bộ TT&TT, Bộ VHTTDL và cơ quan Công an hiện nay đã tạo ra sự linh hoạt và đồng bộ trong xử lý vi phạm SHTT trực tuyến chưa?')
    doc.add_paragraph('- Có hiện tượng chồng chéo hoặc đùn đẩy trách nhiệm giữa các cơ quan khi xử lý một vụ việc vi phạm SHTT trên Internet không?')

    doc.add_heading('3. Nhóm câu hỏi về Thực thi quyền SHTT trên môi trường số', level=2)
    doc.add_paragraph('- Tại sao việc yêu cầu gỡ bỏ nội dung vi phạm và xử phạt trên các nền tảng xuyên biên giới (Facebook, YouTube, TikTok) thường chậm trễ và kém hiệu quả?')
    doc.add_paragraph('- Theo Ông/Bà, các cơ chế như "Notice and Takedown" hiện tại ở Việt Nam đã hoạt động thực sự hiệu quả chưa?')

    doc.add_heading('4. Nhóm câu hỏi về Dữ liệu, Nền tảng số và Công nghệ', level=2)
    doc.add_paragraph('- Mức độ số hóa dữ liệu và mức độ sẵn sàng của hạ tầng công nghệ tại các cơ quan quản lý nhà nước về SHTT hiện nay ra sao?')
    doc.add_paragraph('- Chúng ta đang đối mặt với những rủi ro gì về bảo mật dữ liệu và làm thế nào để ứng dụng hiệu quả AI, Big Data trong việc tự động quét, phát hiện vi phạm SHTT?')

    doc.add_heading('5. Nhóm câu hỏi về Bất cập và Nguyên nhân cốt lõi', level=2)
    doc.add_paragraph('- Nếu chỉ ra căn nguyên cốt lõi nhất dẫn đến tình trạng xâm phạm SHTT tràn lan trên không gian mạng hiện nay, Ông/Bà sẽ chỉ ra những điểm gì?')
    doc.add_paragraph('- Ý thức của chủ thể quyền, doanh nghiệp và người tiêu dùng có phải là một trong những nguyên nhân lớn nhất không?')

    doc.add_heading('6. Nhóm câu hỏi về Giải pháp và Kiến nghị chính sách', level=2)
    doc.add_paragraph('- Trong ngắn hạn 1-3 năm tới, cơ quan quản lý cần ưu tiên hành động đột phá gì để cải thiện tình hình bảo vệ SHTT trong nền kinh tế số?')
    doc.add_paragraph('- Về dài hạn, để xây dựng một môi trường kinh doanh số an toàn và thúc đẩy đổi mới sáng tạo, mô hình quản lý nhà nước cần thay đổi như thế nào?')

    doc.save(os.path.join(OUTPUT_DIR, "INTERVIEW_PROTOCOL.docx"))

create_interview_docx()

print("Files generated successfully.")
