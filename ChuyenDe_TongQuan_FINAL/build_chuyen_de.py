# Dựng "Chuyên đề tổng quan" hoàn chỉnh theo mẫu của cơ sở đào tạo.
# Nguồn: các file .md trong thư mục này + refs.py; hình trong ../hinh_de_cuong
# Xuất: ../ChuyenDe_TongQuan_FINAL.docx (chuyende) hoặc ../DeCuong_ChiTiet_FINAL.docx (decuong); sau đó chạy cap_nhat_muc_luc.ps1
import os
import re
import sys
import unicodedata
import docx
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION, WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

sys.stdout.reconfigure(encoding="utf-8")
HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, HERE)
from refs import REFS  # noqa: E402

# Chạy: python build_chuyen_de.py chuyende | decuong | chuong2
DOC = sys.argv[1] if len(sys.argv) > 1 else "chuyende"
if DOC == "chuyende":
    MD_FILES = ["01_MoDau.md", "02_Phan1_111.md", "03_Phan1_112.md", "04_Phan1_113.md",
                "05_Phan1_12_13.md", "07_KetLuan.md"]
    OUT = os.path.join(ROOT, "ChuyenDe_TongQuan_FINAL.docx")
    REF_TITLE = "TÀI LIỆU THAM KHẢO"
elif DOC == "chuong2":
    MD_FILES = ["20_Chuong2_21.md", "21_Chuong2_22.md", "22_Chuong2_23.md"]
    OUT = os.path.join(ROOT, "Chuong2_CoSoLyLuan.docx")
    REF_TITLE = "TÀI LIỆU THAM KHẢO CHƯƠNG 2"
else:
    MD_FILES = ["10_DeCuong.md"]
    OUT = os.path.join(ROOT, "DeCuong_ChiTiet_FINAL.docx")
    REF_TITLE = "DANH MỤC TÀI LIỆU VIỆN DẪN"
CAPTION_STYLE = "Chu thich hinh bang"

ABBR = [("CSIRO", "Cơ quan Nghiên cứu Khoa học và Công nghiệp Khối thịnh vượng chung Úc"),
        ("EPO", "Cơ quan Sáng chế châu Âu (European Patent Office)"),
        ("EUIPO", "Cơ quan Sở hữu trí tuệ Liên minh châu Âu (European Union Intellectual Property Office)"),
        ("FDI", "Đầu tư trực tiếp nước ngoài (Foreign Direct Investment)"),
        ("GDP", "Tổng sản phẩm trong nước (Gross Domestic Product)"),
        ("IMF", "Quỹ Tiền tệ Quốc tế (International Monetary Fund)"),
        ("OECD", "Tổ chức Hợp tác và Phát triển Kinh tế (Organisation for Economic Co-operation and Development)"),
        ("QLNN", "Quản lý nhà nước"),
        ("SHTT", "Sở hữu trí tuệ"),
        ("UNCTAD", "Hội nghị Liên hợp quốc về Thương mại và Phát triển (United Nations Conference on Trade and Development)"),
        ("WIPO", "Tổ chức Sở hữu trí tuệ Thế giới (World Intellectual Property Organization)")]

# ------------------------------------------------------------------ sắp xếp tiếng Việt
VN_ORDER = "aăâbcdđeêfghijklmnoôơpqrstuưvwxyz"
TONE_MARKS = {"̀", "́", "̃", "̉", "̣"}


def vn_key(s):
    out = []
    for ch in s.lower():
        dec = unicodedata.normalize("NFD", ch)
        base = "".join(c for c in dec if c not in TONE_MARKS)
        base = unicodedata.normalize("NFC", base)
        tone = sum(ord(c) for c in dec if c in TONE_MARKS)
        if base in VN_ORDER:
            out.append((VN_ORDER.index(base), tone))
        elif base.isdigit():
            out.append((-1, int(base)))
        else:
            out.append((-2, 0))
    return out


# ------------------------------------------------------------------ đọc nội dung, đánh số trích dẫn
def read_md(name):
    out = []
    for line in open(os.path.join(HERE, name), encoding="utf-8").read().split("\n"):
        if line.startswith("!INCLUDE "):
            out.append(read_md(line[9:].strip()))
        else:
            out.append(line)
    return "\n".join(out)


text = "\n".join(read_md(f) for f in MD_FILES)
cited = []
for m in re.finditer(r"\[@([^\]]+)\]", text):
    for k in m.group(1).split(";"):
        k = k.strip().lstrip("@")
        if k not in REFS:
            raise SystemExit(f"Thiếu tài liệu: {k}")
        if k not in cited:
            cited.append(k)
unused = [k for k in REFS if k not in cited]
vi = sorted([k for k in cited if REFS[k][0] == "vi"], key=lambda k: vn_key(REFS[k][1]))
en = sorted([k for k in cited if REFS[k][0] == "en"], key=lambda k: REFS[k][1].lower())
NUM = {k: i + 1 for i, k in enumerate(vi + en)}


def cite_repl(m):
    ks = [k.strip().lstrip("@") for k in m.group(1).split(";")]
    return ", ".join(f"[{n}]" for n in sorted(NUM[k] for k in ks))


text = re.sub(r"\[@([^\]]+)\]", cite_repl, text)

# ------------------------------------------------------------------ tiện ích docx
BLACK = RGBColor(0, 0, 0)


def fmt_run(run, size=13, bold=None, italic=None):
    run.font.name = "Times New Roman"
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(a), "Times New Roman")
    run.font.size = Pt(size)
    run.font.color.rgb = BLACK
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic


def add_inline(p, s, size=13, bold=False, italic=False):
    """Hỗ trợ *nghiêng* trong dòng."""
    parts = re.split(r"(\*[^*]+\*)", s)
    for part in parts:
        if not part:
            continue
        if part.startswith("*") and part.endswith("*") and len(part) > 2:
            fmt_run(p.add_run(part[1:-1]), size, bold, not italic)
        else:
            fmt_run(p.add_run(part), size, bold, italic)


def para(doc, s="", align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent=True, size=13, bold=False, italic=False,
         before=0, after=0, style=None, left=None, hanging=False, spacing=1.5, keep_next=False):
    p = doc.add_paragraph(style=style) if style else doc.add_paragraph()
    p.alignment = align
    pf = p.paragraph_format
    pf.space_before, pf.space_after = Pt(before), Pt(after)
    pf.line_spacing = spacing
    if hanging:
        pf.left_indent, pf.first_line_indent = Cm(1), Cm(-1)
    elif indent:
        pf.first_line_indent = Cm(1.27)
    else:
        pf.first_line_indent = Cm(0)
    if left is not None:
        pf.left_indent = Cm(left)
    pf.keep_with_next = keep_next
    if s:
        add_inline(p, s, size, bold, italic)
    return p


def field(p, instr):
    r = p.add_run()
    for kind, val in (("begin", None), ("instr", instr), ("separate", None), ("text", "Nhấn chuột phải, chọn Update Field để cập nhật."), ("end", None)):
        if kind == "instr":
            el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = val; r._r.append(el)
        elif kind == "text":
            t = OxmlElement("w:t"); t.text = val; r._r.append(t)
        else:
            el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), kind); r._r.append(el)
    fmt_run(r, 13)


def page_number_header(section):
    section.header.is_linked_to_previous = False
    hp = section.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = hp.add_run()
    for kind in ("begin", "instr", "end"):
        if kind == "instr":
            el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = "PAGE"
        else:
            el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), kind)
        r._r.append(el)
    fmt_run(r, 13)


def empty_header(section):
    section.header.is_linked_to_previous = False
    for p in section.header.paragraphs:
        for r in p.runs:
            r.text = ""


def set_pgnum(section, fmt=None, start=None):
    sp = section._sectPr
    pg = sp.find(qn("w:pgNumType"))
    if pg is None:
        pg = OxmlElement("w:pgNumType"); sp.append(pg)
    for a in ("w:fmt", "w:start"):
        if pg.get(qn(a)) is not None:
            del pg.attrib[qn(a)]
    if fmt:
        pg.set(qn("w:fmt"), fmt)
    if start is not None:
        pg.set(qn("w:start"), str(start))


def portrait(section):
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width, section.page_height = Cm(21.0), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(3.5), Cm(3.0)
    section.left_margin, section.right_margin = Cm(3.5), Cm(2.0)


# ------------------------------------------------------------------ khởi tạo tài liệu và kiểu chữ
doc = docx.Document()
portrait(doc.sections[0])
st = doc.styles["Normal"]
st.font.name = "Times New Roman"
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
st.font.size = Pt(13)
st.paragraph_format.line_spacing = 1.5
st.paragraph_format.space_after = Pt(0)

HSPEC = {1: (14, True, False, WD_ALIGN_PARAGRAPH.CENTER, 12, 12),
         2: (13, True, False, WD_ALIGN_PARAGRAPH.JUSTIFY, 8, 4),
         3: (13, True, False, WD_ALIGN_PARAGRAPH.JUSTIFY, 6, 3),
         4: (13, True, True, WD_ALIGN_PARAGRAPH.JUSTIFY, 4, 2)}
for lvl, (sz, b, it, al, bf, af) in HSPEC.items():
    hs = doc.styles[f"Heading {lvl}"]
    hs.font.name = "Times New Roman"
    rpr = hs.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(a), "Times New Roman")
    hs.font.size, hs.font.bold, hs.font.italic = Pt(sz), b, it
    hs.font.color.rgb = BLACK
    hs.paragraph_format.alignment = al
    hs.paragraph_format.space_before, hs.paragraph_format.space_after = Pt(bf), Pt(af)
    hs.paragraph_format.line_spacing = 1.5
    hs.paragraph_format.first_line_indent = Cm(0)
    hs.paragraph_format.keep_with_next = True
cap = doc.styles.add_style(CAPTION_STYLE, 1)
cap.base_style = doc.styles["Normal"]
cap.font.name, cap.font.size, cap.font.bold = "Times New Roman", Pt(13), True
cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER


def heading(s, lvl, page_break=False):
    p = doc.add_paragraph(style=f"Heading {lvl}")
    if page_break:
        p.paragraph_format.page_break_before = True
    add_inline(p, s.upper() if lvl == 1 else s, HSPEC[lvl][0], True, HSPEC[lvl][2])
    return p


def front_matter():
  # ------------------------------------------------------------------ Trang bìa (không đánh số)
  for s, sz, b, bf in [("[CƠ SỞ ĐÀO TẠO]", 13, True, 0), ("", 13, False, 40), ("[HỌ VÀ TÊN NGHIÊN CỨU SINH]", 14, True, 40),
                     ("", 13, False, 30), ("BÁO CÁO CHUYÊN ĐỀ TỔNG QUAN", 16, True, 30),
                     ("TỔNG QUAN TÌNH HÌNH NGHIÊN CỨU LIÊN QUAN ĐẾN ĐỀ TÀI LUẬN ÁN", 14, True, 6),
                     ("", 13, False, 20), ("Đề tài luận án:", 13, False, 10),
                     ("QUẢN LÝ NHÀ NƯỚC VỀ QUYỀN SỞ HỮU TRÍ TUỆ\nTRONG NỀN KINH TẾ SỐ Ở VIỆT NAM", 15, True, 6),
                     ("", 13, False, 30), ("Chuyên ngành: Quản lý kinh tế", 13, False, 10), ("Mã số: 9310110", 13, False, 0),
                     ("", 13, False, 30), ("Người hướng dẫn khoa học: [Học hàm, học vị, họ và tên]", 13, False, 10),
                     ("", 13, False, 40), ("Hà Nội, năm 2026", 13, True, 10)]:
    para(doc, s, WD_ALIGN_PARAGRAPH.CENTER, False, sz, b, before=bf, spacing=1.2)

  # ------------------------------------------------------------------ Phần phụ đầu (La Mã)
  sec = doc.add_section(WD_SECTION.NEW_PAGE)
  portrait(sec)
  page_number_header(sec)
  set_pgnum(sec, "lowerRoman", 1)
  para(doc, "MỤC LỤC", WD_ALIGN_PARAGRAPH.CENTER, False, 14, True, after=6)
  field(doc.add_paragraph(), 'TOC \\o "1-4" \\h \\z \\u')
  p = para(doc, "BẢNG BIỂU, SƠ ĐỒ", WD_ALIGN_PARAGRAPH.CENTER, False, 14, True, after=6)
  p.paragraph_format.page_break_before = True
  field(doc.add_paragraph(), f'TOC \\h \\z \\t "{CAPTION_STYLE},1"')
  p = para(doc, "DANH MỤC CHỮ VIẾT TẮT", WD_ALIGN_PARAGRAPH.CENTER, False, 14, True, after=6)
  p.paragraph_format.page_break_before = True
  t = doc.add_table(rows=1, cols=2); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
  for i, h in enumerate(["Chữ viết tắt", "Nghĩa đầy đủ"]):
    c = t.rows[0].cells[i]; c.width = Cm([3.0, 12.5][i])
    add_inline(c.paragraphs[0], h, 13, True); c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
  for a, b in ABBR:
    cells = t.add_row().cells
    cells[0].width, cells[1].width = Cm(3.0), Cm(12.5)
    add_inline(cells[0].paragraphs[0], a, 13); add_inline(cells[1].paragraphs[0], b, 13)
    for c in cells:
        c.paragraphs[0].paragraph_format.line_spacing = 1.2
  # ------------------------------------------------------------------ Phần chính (Ả Rập, từ 1)
  sec = doc.add_section(WD_SECTION.NEW_PAGE)
  portrait(sec)
  page_number_header(sec)
  set_pgnum(sec, "decimal", 1)


if DOC == "chuyende":
    front_matter()
else:
    page_number_header(doc.sections[0])
    set_pgnum(doc.sections[0], "decimal", 1)

lines = text.split("\n")
i = 0
first_h1 = True
while i < len(lines):
    ln = lines[i].rstrip()
    if not ln.strip():
        i += 1; continue
    if ln.startswith("!TITLE "):
        para(doc, ln[7:], WD_ALIGN_PARAGRAPH.CENTER, False, 14, True, before=4, spacing=1.3)
    elif ln.startswith("!SUBTITLE "):
        para(doc, ln[10:], WD_ALIGN_PARAGRAPH.CENTER, False, 13, False, after=8, spacing=1.3)
    elif ln.startswith("!NOTE "):
        para(doc, ln[6:], italic=True, before=8, after=6)
    elif ln.startswith("!CHUONG "):
        para(doc, ln[8:], WD_ALIGN_PARAGRAPH.CENTER, False, 13, True, before=12, after=4, keep_next=True)
    elif ln.startswith("#### "):
        heading(ln[5:], 4)
    elif ln.startswith("### "):
        heading(ln[4:], 3)
    elif ln.startswith("## "):
        heading(ln[3:], 2)
    elif ln.startswith("# "):
        heading(ln[2:], 1, page_break=not first_h1)
        first_h1 = False
    elif ln.startswith("!FIGWIDE ") or ln.startswith("!FIG "):
        wide = ln.startswith("!FIGWIDE ")
        path, capt, src = [x.strip() for x in ln.split(" ", 1)[1].split("|")]
        if wide:
            s2 = doc.add_section(WD_SECTION.NEW_PAGE)
            s2.orientation = WD_ORIENT.LANDSCAPE
            s2.page_width, s2.page_height = Cm(29.7), Cm(21.0)
            s2.top_margin, s2.bottom_margin = Cm(3.5), Cm(2.0)
            s2.left_margin, s2.right_margin = Cm(3.0), Cm(3.0)
            set_pgnum(s2)
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(os.path.join(ROOT, path), width=Cm(20.5 if wide else 15.5))
        cp = doc.add_paragraph(style=CAPTION_STYLE); add_inline(cp, capt, 13, True)
        cp.paragraph_format.line_spacing = 1.5
        para(doc, src, WD_ALIGN_PARAGRAPH.RIGHT, False, 12, italic=True, after=6)
        if wide:
            s3 = doc.add_section(WD_SECTION.NEW_PAGE)
            portrait(s3); set_pgnum(s3)
    elif ln.startswith("!TABLE "):
        spec = [x.strip() for x in ln[7:].split("|")]
        capt, src = spec[0], spec[1]
        wspec = [float(x) for x in spec[2].split(",")] if len(spec) > 2 and spec[2] else None
        aspec = spec[3] if len(spec) > 3 else None
        cp = doc.add_paragraph(style=CAPTION_STYLE); add_inline(cp, capt, 13, True)
        cp.paragraph_format.keep_with_next = True; cp.paragraph_format.space_before = Pt(6)
        rows = []
        i += 1
        while i < len(lines) and lines[i].strip().startswith("|"):
            cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
            if not all(set(c) <= set("-: ") for c in cells):
                rows.append(cells)
            i += 1
        ncol = len(rows[0])
        widths = [Cm(w) for w in (wspec or [15.5 / ncol] * ncol)]
        aspec = (aspec or "L" * ncol).ljust(ncol, aspec[-1] if aspec else "L")
        tb = doc.add_table(rows=0, cols=ncol); tb.style = "Table Grid"; tb.alignment = WD_TABLE_ALIGNMENT.CENTER
        tb.autofit = False
        for ci, w in enumerate(widths):
            tb.columns[ci].width = w
        for ri, r in enumerate(rows):
            cells = tb.add_row().cells
            bold_row = ri == 0 or r[0].startswith("Tổng")
            for ci, v in enumerate(r):
                cells[ci].width = widths[ci]
                cpp = cells[ci].paragraphs[0]
                cpp.paragraph_format.line_spacing = 1.15
                cpp.alignment = WD_ALIGN_PARAGRAPH.CENTER if (ri == 0 or aspec[ci] == "C") else WD_ALIGN_PARAGRAPH.LEFT
                add_inline(cpp, v, 12, bold_row)
        trpr = tb.rows[0]._tr.get_or_add_trPr()
        th = OxmlElement("w:tblHeader"); th.set(qn("w:val"), "true"); trpr.append(th)
        para(doc, src, WD_ALIGN_PARAGRAPH.RIGHT, False, 12, italic=True, after=6)
        continue
    elif ln.startswith("!OUT3 "):
        para(doc, ln[6:], indent=False, left=1.5, italic=True, spacing=1.3)
    elif ln.startswith("!OUT2 "):
        para(doc, ln[6:], indent=False, left=0.75, spacing=1.3)
    elif ln.startswith("!OUT "):
        para(doc, ln[5:], indent=False, bold=True, spacing=1.3, before=2)
    elif ln.startswith("*") and ln.endswith("*") and ln.count("*") == 2:
        para(doc, ln[1:-1], italic=True, keep_next=True, before=4)
    else:
        para(doc, ln)
    i += 1

# ------------------------------------------------------------------ Tài liệu tham khảo (không đánh số trang)
sec = doc.add_section(WD_SECTION.NEW_PAGE)
portrait(sec)
empty_header(sec)
set_pgnum(sec)
heading(REF_TITLE, 1)
for title, keys in (("Tiếng Việt", vi), ("Tiếng Anh", en)):
    para(doc, title, indent=False, bold=True, before=6, keep_next=True)
    for k in keys:
        para(doc, f"[{NUM[k]}] " + REFS[k][2], hanging=True, spacing=1.3, after=2)

doc.save(OUT)
print("Đã tạo:", OUT)
print("Số tài liệu trích dẫn:", len(cited), "(tiếng Việt", len(vi), ", tiếng Anh", len(en), ")")
if unused:
    print("Tài liệu có trong refs.py nhưng không trích dẫn (không đưa vào danh mục):", ", ".join(unused))
