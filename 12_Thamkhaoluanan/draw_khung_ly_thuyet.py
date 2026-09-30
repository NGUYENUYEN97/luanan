from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

prs = Presentation()
slide = prs.slides.add_slide(prs.slide_layouts[6])

def add_box(slide, text, left, top, width, height, font_size=10, bold=False):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(255, 255, 255)
    shape.line.color.rgb = RGBColor(0, 0, 0)
    shape.line.width = Pt(1.5)
    
    text_frame = shape.text_frame
    text_frame.text = text
    text_frame.word_wrap = True
    
    for paragraph in text_frame.paragraphs:
        paragraph.alignment = PP_ALIGN.CENTER
        for run in paragraph.runs:
            run.font.name = 'Times New Roman'
            run.font.size = Pt(font_size)
            run.font.color.rgb = RGBColor(0, 0, 0)
            run.font.bold = bold
    return shape

def add_arrow(slide, left, top, width, height, direction='DOWN'):
    shape_type = MSO_SHAPE.DOWN_ARROW
    if direction == 'RIGHT': shape_type = MSO_SHAPE.RIGHT_ARROW
    elif direction == 'UP': shape_type = MSO_SHAPE.UP_ARROW
    elif direction == 'LEFT': shape_type = MSO_SHAPE.LEFT_ARROW
    
    shape = slide.shapes.add_shape(shape_type, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(255, 255, 255)
    shape.line.color.rgb = RGBColor(0, 0, 0)
    return shape

# Title
title_shape = slide.shapes.add_textbox(Inches(0.5), Inches(0.1), Inches(9), Inches(0.5))
tf = title_shape.text_frame
tf.text = "HÌNH 2.x: KHUNG LÝ THUYẾT NGHIÊN CỨU"
tf.paragraphs[0].alignment = PP_ALIGN.CENTER
tf.paragraphs[0].runs[0].font.name = 'Times New Roman'
tf.paragraphs[0].runs[0].font.size = Pt(14)
tf.paragraphs[0].runs[0].font.bold = True

# Row 1: Theory Base
text_theory = "CƠ SỞ LÝ THUYẾT NỀN\n(Lý thuyết quản lý nhà nước, Lý thuyết quản trị công, Lý thuyết thể chế,\nLý thuyết kinh tế số, Lý thuyết đổi mới sáng tạo, Các nghiên cứu trước)"
add_box(slide, text_theory, Inches(1.0), Inches(0.6), Inches(8.0), Inches(0.8), bold=True, font_size=11)

add_arrow(slide, Inches(5.0), Inches(1.4), Inches(0.3), Inches(0.3), 'DOWN')

# Row 2: Variables
text_nd = "NỘI DUNG QUẢN LÝ NHÀ NƯỚC\n(Biến độc lập)\n\n1. Xây dựng chiến lược, chính sách và pháp luật\n2. Tổ chức thực hiện chiến lược, chính sách và pháp luật\n3. Hỗ trợ xác lập, khai thác, chuyển giao và thương mại hóa tài sản trí tuệ\n4. Thanh tra, kiểm tra, giám sát và thực thi quyền trong môi trường số"
add_box(slide, text_nd, Inches(0.5), Inches(1.8), Inches(3.6), Inches(1.8), font_size=10, bold=True)

add_arrow(slide, Inches(4.1), Inches(2.5), Inches(0.4), Inches(0.3), 'RIGHT')

text_cl = "CHẤT LƯỢNG QUẢN LÝ NHÀ NƯỚC\n(Biến phụ thuộc)\n\n1. Hiệu lực\n2. Hiệu quả\n3. Tính phù hợp\n4. Tính đồng bộ, thống nhất\n5. Tính minh bạch"
add_box(slide, text_cl, Inches(4.5), Inches(1.8), Inches(2.6), Inches(1.8), font_size=10, bold=True)

add_arrow(slide, Inches(7.1), Inches(2.5), Inches(0.4), Inches(0.3), 'LEFT')

text_yt = "CÁC YẾU TỐ ẢNH HƯỞNG\n(Biến điều tiết/tác động)\n\n- Nhóm yếu tố bên trong\n- Nhóm yếu tố bên ngoài"
add_box(slide, text_yt, Inches(7.5), Inches(1.8), Inches(2.2), Inches(1.8), font_size=10, bold=True)

# Row 3: Evaluation
add_arrow(slide, Inches(5.65), Inches(3.6), Inches(0.3), Inches(0.4), 'DOWN')

text_tt = "ĐÁNH GIÁ THỰC TRẠNG\n(Phân tích Nội dung QLNN & Đo lường Chất lượng QLNN)\n\nXác định: KẾT QUẢ ĐẠT ĐƯỢC - HẠN CHẾ - NGUYÊN NHÂN"
add_box(slide, text_tt, Inches(2.0), Inches(4.0), Inches(6.0), Inches(1.0), font_size=11, bold=True)

# Row 4: Solutions
add_arrow(slide, Inches(5.0), Inches(5.0), Inches(0.3), Inches(0.4), 'DOWN')

text_gp = "ĐỀ XUẤT GIẢI PHÁP HOÀN THIỆN\n(Khắc phục hạn chế, giải quyết nguyên nhân, nâng cao chất lượng QLNN)"
add_box(slide, text_gp, Inches(2.0), Inches(5.4), Inches(6.0), Inches(0.8), font_size=11, bold=True)

# Dashed line surrounding the core variables to denote the "Central Research Problem"
shape_bg = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.2), Inches(1.5), Inches(9.6), Inches(2.3))
shape_bg.fill.background()
shape_bg.line.color.rgb = RGBColor(0, 0, 0)
shape_bg.line.width = Pt(1.5)
shape_bg.line.dash_style = 4
# Move the dashed box to back
slide.shapes._spTree.remove(shape_bg._element)
slide.shapes._spTree.insert(2, shape_bg._element)

# Add a small label for the dashed box
t_label = slide.shapes.add_textbox(Inches(0.3), Inches(1.55), Inches(3.0), Inches(0.3))
t_label.text_frame.text = "VẤN ĐỀ NGHIÊN CỨU TRUNG TÂM"
t_label.text_frame.paragraphs[0].runs[0].font.name = 'Times New Roman'
t_label.text_frame.paragraphs[0].runs[0].font.size = Pt(10)
t_label.text_frame.paragraphs[0].runs[0].font.bold = True
t_label.text_frame.paragraphs[0].runs[0].font.italic = True

prs.save("KhungLyThuyetNghienCuu.pptx")
print("Saved KhungLyThuyetNghienCuu.pptx")
