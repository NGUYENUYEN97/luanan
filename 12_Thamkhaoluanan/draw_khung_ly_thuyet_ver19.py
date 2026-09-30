from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_AUTO_SIZE
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls

prs = Presentation()
prs.slide_width = Inches(11.69)
prs.slide_height = Inches(8.27)
slide = prs.slides.add_slide(prs.slide_layouts[6])

def add_box(slide, text, left, top, width, height, font_size=10, fill_color=(255,255,255)):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(*fill_color)
    shape.line.color.rgb = RGBColor(0, 0, 0)
    shape.line.width = Pt(1.5)
    
    tf = shape.text_frame
    tf.text = text
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    
    for i, p in enumerate(tf.paragraphs):
        if i < 2:
            p.alignment = PP_ALIGN.CENTER
        else:
            p.alignment = PP_ALIGN.LEFT
            p.level = 0
            
        for run in p.runs:
            run.font.name = 'Times New Roman'
            run.font.size = Pt(font_size)
            run.font.color.rgb = RGBColor(0, 0, 0)
            
            # Bold the headers (first two lines)
            if i < 2:
                run.font.bold = True
    return shape

def add_line_arrow(slide, begin_x, begin_y, end_x, end_y, dashed=False, arrow_end=True, arrow_start=False, thickness=1.5):
    connector = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, begin_x, begin_y, end_x, end_y)
    connector.line.color.rgb = RGBColor(0,0,0)
    connector.line.width = Pt(thickness)
    if dashed:
        connector.line.dash_style = 4
        
    ln = connector.element.spPr.ln
    if ln is not None:
        if arrow_end:
            tailEnd = parse_xml(r'<a:tailEnd type="triangle" w="med" len="med" %s/>' % nsdecls('a'))
            ln.append(tailEnd)
        if arrow_start:
            headEnd = parse_xml(r'<a:headEnd type="triangle" w="med" len="med" %s/>' % nsdecls('a'))
            ln.append(headEnd)
    return connector


# --- TITLE ---
t_title = slide.shapes.add_textbox(Inches(0.5), Inches(0.05), Inches(10.69), Inches(0.4))
t_title.text_frame.text = "KHUNG LÝ THUYẾT NGHIÊN CỨU"
t_title.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
t_title.text_frame.paragraphs[0].runs[0].font.name = 'Times New Roman'
t_title.text_frame.paragraphs[0].runs[0].font.size = Pt(16)
t_title.text_frame.paragraphs[0].runs[0].font.bold = True


# --- LÝ THUYẾT NỀN TẢNG (Top) ---
t1 = "LÝ THUYẾT NỀN TẢNG\nLý thuyết quản lý công  •  Lý thuyết thể chế  •  Lý thuyết tăng trưởng nội sinh  •  Lý thuyết về quyền sở hữu trí tuệ"
add_box(slide, t1, Inches(1.5), Inches(0.4), Inches(8.69), Inches(0.55), font_size=10, fill_color=(250, 250, 250))


# --- TOP ROW: ĐIỀU TIẾT & KIỂM SOÁT ---
# Điều tiết (trên Trung gian)
tdt = "BIẾN ĐIỀU TIẾT\n(MÔI TRƯỜNG & HỘI NHẬP)\n\n• Trình độ công nghệ (AI...)\n• Mức độ hội nhập quốc tế\n• Nhận thức pháp luật xã hội"
add_box(slide, tdt, Inches(4.35), Inches(1.2), Inches(3.0), Inches(1.3), font_size=10, fill_color=(250, 250, 250))
# Arrow down to Trung gian
add_line_arrow(slide, Inches(5.85), Inches(2.5), Inches(5.85), Inches(3.0), dashed=True, thickness=1.5)

# Kiểm soát (trên Phụ thuộc)
tks = "BIẾN KIỂM SOÁT\n(ĐẶC ĐIỂM KINH TẾ - XÃ HỘI)\n\n• Quy mô nền kinh tế\n• Đặc thù ngành nghề\n• Nhân khẩu học"
add_box(slide, tks, Inches(7.9), Inches(1.2), Inches(3.0), Inches(1.3), font_size=10, fill_color=(250, 250, 250))
# Arrow down to Phụ thuộc
add_line_arrow(slide, Inches(9.4), Inches(2.5), Inches(9.4), Inches(3.0), dashed=True, thickness=1.5)


# --- MIDDLE ROW: CÁC BIẾN CHÍNH ---
# B1 (Độc lập)
t2 = "BIẾN ĐỘC LẬP\n(QUẢN LÝ NHÀ NƯỚC VỀ SHTT)\n\n• Xây dựng chiến lược & pháp luật\n• Tổ chức thực thi chính sách\n• Quản lý xác lập & bảo vệ TSTT\n• Thanh tra, giám sát & xử lý"
add_box(slide, t2, Inches(0.8), Inches(3.0), Inches(3.0), Inches(1.9), font_size=10, fill_color=(250, 250, 250))

# B2 (Trung gian)
t3 = "BIẾN TRUNG GIAN\n(SỰ GIA TĂNG TÀI SẢN TRÍ TUỆ)\n\n• Tăng trưởng văn bằng bảo hộ\n• Tỷ lệ thương mại hóa & chuyển giao\n• An toàn môi trường sở hữu trí tuệ"
add_box(slide, t3, Inches(4.35), Inches(3.0), Inches(3.0), Inches(1.9), font_size=10, fill_color=(250, 250, 250))

# B3 (Phụ thuộc)
t4 = "BIẾN PHỤ THUỘC\n(PHÁT TRIỂN KINH TẾ SỐ)\n\n• Thúc đẩy Đổi mới sáng tạo\n• Tăng trưởng kinh tế tri thức\n• Nâng cao năng lực cạnh tranh"
add_box(slide, t4, Inches(7.9), Inches(3.0), Inches(3.0), Inches(1.9), font_size=10, fill_color=(250, 250, 250))


# --- ARROWS BETWEEN CORE BOXES (Horizontal) ---
add_line_arrow(slide, Inches(3.8), Inches(3.95), Inches(4.35), Inches(3.95), thickness=2.0) # B1 -> B2
add_line_arrow(slide, Inches(7.35), Inches(3.95), Inches(7.9), Inches(3.95), thickness=2.0) # B2 -> B3


# --- BOTTOM ROW: PHẢN HỒI ---
tfb = "PHẢN HỒI CHÍNH SÁCH\nThực tiễn kinh tế ➔ Đánh giá ➔ Điều chỉnh ➔ Hoàn thiện Quản lý nhà nước"
# Hộp phản hồi kéo dài từ B1 sang B3
add_box(slide, tfb, Inches(0.8), Inches(5.6), Inches(10.1), Inches(0.6), font_size=10, fill_color=(255, 255, 255))

# Mũi tên từ Phụ thuộc (B3) thả thẳng xuống Phản hồi
add_line_arrow(slide, Inches(9.4), Inches(4.9), Inches(9.4), Inches(5.6), dashed=True, arrow_end=True, thickness=1.5)

# Mũi tên từ Phản hồi đâm ngược thẳng lên Độc lập (B1)
add_line_arrow(slide, Inches(2.3), Inches(5.6), Inches(2.3), Inches(4.9), dashed=True, arrow_end=True, thickness=1.5)


prs.save("KhungLyThuyet_KinhTe_SieuGonGang.pptx")
print("Saved ultra-clean non-cluttered layout.")
