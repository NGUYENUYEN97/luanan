from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_AUTO_SIZE
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls

prs = Presentation()
prs.slide_width = Inches(11.69)
prs.slide_height = Inches(8.27)
slide = prs.slides.add_slide(prs.slide_layouts[6])

def add_box(slide, text, left, top, width, height, font_size=10, bold_first_line=True, bold_bracket_lines=False):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(255, 255, 255)
    shape.line.color.rgb = RGBColor(0, 0, 0)
    shape.line.width = Pt(1.5)
    
    tf = shape.text_frame
    tf.text = text
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    
    for i, p in enumerate(tf.paragraphs):
        if "•" in p.text or p.text.startswith("1.") or p.text.startswith("2.") or p.text.startswith("3.") or p.text.startswith("4.") or p.text.startswith("["):
            p.alignment = PP_ALIGN.LEFT
        else:
            p.alignment = PP_ALIGN.CENTER
            
        for run in p.runs:
            run.font.name = 'Times New Roman'
            run.font.size = Pt(font_size)
            run.font.color.rgb = RGBColor(0, 0, 0)
            
            if bold_first_line and (i == 0 or (i == 1 and p.text.isupper())):
                run.font.bold = True
            if bold_bracket_lines and p.text.startswith("["):
                run.font.bold = True
                run.font.italic = True
    return shape

def add_line_arrow(slide, begin_x, begin_y, end_x, end_y, dashed=False, arrow_end=True, arrow_start=False):
    connector = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, begin_x, begin_y, end_x, end_y)
    connector.line.color.rgb = RGBColor(0,0,0)
    connector.line.width = Pt(1.5)
    if dashed:
        connector.line.dash_style = 4
        
    ln = connector.element.spPr.ln
    if ln is not None:
        if arrow_end:
            tailEnd = parse_xml(r'<a:tailEnd type="triangle" w="med" len="med" %s/>' % nsdecls('a'))
            ln.append(tailEnd)
        if arrow_start:
            headEnd = parse_xml(r'<a:headEnd type="triangle" w="med" len="med" %s/>' % nsdecls('a'))
            ln.append(headEnd)
    return connector


# Title
t_title = slide.shapes.add_textbox(Inches(0.5), Inches(0.05), Inches(10.69), Inches(0.4))
t_title.text_frame.text = "KHUNG LÝ THUYẾT NGHIÊN CỨU"
t_title.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
t_title.text_frame.paragraphs[0].runs[0].font.name = 'Times New Roman'
t_title.text_frame.paragraphs[0].runs[0].font.size = Pt(16)
t_title.text_frame.paragraphs[0].runs[0].font.bold = True


# --- LÝ THUYẾT NỀN TẢNG ---
t1 = "LÝ THUYẾT NỀN TẢNG\nLý thuyết quản lý công  •  Lý thuyết quản lý công mới  •  Lý thuyết quản trị số  •  Lý thuyết quản trị thích ứng  •  Lý thuyết thể chế  •  Lý thuyết về quyền sở hữu trí tuệ"
add_box(slide, t1, Inches(1.5), Inches(0.4), Inches(8.69), Inches(0.6), font_size=10)

add_line_arrow(slide, Inches(5.845), Inches(1.0), Inches(5.845), Inches(1.2))

# --- KHÁI NIỆM CỐT LÕI ---
t2 = "KHÁI NIỆM CỐT LÕI\nQuyền sở hữu trí tuệ  •  Quản lý nhà nước  •  Kinh tế số  •  Quản trị số  •  Quản lý nhà nước về quyền SHTT"
add_box(slide, t2, Inches(1.5), Inches(1.2), Inches(8.69), Inches(0.5), font_size=10)

add_line_arrow(slide, Inches(5.845), Inches(1.7), Inches(5.845), Inches(1.9))

# --- CÁC YẾU TỐ ĐẦU VÀO (Left) ---
t3 = "CÁC YẾU TỐ ĐẦU VÀO\n\n• Nhóm thể chế: Pháp luật, Chính sách, Phối hợp\n• Nhóm nguồn lực: Nguồn nhân lực, Tài chính, Hạ tầng\n• Nhóm công nghệ: Dữ liệu số, AI, Nền tảng số"
add_box(slide, t3, Inches(0.3), Inches(2.2), Inches(3.3), Inches(1.6), font_size=10)

# Mũi tên từ Đầu vào sang Quá trình (Center)
add_line_arrow(slide, Inches(3.6), Inches(3.0), Inches(4.0), Inches(3.0))


# --- QUÁ TRÌNH QUẢN LÝ (Center Combined Block) ---
# Tăng chiều cao để chứa đủ 4 nội dung gốc
t4 = "QUẢN LÝ NHÀ NƯỚC VỀ QUYỀN SHTT TRONG NỀN KINH TẾ SỐ\n\n[Nội dung thực hiện]\n1. Xây dựng chiến lược, chính sách và pháp luật\n2. Tổ chức thực hiện\n3. Hỗ trợ xác lập, khai thác, chuyển giao, thương mại hóa TSTT\n4. Thanh tra, kiểm tra, giám sát và thực thi quyền SHTT\n\n[Năng lực thực thi]\n• Năng lực hoạch định và tổ chức thực hiện\n• Năng lực phối hợp liên ngành và chuyển đổi số"
add_box(slide, t4, Inches(4.0), Inches(1.9), Inches(4.2), Inches(2.7), font_size=10, bold_bracket_lines=True)

add_line_arrow(slide, Inches(6.1), Inches(4.6), Inches(6.1), Inches(5.0))


# --- CHẤT LƯỢNG (Center) ---
t6 = "CHẤT LƯỢNG QUẢN LÝ NHÀ NƯỚC VỀ QUYỀN SHTT\n\n• Hiệu lực   • Hiệu quả   • Tính phù hợp   • Tính đồng bộ   • Tính minh bạch"
add_box(slide, t6, Inches(4.3), Inches(5.0), Inches(3.6), Inches(0.9), font_size=10)

add_line_arrow(slide, Inches(6.1), Inches(5.9), Inches(6.1), Inches(6.3))


# --- KẾT QUẢ/TÁC ĐỘNG (Center) ---
t7 = "KẾT QUẢ VÀ TÁC ĐỘNG\n\n• Bảo vệ hiệu quả quyền SHTT       • Thúc đẩy đổi mới sáng tạo\n• Phát triển tài sản trí tuệ         • Nâng cao năng lực cạnh tranh quốc gia\n• Thúc đẩy phát triển kinh tế số"
add_box(slide, t7, Inches(4.1), Inches(6.3), Inches(4.0), Inches(1.0), font_size=10)


# --- YẾU TỐ ẢNH HƯỞNG (Right) ---
tdt = "CÁC YẾU TỐ ẢNH HƯỞNG (ĐIỀU TIẾT)\n\n• Mức độ phát triển kinh tế số\n• Năng lực đổi mới sáng tạo\n• Hội nhập quốc tế\n• Nhận thức xã hội\n• Mức độ cạnh tranh thị trường"
add_box(slide, tdt, Inches(8.4), Inches(3.1), Inches(3.0), Inches(1.8), font_size=10)

# Mũi tên điều tiết nét đứt chỉ vào mũi tên Center -> Quality (Y=4.8)
add_line_arrow(slide, Inches(8.4), Inches(4.8), Inches(6.1), Inches(4.8), dashed=True)


# --- FEEDBACK LOOP ---
tfb = "PHẢN HỒI CHÍNH SÁCH\nKết quả thực hiện  ➔  Đánh giá  ➔  Điều chỉnh chính sách, thể chế, công cụ  ➔  Hoàn thiện QLNN"
add_box(slide, tfb, Inches(1.5), Inches(7.5), Inches(8.69), Inches(0.5), font_size=10)

# Từ Kết quả xuống Feedback
add_line_arrow(slide, Inches(6.1), Inches(7.3), Inches(6.1), Inches(7.5), dashed=True)

# Từ khối Feedback vòng lên Quá trình
add_line_arrow(slide, Inches(1.5), Inches(7.75), Inches(0.15), Inches(7.75), dashed=True, arrow_end=False)
add_line_arrow(slide, Inches(0.15), Inches(7.75), Inches(0.15), Inches(1.75), dashed=True, arrow_end=False)
add_line_arrow(slide, Inches(0.15), Inches(1.75), Inches(4.8), Inches(1.75), dashed=True, arrow_end=False)
add_line_arrow(slide, Inches(4.8), Inches(1.75), Inches(4.8), Inches(1.9), dashed=True)

prs.save("KhungLyThuyet_GS_BW_ChuanXac_PhucHoiND.pptx")
print("Saved PPTX with fully restored central content.")
