from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

prs = Presentation()
slide = prs.slides.add_slide(prs.slide_layouts[6])

def add_box(slide, text, left, top, width, height, font_size=11, bold=False):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(255, 255, 255)
    shape.line.color.rgb = RGBColor(0, 0, 0)
    shape.line.width = Pt(1.5)
    
    text_frame = shape.text_frame
    text_frame.text = text
    text_frame.word_wrap = True
    
    for paragraph in text_frame.paragraphs:
        paragraph.alignment = PP_ALIGN.CENTER
        for run in paragraph.runs:
            run.font.name = 'Times New Roman'
            run.font.size = Pt(font_size)
            run.font.color.rgb = RGBColor(0, 0, 0)
            run.font.bold = bold
    return shape

def add_arrow(slide, left, top, width, height, direction='DOWN'):
    shape_type = MSO_SHAPE.DOWN_ARROW
    if direction == 'RIGHT': shape_type = MSO_SHAPE.RIGHT_ARROW
    elif direction == 'UP': shape_type = MSO_SHAPE.UP_ARROW
    
    shape = slide.shapes.add_shape(shape_type, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(255, 255, 255)
    shape.line.color.rgb = RGBColor(0, 0, 0)
    return shape

# Title
title_shape = slide.shapes.add_textbox(Inches(0.5), Inches(0.1), Inches(9), Inches(0.5))
tf = title_shape.text_frame
tf.text = "KHUNG LÝ THUYẾT NGHIÊN CỨU"
tf.paragraphs[0].alignment = PP_ALIGN.CENTER
tf.paragraphs[0].runs[0].font.name = 'Times New Roman'
tf.paragraphs[0].runs[0].font.size = Pt(16)
tf.paragraphs[0].runs[0].font.bold = True

# --- THEORIES SECTION ---
theory_container = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.2), Inches(0.7), Inches(9.6), Inches(1.3))
theory_container.fill.solid()
theory_container.fill.fore_color.rgb = RGBColor(255, 255, 255)
theory_container.line.color.rgb = RGBColor(0, 0, 0)
theory_container.line.width = Pt(1.5)

t_th = slide.shapes.add_textbox(Inches(0.5), Inches(0.75), Inches(9.0), Inches(0.3))
t_th.text_frame.text = "CƠ SỞ LÝ THUYẾT NỀN"
t_th.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
t_th.text_frame.paragraphs[0].runs[0].font.name = 'Times New Roman'
t_th.text_frame.paragraphs[0].runs[0].font.size = Pt(11)
t_th.text_frame.paragraphs[0].runs[0].font.bold = True

add_box(slide, "Lý thuyết\nquản lý nhà nước", Inches(0.4), Inches(1.1), Inches(1.7), Inches(0.7), font_size=10, bold=True)
add_box(slide, "Lý thuyết\nquản trị công", Inches(2.2), Inches(1.1), Inches(1.7), Inches(0.7), font_size=10, bold=True)
add_box(slide, "Lý thuyết\nthể chế", Inches(4.0), Inches(1.1), Inches(1.7), Inches(0.7), font_size=10, bold=True)
add_box(slide, "Lý thuyết\nkinh tế số", Inches(5.8), Inches(1.1), Inches(1.7), Inches(0.7), font_size=10, bold=True)
add_box(slide, "Lý thuyết\nđổi mới sáng tạo", Inches(7.6), Inches(1.1), Inches(1.7), Inches(0.7), font_size=10, bold=True)

# Dashed box for Model
model_container = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.2), Inches(2.2), Inches(9.6), Inches(4.7))
model_container.fill.background()
model_container.line.color.rgb = RGBColor(0, 0, 0)
model_container.line.width = Pt(1.5)
model_container.line.dash_style = 4
slide.shapes._spTree.remove(model_container._element)
slide.shapes._spTree.insert(2, model_container._element)

t_md = slide.shapes.add_textbox(Inches(0.3), Inches(2.25), Inches(4.0), Inches(0.3))
t_md.text_frame.text = "MÔ HÌNH NHÂN QUẢ GIỮA CÁC BIẾN"
t_md.text_frame.paragraphs[0].runs[0].font.name = 'Times New Roman'
t_md.text_frame.paragraphs[0].runs[0].font.size = Pt(11)
t_md.text_frame.paragraphs[0].runs[0].font.bold = True
t_md.text_frame.paragraphs[0].runs[0].font.italic = True


# --- VARIABLES SECTION ---

# Independent Variable
text_nd = "NỘI DUNG QUẢN LÝ NHÀ NƯỚC\n(Biến độc lập)\n\n1. Xây dựng chiến lược, chính sách, pháp luật\n2. Tổ chức thực hiện\n3. Hỗ trợ xác lập, khai thác, thương mại hóa\n4. Thanh tra, kiểm tra, giám sát, thực thi quyền"
add_box(slide, text_nd, Inches(0.5), Inches(3.6), Inches(3.6), Inches(1.8), font_size=11, bold=True)

# Dependent Variable
text_cl = "CHẤT LƯỢNG QUẢN LÝ NHÀ NƯỚC\n(Biến phụ thuộc)\n\n1. Hiệu lực\n2. Hiệu quả\n3. Tính phù hợp\n4. Tính đồng bộ\n5. Tính minh bạch"
add_box(slide, text_cl, Inches(5.8), Inches(3.6), Inches(3.6), Inches(1.8), font_size=11, bold=True)

# Main Arrow
add_arrow(slide, Inches(4.1), Inches(4.3), Inches(1.7), Inches(0.4), 'RIGHT')

# Influencing Variables
add_box(slide, "CÁC YẾU TỐ BÊN TRONG", Inches(6.3), Inches(2.5), Inches(2.6), Inches(0.6), font_size=11, bold=True)
add_arrow(slide, Inches(7.45), Inches(3.1), Inches(0.3), Inches(0.5), 'DOWN')

add_box(slide, "CÁC YẾU TỐ BÊN NGOÀI", Inches(6.3), Inches(5.9), Inches(2.6), Inches(0.6), font_size=11, bold=True)
add_arrow(slide, Inches(7.45), Inches(5.4), Inches(0.3), Inches(0.5), 'UP')


prs.save("KhungLyThuyetNghienCuu_GS_Final.pptx")
print("Saved final theory framework pptx.")
