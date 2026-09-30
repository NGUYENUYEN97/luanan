from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_AUTO_SIZE
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls

prs = Presentation()
prs.slide_width = Inches(11.69)
prs.slide_height = Inches(8.27)
slide = prs.slides.add_slide(prs.slide_layouts[6])

def add_box(slide, text, left, top, width, height, font_size=11, fill_color=(255, 255, 255), border_color=(0,0,0), bold=True):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(*fill_color)
    shape.line.color.rgb = RGBColor(*border_color)
    shape.line.width = Pt(1.5)
    
    tf = shape.text_frame
    tf.text = text
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    
    for i, p in enumerate(tf.paragraphs):
        p.alignment = PP_ALIGN.CENTER
        for run in p.runs:
            run.font.name = 'Times New Roman'
            run.font.size = Pt(font_size)
            run.font.color.rgb = RGBColor(0, 0, 0)
            if bold and i == 0:
                run.font.bold = True
    return shape

def add_line_arrow(slide, begin_x, begin_y, end_x, end_y, dashed=False, arrow_end=True, arrow_start=False, thickness=1.5):
    connector = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, begin_x, begin_y, end_x, end_y)
    connector.line.color.rgb = RGBColor(0,0,0)
    connector.line.width = Pt(thickness)
    if dashed:
        connector.line.dash_style = 4
        
    ln = connector.element.spPr.ln
    if ln is not None:
        if arrow_end:
            tailEnd = parse_xml(r'<a:tailEnd type="triangle" w="med" len="med" %s/>' % nsdecls('a'))
            ln.append(tailEnd)
        if arrow_start:
            headEnd = parse_xml(r'<a:headEnd type="triangle" w="med" len="med" %s/>' % nsdecls('a'))
            ln.append(headEnd)
    return connector

# TITLE
t_title = slide.shapes.add_textbox(Inches(0.5), Inches(0.1), Inches(10.69), Inches(0.5))
t_title.text_frame.text = "MÔ HÌNH NGHIÊN CỨU ĐỊNH LƯỢNG ĐỀ TÀI\nQuản lý nhà nước về quyền sở hữu trí tuệ trong nền kinh tế số"
t_title.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
t_title.text_frame.paragraphs[0].runs[0].font.name = 'Times New Roman'
t_title.text_frame.paragraphs[0].runs[0].font.size = Pt(15)
t_title.text_frame.paragraphs[0].runs[0].font.bold = True
t_title.text_frame.paragraphs[1].alignment = PP_ALIGN.CENTER
t_title.text_frame.paragraphs[1].runs[0].font.name = 'Times New Roman'
t_title.text_frame.paragraphs[1].runs[0].font.size = Pt(12)
t_title.text_frame.paragraphs[1].runs[0].font.italic = True

# THEORY BOX
t_theory = "LÝ THUYẾT NỀN TẢNG\nLý thuyết quản lý công   •   Lý thuyết thể chế   •   Lý thuyết quyền SHTT   •   Lý thuyết KT số"
add_box(slide, t_theory, Inches(1.5), Inches(0.7), Inches(8.69), Inches(0.6), font_size=11, fill_color=(245, 245, 245))


# 6 INDEPENDENT VARIABLES
ivs = [
    "Khung pháp lý",
    "Năng lực cán bộ",
    "Nguồn lực tài chính",
    "Hạ tầng số",
    "Phối hợp thể chế",
    "Nhận thức xã hội"
]

iv_boxes = []
start_y = 1.8
for i, iv in enumerate(ivs):
    y = start_y + i * 0.9
    shape = add_box(slide, iv, Inches(0.5), Inches(y), Inches(2.6), Inches(0.6), font_size=11, bold=False)
    iv_boxes.append(y + 0.3) # Store center Y for arrows
    
# Nhãn nhóm Biến Độc Lập
lbl_iv = slide.shapes.add_textbox(Inches(0.5), Inches(1.4), Inches(2.6), Inches(0.4))
lbl_iv.text_frame.text = "CÁC BIẾN ĐỘC LẬP"
lbl_iv.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
lbl_iv.text_frame.paragraphs[0].runs[0].font.name = 'Times New Roman'
lbl_iv.text_frame.paragraphs[0].runs[0].font.bold = True


# MEDIATING VARIABLE
t_mv = "BIẾN TRUNG GIAN\n\nMức độ chuyển đổi số trong quản lý"
add_box(slide, t_mv, Inches(4.5), Inches(3.9), Inches(2.6), Inches(1.2), font_size=11)
mv_center_y = 4.5


# DEPENDENT VARIABLE
t_dv = "BIẾN PHỤ THUỘC\n\nHiệu quả quản lý nhà nước về quyền SHTT"
add_box(slide, t_dv, Inches(8.5), Inches(3.9), Inches(2.6), Inches(1.2), font_size=11)


# ARROWS FROM IV TO MV
for y in iv_boxes:
    add_line_arrow(slide, Inches(3.1), Inches(y), Inches(4.5), Inches(mv_center_y), thickness=1.5)
    

# ARROW FROM MV TO DV
add_line_arrow(slide, Inches(7.1), Inches(mv_center_y), Inches(8.5), Inches(mv_center_y), thickness=2.0)


# MODERATING VARIABLE
t_mod = "BIẾN ĐIỀU TIẾT\n\nMức độ phát triển kinh tế số"
add_box(slide, t_mod, Inches(6.5), Inches(2.0), Inches(2.6), Inches(1.0), font_size=11)
# Arrow Mod -> MV-DV path
add_line_arrow(slide, Inches(7.8), Inches(3.0), Inches(7.8), Inches(4.5), dashed=True, thickness=1.5)


# CONTROL VARIABLE
t_ctrl = "BIẾN KIỂM SOÁT\n\n• Loại địa phương\n• Quy mô kinh tế\n• Ngành nghề"
shape_ctrl = add_box(slide, t_ctrl, Inches(8.5), Inches(5.8), Inches(2.6), Inches(1.3), font_size=11)
# Align bullet points in Control variable left
tf = shape_ctrl.text_frame
for i, p in enumerate(tf.paragraphs):
    if p.text.startswith("•"):
        p.alignment = PP_ALIGN.LEFT
        p.level = 0
# Arrow Control -> DV
add_line_arrow(slide, Inches(9.8), Inches(5.8), Inches(9.8), Inches(5.1), dashed=True, thickness=1.5)


# ARROWS FROM THEORY
add_line_arrow(slide, Inches(1.8), Inches(1.3), Inches(1.8), Inches(1.8), dashed=True)
add_line_arrow(slide, Inches(5.8), Inches(1.3), Inches(5.8), Inches(3.9), dashed=True)
add_line_arrow(slide, Inches(9.8), Inches(1.3), Inches(9.8), Inches(3.9), dashed=True)


prs.save("MoHinhNghienCuu_DinhLuong.pptx")
print("Saved quantitative SEM model.")
