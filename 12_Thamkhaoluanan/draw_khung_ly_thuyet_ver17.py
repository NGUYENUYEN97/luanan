from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_AUTO_SIZE
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls

prs = Presentation()
prs.slide_width = Inches(11.69)
prs.slide_height = Inches(8.27)
slide = prs.slides.add_slide(prs.slide_layouts[6])

def add_box(slide, text, left, top, width, height, font_size=10, bold_first_line=True):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(255, 255, 255)
    shape.line.color.rgb = RGBColor(0, 0, 0)
    shape.line.width = Pt(1.5)
    
    tf = shape.text_frame
    tf.text = text
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    
    for i, p in enumerate(tf.paragraphs):
        if "•" in p.text or p.text.startswith("1.") or p.text.startswith("2.") or p.text.startswith("3.") or p.text.startswith("4."):
            p.alignment = PP_ALIGN.LEFT
        else:
            p.alignment = PP_ALIGN.CENTER
            
        for run in p.runs:
            run.font.name = 'Times New Roman'
            run.font.size = Pt(font_size)
            run.font.color.rgb = RGBColor(0, 0, 0)
            
            if bold_first_line and (i == 0 or (i == 1 and p.text.isupper())):
                run.font.bold = True
    return shape

def add_line_arrow(slide, begin_x, begin_y, end_x, end_y, dashed=False, arrow_end=True, arrow_start=False):
    connector = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, begin_x, begin_y, end_x, end_y)
    connector.line.color.rgb = RGBColor(0,0,0)
    connector.line.width = Pt(1.5)
    if dashed:
        connector.line.dash_style = 4
        
    ln = connector.element.spPr.ln
    if ln is not None:
        if arrow_end:
            tailEnd = parse_xml(r'<a:tailEnd type="triangle" w="med" len="med" %s/>' % nsdecls('a'))
            ln.append(tailEnd)
        if arrow_start:
            headEnd = parse_xml(r'<a:headEnd type="triangle" w="med" len="med" %s/>' % nsdecls('a'))
            ln.append(headEnd)
    return connector

# Title
t_title = slide.shapes.add_textbox(Inches(0.5), Inches(0.05), Inches(10.69), Inches(0.4))
t_title.text_frame.text = "KHUNG LÝ THUYẾT NGHIÊN CỨU (MÔ HÌNH NHÂN QUẢ)"
t_title.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
t_title.text_frame.paragraphs[0].runs[0].font.name = 'Times New Roman'
t_title.text_frame.paragraphs[0].runs[0].font.size = Pt(16)
t_title.text_frame.paragraphs[0].runs[0].font.bold = True

# --- BỐI CẢNH NỀN KINH TẾ SỐ ---
ctx_shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.15), Inches(1.15), Inches(11.39), Inches(6.0))
ctx_shape.fill.solid()
ctx_shape.fill.fore_color.rgb = RGBColor(255, 255, 255)
ctx_shape.line.color.rgb = RGBColor(0, 0, 0)
ctx_shape.line.width = Pt(1.5)
ctx_shape.line.dash_style = 4

lbl = slide.shapes.add_textbox(Inches(0.25), Inches(1.2), Inches(5.0), Inches(0.4))
lbl.text_frame.text = "BỐI CẢNH: NỀN KINH TẾ SỐ (Kỷ nguyên dữ liệu & AI)"
lbl.text_frame.paragraphs[0].runs[0].font.name = 'Times New Roman'
lbl.text_frame.paragraphs[0].runs[0].font.size = Pt(12)
lbl.text_frame.paragraphs[0].runs[0].font.bold = True
lbl.text_frame.paragraphs[0].runs[0].font.italic = True


# --- LÝ THUYẾT NỀN TẢNG ---
t1 = "LÝ THUYẾT NỀN TẢNG\nLý thuyết quản lý công  •  Lý thuyết thể chế  •  Lý thuyết tăng trưởng nội sinh  •  Lý thuyết về quyền sở hữu trí tuệ"
add_box(slide, t1, Inches(1.5), Inches(0.45), Inches(8.69), Inches(0.5), font_size=10)


# --- BIẾN ĐỘC LẬP (Left) ---
t2 = "BIẾN ĐỘC LẬP (NGUYÊN NHÂN)\nQUẢN LÝ NHÀ NƯỚC VỀ QUYỀN SHTT\n\n1. Xây dựng chiến lược, chính sách, pháp luật\n2. Tổ chức thực hiện\n3. Quản lý hoạt động xác lập và bảo vệ TSTT\n4. Thanh tra, kiểm tra, giám sát, thực thi quyền"
add_box(slide, t2, Inches(0.3), Inches(3.0), Inches(3.5), Inches(2.0), font_size=10)


# --- BIẾN TRUNG GIAN (Center) ---
t3 = "BIẾN TRUNG GIAN (CƠ CHẾ TRUYỀN DẪN)\nSỰ GIA TĂNG & THƯƠNG MẠI HÓA TÀI SẢN TRÍ TUỆ\n\n• Tăng trưởng số lượng văn bằng bảo hộ\n• Tỷ lệ thương mại hóa và chuyển giao\n• Đảm bảo an toàn môi trường Sở hữu trí tuệ"
add_box(slide, t3, Inches(4.2), Inches(3.2), Inches(3.5), Inches(1.6), font_size=10)


# --- BIẾN PHỤ THUỘC (Right) ---
t4 = "BIẾN PHỤ THUỘC (KẾT QUẢ)\nPHÁT TRIỂN KINH TẾ SỐ\n\n• Thúc đẩy Đổi mới sáng tạo toàn diện\n• Tăng trưởng kinh tế tri thức\n• Nâng cao năng lực cạnh tranh quốc gia"
add_box(slide, t4, Inches(8.1), Inches(3.2), Inches(3.3), Inches(1.6), font_size=10)


# --- ARROWS: Độc lập -> Trung gian -> Phụ thuộc ---
add_line_arrow(slide, Inches(3.8), Inches(4.0), Inches(4.2), Inches(4.0)) # T2 -> T3
add_line_arrow(slide, Inches(7.7), Inches(4.0), Inches(8.1), Inches(4.0)) # T3 -> T4


# --- BIẾN ĐIỀU TIẾT (Top Center) ---
tdt = "BIẾN ĐIỀU TIẾT (MỨC ĐỘ TÁC ĐỘNG)\nCÁC YẾU TỐ MÔI TRƯỜNG & HỘI NHẬP\n\n• Trình độ công nghệ (AI, Big Data)\n• Mức độ hội nhập quốc tế\n• Nhận thức pháp luật của xã hội"
add_box(slide, tdt, Inches(3.7), Inches(1.3), Inches(3.5), Inches(1.4), font_size=10)

# Mũi tên điều tiết chỉ vào mũi tên Độc lập -> Trung gian (X=4.0, Y=4.0)
add_line_arrow(slide, Inches(4.8), Inches(2.7), Inches(4.0), Inches(4.0), dashed=True)
# Mũi tên điều tiết chỉ vào mũi tên Trung gian -> Phụ thuộc (X=7.9, Y=4.0)
add_line_arrow(slide, Inches(6.1), Inches(2.7), Inches(7.9), Inches(4.0), dashed=True)


# --- BIẾN KIỂM SOÁT (Bottom Right) ---
tks = "BIẾN KIỂM SOÁT (CỐ ĐỊNH)\nĐẶC ĐIỂM KINH TẾ - XÃ HỘI\n\n• Quy mô và cơ cấu nền kinh tế\n• Đặc thù ngành nghề (Công nghệ, Dịch vụ)\n• Đặc điểm nhân khẩu học"
add_box(slide, tks, Inches(7.8), Inches(5.3), Inches(3.5), Inches(1.4), font_size=10)

# Mũi tên kiểm soát chỉ ngược lên Biến phụ thuộc
add_line_arrow(slide, Inches(9.5), Inches(5.3), Inches(9.5), Inches(4.8), dashed=True)


# --- PHẢN HỒI CHÍNH SÁCH ---
tfb = "PHẢN HỒI CHÍNH SÁCH\nThực tiễn kinh tế số ➔ Đánh giá ➔ Điều chỉnh ➔ Hoàn thiện QLNN"
add_box(slide, tfb, Inches(1.5), Inches(7.35), Inches(8.69), Inches(0.5), font_size=10)

# Từ Phụ thuộc xuống Phản hồi
add_line_arrow(slide, Inches(9.7), Inches(4.8), Inches(9.7), Inches(7.6), dashed=True, arrow_end=False)
add_line_arrow(slide, Inches(9.7), Inches(7.6), Inches(10.19), Inches(7.6), dashed=True)

# Từ Phản hồi vòng về Độc lập
add_line_arrow(slide, Inches(1.5), Inches(7.6), Inches(0.05), Inches(7.6), dashed=True, arrow_end=False)
add_line_arrow(slide, Inches(0.05), Inches(7.6), Inches(0.05), Inches(4.0), dashed=True, arrow_end=False)
add_line_arrow(slide, Inches(0.05), Inches(4.0), Inches(0.3), Inches(4.0), dashed=True)


prs.save("KhungLyThuyet_KinhTe_5Bien_HoanHao.pptx")
print("Saved PPTX with strict 5-variable SEM layout.")
