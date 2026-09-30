from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_AUTO_SIZE
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls

prs = Presentation()
prs.slide_width = Inches(11.69)
prs.slide_height = Inches(8.27)
slide = prs.slides.add_slide(prs.slide_layouts[6])

def add_box(slide, text, left, top, width, height, font_size=10, bold_first_line=True):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(255, 255, 255)
    shape.line.color.rgb = RGBColor(0, 0, 0)
    shape.line.width = Pt(1.5)
    
    tf = shape.text_frame
    tf.text = text
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    
    for i, p in enumerate(tf.paragraphs):
        if "•" in p.text or p.text.startswith("1.") or p.text.startswith("2.") or p.text.startswith("3.") or p.text.startswith("4."):
            p.alignment = PP_ALIGN.LEFT
        else:
            p.alignment = PP_ALIGN.CENTER
            
        for run in p.runs:
            run.font.name = 'Times New Roman'
            run.font.size = Pt(font_size)
            run.font.color.rgb = RGBColor(0, 0, 0)
            
            if bold_first_line and (i == 0 or (i == 1 and p.text.isupper())):
                run.font.bold = True
    return shape

def add_line_arrow(slide, begin_x, begin_y, end_x, end_y, dashed=False, arrow_end=True, arrow_start=False):
    connector = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, begin_x, begin_y, end_x, end_y)
    connector.line.color.rgb = RGBColor(0,0,0)
    connector.line.width = Pt(1.5)
    if dashed:
        connector.line.dash_style = 4
        
    ln = connector.element.spPr.ln
    if ln is not None:
        if arrow_end:
            tailEnd = parse_xml(r'<a:tailEnd type="triangle" w="med" len="med" %s/>' % nsdecls('a'))
            ln.append(tailEnd)
        if arrow_start:
            headEnd = parse_xml(r'<a:headEnd type="triangle" w="med" len="med" %s/>' % nsdecls('a'))
            ln.append(headEnd)
    return connector

# Title
t_title = slide.shapes.add_textbox(Inches(0.5), Inches(0.05), Inches(10.69), Inches(0.4))
t_title.text_frame.text = "KHUNG LÝ THUYẾT NGHIÊN CỨU"
t_title.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
t_title.text_frame.paragraphs[0].runs[0].font.name = 'Times New Roman'
t_title.text_frame.paragraphs[0].runs[0].font.size = Pt(16)
t_title.text_frame.paragraphs[0].runs[0].font.bold = True

# --- LÝ THUYẾT NỀN TẢNG ---
t1 = "LÝ THUYẾT NỀN TẢNG\nLý thuyết quản lý công  •  Lý thuyết quản lý công mới  •  Lý thuyết quản trị số  •  Lý thuyết quản trị thích ứng  •  Lý thuyết thể chế  •  Lý thuyết về quyền sở hữu trí tuệ"
add_box(slide, t1, Inches(1.5), Inches(0.5), Inches(8.69), Inches(0.6), font_size=10)
# Arrow down to the center box
add_line_arrow(slide, Inches(6.0), Inches(1.1), Inches(6.0), Inches(1.5))


# --- BIẾN ĐỘC LẬP (Left) ---
t2 = "BIẾN ĐỘC LẬP\n(CÁC YẾU TỐ ĐẦU VÀO)\n\n• Nhóm thể chế: Pháp luật, Chính sách, Phối hợp\n• Nhóm nguồn lực: Nguồn nhân lực, Tài chính, Hạ tầng\n• Nhóm công nghệ: Dữ liệu số, AI, Nền tảng số"
add_box(slide, t2, Inches(0.4), Inches(2.0), Inches(3.2), Inches(1.6), font_size=10)
# Mũi tên từ Độc lập sang Trung gian
add_line_arrow(slide, Inches(3.6), Inches(2.8), Inches(4.0), Inches(2.8))


# --- BIẾN TRUNG GIAN (Center) ---
t3 = "BIẾN TRUNG GIAN\n(QUẢN LÝ NHÀ NƯỚC VỀ QUYỀN SHTT TRONG NỀN KINH TẾ SỐ)\n\n1. Xây dựng chiến lược, chính sách và pháp luật\n2. Tổ chức thực hiện\n3. Quản lý hoạt động xác lập, bảo vệ và hỗ trợ khai thác, thương mại hóa TSTT\n4. Thanh tra, kiểm tra, giám sát và thực thi quyền SHTT"
add_box(slide, t3, Inches(4.0), Inches(1.5), Inches(4.0), Inches(1.8), font_size=10)
add_line_arrow(slide, Inches(6.0), Inches(3.3), Inches(6.0), Inches(4.0))


# --- BIẾN PHỤ THUỘC (Center) ---
t4 = "BIẾN PHỤ THUỘC\n(CHẤT LƯỢNG QUẢN LÝ NHÀ NƯỚC VỀ QUYỀN SHTT)\n\n• Hiệu lực   • Hiệu quả   • Tính phù hợp   • Tính đồng bộ   • Tính minh bạch"
add_box(slide, t4, Inches(4.1), Inches(4.0), Inches(3.8), Inches(1.0), font_size=10)
add_line_arrow(slide, Inches(6.0), Inches(5.0), Inches(6.0), Inches(5.5))


# --- KẾT QUẢ/TÁC ĐỘNG (Center) ---
t5 = "KẾT QUẢ VÀ TÁC ĐỘNG\n\n• Bảo vệ hiệu quả quyền SHTT       • Thúc đẩy đổi mới sáng tạo\n• Phát triển tài sản trí tuệ         • Nâng cao năng lực cạnh tranh quốc gia\n• Thúc đẩy phát triển kinh tế số"
add_box(slide, t5, Inches(4.0), Inches(5.5), Inches(4.0), Inches(1.1), font_size=10)


# --- BIẾN ĐIỀU TIẾT (Right) ---
tdt = "BIẾN ĐIỀU TIẾT\n(CÁC YẾU TỐ ẢNH HƯỞNG)\n\n• Mức độ phát triển kinh tế số\n• Năng lực đổi mới sáng tạo\n• Hội nhập quốc tế\n• Nhận thức xã hội\n• Mức độ cạnh tranh thị trường"
add_box(slide, tdt, Inches(8.4), Inches(2.2), Inches(2.9), Inches(1.8), font_size=10)

# Mũi tên điều tiết nét đứt chỉ vào mũi tên Độc lập -> Trung gian (Y=2.8) - Optional? Usually it affects the path to Dependent.
# Let's make it affect Trung gian -> Phụ thuộc (Y=3.65)
add_line_arrow(slide, Inches(8.4), Inches(3.65), Inches(6.0), Inches(3.65), dashed=True)


# --- FEEDBACK LOOP ---
tfb = "PHẢN HỒI CHÍNH SÁCH\nKết quả thực hiện  ➔  Đánh giá  ➔  Điều chỉnh chính sách, thể chế, công cụ  ➔  Hoàn thiện QLNN"
add_box(slide, tfb, Inches(1.5), Inches(7.2), Inches(8.69), Inches(0.6), font_size=10)

# Từ Kết quả xuống Feedback
add_line_arrow(slide, Inches(6.0), Inches(6.6), Inches(6.0), Inches(7.2), dashed=True)

# Từ khối Feedback vòng lên Biến Trung gian (T3)
add_line_arrow(slide, Inches(1.5), Inches(7.5), Inches(0.15), Inches(7.5), dashed=True, arrow_end=False)
add_line_arrow(slide, Inches(0.15), Inches(7.5), Inches(0.15), Inches(1.2), dashed=True, arrow_end=False)
add_line_arrow(slide, Inches(0.15), Inches(1.2), Inches(4.5), Inches(1.2), dashed=True, arrow_end=False)
add_line_arrow(slide, Inches(4.5), Inches(1.2), Inches(4.5), Inches(1.5), dashed=True)

prs.save("KhungLyThuyet_GS_BW_PhienBanBien_Moi.pptx")
print("Saved PPTX with full variables requested.")
