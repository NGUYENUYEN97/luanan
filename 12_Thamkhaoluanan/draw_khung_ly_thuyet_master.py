from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

prs = Presentation()
prs.slide_width = Inches(13.33)
prs.slide_height = Inches(7.5)

slide = prs.slides.add_slide(prs.slide_layouts[6])

def add_box(slide, text, left, top, width, height, font_size=10, bold=False):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(255, 255, 255)
    shape.line.color.rgb = RGBColor(0, 0, 0)
    shape.line.width = Pt(1.5)
    
    text_frame = shape.text_frame
    text_frame.text = text
    text_frame.word_wrap = True
    
    for paragraph in text_frame.paragraphs:
        paragraph.alignment = PP_ALIGN.CENTER
        for run in paragraph.runs:
            run.font.name = 'Times New Roman'
            run.font.size = Pt(font_size)
            run.font.color.rgb = RGBColor(0, 0, 0)
            run.font.bold = bold
    return shape

def add_arrow(slide, left, top, width, height, direction='DOWN'):
    shape_type = MSO_SHAPE.DOWN_ARROW
    if direction == 'RIGHT': shape_type = MSO_SHAPE.RIGHT_ARROW
    elif direction == 'UP': shape_type = MSO_SHAPE.UP_ARROW
    elif direction == 'LEFT': shape_type = MSO_SHAPE.LEFT_ARROW
    
    shape = slide.shapes.add_shape(shape_type, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(255, 255, 255)
    shape.line.color.rgb = RGBColor(0, 0, 0)
    return shape

# --- TITLE ---
t_title = slide.shapes.add_textbox(Inches(0.5), Inches(0.1), Inches(12.33), Inches(0.4))
t_title.text_frame.text = "KHUNG LÝ THUYẾT NGHIÊN CỨU"
t_title.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
t_title.text_frame.paragraphs[0].runs[0].font.name = 'Times New Roman'
t_title.text_frame.paragraphs[0].runs[0].font.size = Pt(16)
t_title.text_frame.paragraphs[0].runs[0].font.bold = True

# --- ROW 1: THEORIES ---
t_th = slide.shapes.add_textbox(Inches(0.2), Inches(0.5), Inches(12.93), Inches(0.3))
t_th.text_frame.text = "LÝ THUYẾT NỀN TẢNG"
t_th.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
t_th.text_frame.paragraphs[0].runs[0].font.name = 'Times New Roman'
t_th.text_frame.paragraphs[0].runs[0].font.size = Pt(12)
t_th.text_frame.paragraphs[0].runs[0].font.bold = True

start_x = 0.6
box_w = 1.9
space_w = 0.15
theories = ["Lý thuyết\nquản lý công", "Lý thuyết\nquản lý công mới", "Lý thuyết\nquản trị số", "Lý thuyết\nquản trị thích ứng", "Lý thuyết\nthể chế", "Lý thuyết\nvề sở hữu trí tuệ"]
for i, th in enumerate(theories):
    add_box(slide, th, Inches(start_x + i*(box_w+space_w)), Inches(0.9), Inches(box_w), Inches(0.6), font_size=11, bold=True)

add_arrow(slide, Inches(6.5), Inches(1.5), Inches(0.3), Inches(0.3), 'DOWN')

# --- ROW 2: CORE CONCEPTS ---
t_cc = slide.shapes.add_textbox(Inches(0.2), Inches(1.85), Inches(12.93), Inches(0.3))
t_cc.text_frame.text = "KHÁI NIỆM CỐT LÕI"
t_cc.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
t_cc.text_frame.paragraphs[0].runs[0].font.name = 'Times New Roman'
t_cc.text_frame.paragraphs[0].runs[0].font.size = Pt(12)
t_cc.text_frame.paragraphs[0].runs[0].font.bold = True

start_x2 = 1.4
box_w2 = 1.9
space_w2 = 0.25
concepts = ["Quyền\nsở hữu trí tuệ", "Quản lý\nnhà nước", "Kinh tế số", "Quản trị số", "Quản lý nhà nước về\nquyền sở hữu trí tuệ"]
for i, cc in enumerate(concepts):
    add_box(slide, cc, Inches(start_x2 + i*(box_w2+space_w2)), Inches(2.2), Inches(box_w2), Inches(0.6), font_size=11, bold=True)


add_arrow(slide, Inches(6.5), Inches(2.8), Inches(0.3), Inches(0.4), 'DOWN')

# --- ROW 3: CAUSAL MODEL ---
# Dashed bounding box for the entire model
model_container = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.2), Inches(3.2), Inches(12.93), Inches(4.1))
model_container.fill.background()
model_container.line.color.rgb = RGBColor(0, 0, 0)
model_container.line.width = Pt(1.5)
model_container.line.dash_style = 4
slide.shapes._spTree.remove(model_container._element)
slide.shapes._spTree.insert(2, model_container._element)

t_md = slide.shapes.add_textbox(Inches(0.3), Inches(3.25), Inches(6.0), Inches(0.3))
t_md.text_frame.text = "MÔ HÌNH NHÂN QUẢ VÀ CÁC THÀNH TỐ PHÂN TÍCH"
t_md.text_frame.paragraphs[0].runs[0].font.name = 'Times New Roman'
t_md.text_frame.paragraphs[0].runs[0].font.size = Pt(11)
t_md.text_frame.paragraphs[0].runs[0].font.bold = True
t_md.text_frame.paragraphs[0].runs[0].font.italic = True

# 1. Các yếu tố tác động (Left)
text_factors = "CÁC YẾU TỐ TÁC ĐỘNG\n\n- Thể chế\n- Công nghệ số\n- Trí tuệ nhân tạo (AI)\n- Dữ liệu\n- Nguồn nhân lực\n- Phối hợp liên ngành\n- Hợp tác quốc tế\n- Nhận thức xã hội"
add_box(slide, text_factors, Inches(0.5), Inches(3.7), Inches(2.8), Inches(3.3), font_size=12, bold=True)

add_arrow(slide, Inches(3.3), Inches(5.2), Inches(0.5), Inches(0.3), 'RIGHT')

# 2. Khung phân tích QLNN (Center Left)
text_framework = "KHUNG PHÂN TÍCH QUẢN LÝ NHÀ NƯỚC\n\n1. Chủ thể quản lý\n2. Khách thể quản lý\n3. Mục tiêu quản lý\n4. Nội dung quản lý\n5. Công cụ quản lý"
add_box(slide, text_framework, Inches(3.8), Inches(3.7), Inches(3.2), Inches(3.3), font_size=12, bold=True)

add_arrow(slide, Inches(7.0), Inches(5.2), Inches(0.5), Inches(0.3), 'RIGHT')

# 3. Chất lượng quản lý (Center Right)
text_quality = "CHẤT LƯỢNG QUẢN LÝ\n\n- Tính hiệu lực\n- Tính hiệu quả\n- Tính phù hợp\n- Tính đồng bộ\n- Tính minh bạch"
add_box(slide, text_quality, Inches(7.5), Inches(3.7), Inches(2.5), Inches(3.3), font_size=12, bold=True)

add_arrow(slide, Inches(10.0), Inches(5.2), Inches(0.5), Inches(0.3), 'RIGHT')

# 4. Tác động quản lý (Right)
text_impact = "TÁC ĐỘNG CỦA QUẢN LÝ\n\n- Bảo vệ quyền sở hữu trí tuệ\n- Thúc đẩy đổi mới sáng tạo\n- Phát triển kinh tế số"
add_box(slide, text_impact, Inches(10.5), Inches(3.7), Inches(2.6), Inches(3.3), font_size=12, bold=True)


prs.save("KhungLyThuyetNghienCuu_Master.pptx")
print("Saved Master theoretical framework.")
