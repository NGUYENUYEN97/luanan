from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_AUTO_SIZE
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls

prs = Presentation()
prs.slide_width = Inches(11.69)
prs.slide_height = Inches(8.27)
slide = prs.slides.add_slide(prs.slide_layouts[6])

def add_box(slide, text, left, top, width, height, font_size=11, fill_color=(255, 255, 255), border_color=(0,0,0), bold=True):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(*fill_color)
    shape.line.color.rgb = RGBColor(*border_color)
    shape.line.width = Pt(1.5)
    
    tf = shape.text_frame
    tf.text = text
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    
    for i, p in enumerate(tf.paragraphs):
        p.alignment = PP_ALIGN.CENTER
        for run in p.runs:
            run.font.name = 'Times New Roman'
            run.font.size = Pt(font_size)
            run.font.color.rgb = RGBColor(0, 0, 0)
            if bold and i == 0:
                run.font.bold = True
    return shape

def add_line_arrow(slide, begin_x, begin_y, end_x, end_y, dashed=False, arrow_end=True, arrow_start=False, thickness=1.5):
    connector = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, begin_x, begin_y, end_x, end_y)
    connector.line.color.rgb = RGBColor(0,0,0)
    connector.line.width = Pt(thickness)
    if dashed:
        connector.line.dash_style = 4
        
    ln = connector.element.spPr.ln
    if ln is not None:
        if arrow_end:
            tailEnd = parse_xml(r'<a:tailEnd type="triangle" w="med" len="med" %s/>' % nsdecls('a'))
            ln.append(tailEnd)
        if arrow_start:
            headEnd = parse_xml(r'<a:headEnd type="triangle" w="med" len="med" %s/>' % nsdecls('a'))
            ln.append(headEnd)
    return connector

# TITLE
t_title = slide.shapes.add_textbox(Inches(0.5), Inches(0.05), Inches(10.69), Inches(0.5))
t_title.text_frame.text = "KHUNG NGHIÊN CỨU TỔNG THỂ\n(Tích hợp Mô hình định lượng & Chuỗi tác động vĩ mô)"
t_title.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
t_title.text_frame.paragraphs[0].runs[0].font.name = 'Times New Roman'
t_title.text_frame.paragraphs[0].runs[0].font.size = Pt(15)
t_title.text_frame.paragraphs[0].runs[0].font.bold = True
t_title.text_frame.paragraphs[1].alignment = PP_ALIGN.CENTER
t_title.text_frame.paragraphs[1].runs[0].font.name = 'Times New Roman'
t_title.text_frame.paragraphs[1].runs[0].font.size = Pt(12)
t_title.text_frame.paragraphs[1].runs[0].font.italic = True

# THEORY BOX
t_theory = "LÝ THUYẾT NỀN TẢNG\nLý thuyết quản lý công   •   Lý thuyết thể chế   •   Lý thuyết quyền SHTT   •   Lý thuyết KT số"
add_box(slide, t_theory, Inches(1.5), Inches(0.6), Inches(8.69), Inches(0.6), font_size=11, fill_color=(245, 245, 245))

# 6 INDEPENDENT VARIABLES
lbl_iv = slide.shapes.add_textbox(Inches(0.3), Inches(1.0), Inches(2.4), Inches(0.4))
lbl_iv.text_frame.text = "CÁC BIẾN ĐỘC LẬP"
lbl_iv.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
lbl_iv.text_frame.paragraphs[0].runs[0].font.name = 'Times New Roman'
lbl_iv.text_frame.paragraphs[0].runs[0].font.bold = True

ivs = [
    "Khung pháp lý",
    "Năng lực cán bộ",
    "Nguồn lực tài chính",
    "Hạ tầng số",
    "Phối hợp thể chế",
    "Nhận thức xã hội"
]

iv_y_centers = []
start_y = 1.4
for i, iv in enumerate(ivs):
    y = start_y + i * 0.75
    shape = add_box(slide, iv, Inches(0.3), Inches(y), Inches(2.4), Inches(0.6), font_size=11, bold=False)
    iv_y_centers.append(y + 0.3)
    
# MV, DV, IMPACT (Center Y = 3.65)
t_mv = "BIẾN TRUNG GIAN\n\nMức độ chuyển đổi số trong quản lý"
add_box(slide, t_mv, Inches(3.3), Inches(3.05), Inches(2.3), Inches(1.2), font_size=11)

t_dv = "BIẾN PHỤ THUỘC\n\nHiệu quả QLNN về quyền SHTT"
add_box(slide, t_dv, Inches(6.2), Inches(3.05), Inches(2.3), Inches(1.2), font_size=11)

t_imp = "TÁC ĐỘNG VĨ MÔ\n\n• Thúc đẩy Đổi mới sáng tạo\n• Phát triển kinh tế số"
shape_imp = add_box(slide, t_imp, Inches(9.1), Inches(3.05), Inches(2.3), Inches(1.2), font_size=11)
for i, p in enumerate(shape_imp.text_frame.paragraphs):
    if p.text.startswith("•"):
        p.alignment = PP_ALIGN.LEFT
        p.level = 0

# ARROWS IV -> MV
for y in iv_y_centers:
    add_line_arrow(slide, Inches(2.7), Inches(y), Inches(3.3), Inches(3.65), thickness=1.5)

# ARROWS MV -> DV -> IMPACT
add_line_arrow(slide, Inches(5.6), Inches(3.65), Inches(6.2), Inches(3.65), thickness=2.0)
add_line_arrow(slide, Inches(8.5), Inches(3.65), Inches(9.1), Inches(3.65), thickness=2.0)


# MODERATING VARIABLE (Top Center X = 5.9)
t_mod = "BIẾN ĐIỀU TIẾT\n\nMức độ phát triển kinh tế số"
add_box(slide, t_mod, Inches(4.75), Inches(1.4), Inches(2.3), Inches(1.0), font_size=11)
# Arrow Mod -> MV-DV path
add_line_arrow(slide, Inches(5.9), Inches(2.4), Inches(5.9), Inches(3.65), dashed=True, thickness=1.5)


# CONTROL VARIABLE (Bottom Center X = 7.35)
t_ctrl = "BIẾN KIỂM SOÁT\n\n• Loại địa phương\n• Quy mô kinh tế\n• Ngành nghề"
shape_ctrl = add_box(slide, t_ctrl, Inches(6.2), Inches(4.8), Inches(2.3), Inches(1.2), font_size=11)
for i, p in enumerate(shape_ctrl.text_frame.paragraphs):
    if p.text.startswith("•"):
        p.alignment = PP_ALIGN.LEFT
        p.level = 0
# Arrow Control -> DV
add_line_arrow(slide, Inches(7.35), Inches(4.8), Inches(7.35), Inches(4.25), dashed=True, thickness=1.5)


# FEEDBACK LOOP
t_fb = "PHẢN HỒI CHÍNH SÁCH\nĐiều chỉnh chính sách   •   Hoàn thiện pháp luật   •   Nâng cấp nguồn lực   •   Đổi mới mô hình quản lý"
add_box(slide, t_fb, Inches(1.0), Inches(6.4), Inches(9.5), Inches(0.6), font_size=11)

# Drop arrow from Impact (X = 10.25)
add_line_arrow(slide, Inches(10.25), Inches(4.25), Inches(10.25), Inches(6.4), dashed=True, thickness=1.5)

# Up arrow from Feedback to IVs (X = 1.5)
add_line_arrow(slide, Inches(1.5), Inches(6.4), Inches(1.5), Inches(5.9), dashed=True, thickness=1.5)


# ARROWS FROM THEORY
add_line_arrow(slide, Inches(2.5), Inches(1.2), Inches(2.5), Inches(1.4), dashed=True)
add_line_arrow(slide, Inches(5.9), Inches(1.2), Inches(5.9), Inches(1.4), dashed=True)
add_line_arrow(slide, Inches(8.5), Inches(1.2), Inches(8.5), Inches(3.05), dashed=True)


prs.save("MoHinhNghienCuu_DinhLuong_Va_TacDong.pptx")
print("Saved quantitative SEM model with Impact and Feedback.")
