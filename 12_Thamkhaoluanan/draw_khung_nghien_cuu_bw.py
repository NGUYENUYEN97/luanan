from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

prs = Presentation()
slide = prs.slides.add_slide(prs.slide_layouts[6])

def add_box(slide, text, left, top, width, height, font_size=11, bold=False):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(255, 255, 255) # White background
    shape.line.color.rgb = RGBColor(0, 0, 0) # Black border
    shape.line.width = Pt(1.5)
    
    text_frame = shape.text_frame
    text_frame.text = text
    text_frame.word_wrap = True
    
    for paragraph in text_frame.paragraphs:
        paragraph.alignment = PP_ALIGN.CENTER
        for run in paragraph.runs:
            run.font.name = 'Times New Roman'
            run.font.size = Pt(font_size)
            run.font.color.rgb = RGBColor(0, 0, 0) # Black text
            run.font.bold = bold
    return shape

def add_arrow(slide, left, top, width, height, direction='DOWN'):
    shape_type = MSO_SHAPE.DOWN_ARROW
    if direction == 'RIGHT': shape_type = MSO_SHAPE.RIGHT_ARROW
    elif direction == 'UP': shape_type = MSO_SHAPE.UP_ARROW
    
    shape = slide.shapes.add_shape(shape_type, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(255, 255, 255) # White fill
    shape.line.color.rgb = RGBColor(0, 0, 0) # Black border
    return shape

# Title
title_shape = slide.shapes.add_textbox(Inches(0.5), Inches(0.1), Inches(9), Inches(0.5))
tf = title_shape.text_frame
tf.text = "KHUNG NGHIÊN CỨU LOGIC CỦA LUẬN ÁN (Tiếp cận theo ngành Quản lý Kinh tế)"
tf.paragraphs[0].alignment = PP_ALIGN.CENTER
tf.paragraphs[0].runs[0].font.name = 'Times New Roman'
tf.paragraphs[0].runs[0].font.size = Pt(16)
tf.paragraphs[0].runs[0].font.bold = True

# Level 1: Tổng quan và Phương pháp
b1 = add_box(slide, "TỔNG QUAN NGHIÊN CỨU VÀ PHƯƠNG PHÁP NGHIÊN CỨU (CHƯƠNG 1)\n- Xác định khoảng trống nghiên cứu\n- Đối tượng, phạm vi nghiên cứu\n- Phương pháp thu thập và phân tích dữ liệu", 
             Inches(1.5), Inches(0.6), Inches(7.0), Inches(0.8), bold=True)

add_arrow(slide, Inches(4.8), Inches(1.4), Inches(0.4), Inches(0.3), 'DOWN')

# Level 2: Cơ sở lý luận (Chương 2)
# Create a big container box for Theory
container2 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(1.7), Inches(9.0), Inches(2.3))
container2.fill.solid()
container2.fill.fore_color.rgb = RGBColor(255, 255, 255) # White fill
container2.line.color.rgb = RGBColor(0, 0, 0) # Black border
container2.line.width = Pt(1.5)
tf_c2 = container2.text_frame
tf_c2.text = "CƠ SỞ LÝ LUẬN VÀ KINH NGHIỆM THỰC TIỄN (CHƯƠNG 2)"
for p in tf_c2.paragraphs:
    for run in p.runs:
        run.font.name = 'Times New Roman'
        run.font.size = Pt(12)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0, 0, 0)

# Theory components
add_box(slide, "Chủ thể quản lý\n(Chính phủ, Các Bộ, Ngành cấp TW)", Inches(0.7), Inches(2.2), Inches(2.0), Inches(0.6), font_size=10)
add_box(slide, "Khách thể quản lý\n(Quyền tác giả, quyền liên quan, quyền sở hữu công nghiệp)", Inches(0.7), Inches(3.0), Inches(2.0), Inches(0.8), font_size=10)

add_box(slide, "NỘI DUNG QUẢN LÝ NHÀ NƯỚC\n1. Xây dựng chiến lược, chính sách, pháp luật\n2. Tổ chức thực hiện\n3. Hỗ trợ xác lập, khai thác, thương mại hóa\n4. Thanh tra, kiểm tra, thực thi quyền", Inches(2.9), Inches(2.2), Inches(3.2), Inches(1.6), bold=True, font_size=10)

add_box(slide, "TIÊU CHÍ ĐÁNH GIÁ\n1. Hiệu lực\n2. Hiệu quả\n3. Tính phù hợp\n4. Tính đồng bộ\n5. Tính minh bạch", Inches(6.3), Inches(2.2), Inches(1.4), Inches(1.6), font_size=10)

add_box(slide, "YẾU TỐ ẢNH HƯỞNG\n- Yếu tố bên trong\n- Yếu tố bên ngoài", Inches(7.9), Inches(2.2), Inches(1.4), Inches(1.6), font_size=10)

add_arrow(slide, Inches(4.8), Inches(4.0), Inches(0.4), Inches(0.3), 'DOWN')

# Level 3: Thực trạng (Chương 3)
container3 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(4.3), Inches(9.0), Inches(1.8))
container3.fill.solid()
container3.fill.fore_color.rgb = RGBColor(255, 255, 255) # White fill
container3.line.color.rgb = RGBColor(0, 0, 0) # Black border
container3.line.width = Pt(1.5)
tf_c3 = container3.text_frame
tf_c3.text = "THỰC TRẠNG QUẢN LÝ NHÀ NƯỚC VỀ SỞ HỮU TRÍ TUỆ TRONG KINH TẾ SỐ (CHƯƠNG 3)"
for p in tf_c3.paragraphs:
    for run in p.runs:
        run.font.name = 'Times New Roman'
        run.font.size = Pt(12)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0, 0, 0)

add_box(slide, "ĐÁNH GIÁ THỰC TRẠNG THEO 04 NỘI DUNG QUẢN LÝ NHÀ NƯỚC", Inches(0.7), Inches(4.7), Inches(3.0), Inches(0.5), font_size=10, bold=True)
add_arrow(slide, Inches(3.8), Inches(4.8), Inches(0.3), Inches(0.3), 'RIGHT')
add_box(slide, "ĐO LƯỜNG THEO 05 TIÊU CHÍ\n(Hiệu lực, Hiệu quả, Phù hợp,\nĐồng bộ, Minh bạch)", Inches(4.2), Inches(4.7), Inches(2.3), Inches(0.5), font_size=10, bold=True)
add_arrow(slide, Inches(6.6), Inches(4.8), Inches(0.3), Inches(0.3), 'RIGHT')
add_box(slide, "RÚT RA KẾT LUẬN\n- Kết quả đạt được\n- Hạn chế, yếu kém\n- Nguyên nhân (từ các yếu tố ảnh hưởng)", Inches(7.0), Inches(4.7), Inches(2.3), Inches(1.2), font_size=10, bold=True)

add_arrow(slide, Inches(4.8), Inches(6.1), Inches(0.4), Inches(0.3), 'DOWN')

# Level 4: Giải pháp (Chương 4)
b4 = add_box(slide, "QUAN ĐIỂM, ĐỊNH HƯỚNG VÀ GIẢI PHÁP HOÀN THIỆN (CHƯƠNG 4)\nĐề xuất các nhóm giải pháp tương ứng với 04 nội dung quản lý và khắc phục các nguyên nhân hạn chế", 
             Inches(1.5), Inches(6.4), Inches(7.0), Inches(0.8), bold=True)

prs.save("KhungNghienCuu_QLKT.pptx")
print("Saved BW pptx.")
