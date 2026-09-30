from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

prs = Presentation()
slide = prs.slides.add_slide(prs.slide_layouts[6])

def add_box(slide, text, left, top, width, height, font_size=11, bold=False):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(255, 255, 255)
    shape.line.color.rgb = RGBColor(0, 0, 0)
    shape.line.width = Pt(1.5)
    
    text_frame = shape.text_frame
    text_frame.text = text
    text_frame.word_wrap = True
    
    for paragraph in text_frame.paragraphs:
        paragraph.alignment = PP_ALIGN.CENTER
        for run in paragraph.runs:
            run.font.name = 'Times New Roman'
            run.font.size = Pt(font_size)
            run.font.color.rgb = RGBColor(0, 0, 0)
            run.font.bold = bold
    return shape

def add_arrow(slide, left, top, width, height, direction='DOWN'):
    shape_type = MSO_SHAPE.DOWN_ARROW
    if direction == 'RIGHT': shape_type = MSO_SHAPE.RIGHT_ARROW
    elif direction == 'UP': shape_type = MSO_SHAPE.UP_ARROW
    
    shape = slide.shapes.add_shape(shape_type, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(255, 255, 255)
    shape.line.color.rgb = RGBColor(0, 0, 0)
    return shape

# Title
title_shape = slide.shapes.add_textbox(Inches(0.5), Inches(0.2), Inches(9), Inches(0.5))
tf = title_shape.text_frame
tf.text = "KHUNG LÝ THUYẾT NGHIÊN CỨU"
tf.paragraphs[0].alignment = PP_ALIGN.CENTER
tf.paragraphs[0].runs[0].font.name = 'Times New Roman'
tf.paragraphs[0].runs[0].font.size = Pt(16)
tf.paragraphs[0].runs[0].font.bold = True

# --- TOP SECTION: THEORIES ---
# Container for Theories
theory_container = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(1.0), Inches(9.0), Inches(1.8))
theory_container.fill.solid()
theory_container.fill.fore_color.rgb = RGBColor(255, 255, 255)
theory_container.line.color.rgb = RGBColor(0, 0, 0)
theory_container.line.width = Pt(1.5)

# Separate Textbox for Theory Title
t_th = slide.shapes.add_textbox(Inches(0.5), Inches(1.1), Inches(9.0), Inches(0.4))
tf_th = t_th.text_frame
tf_th.text = "CƠ SỞ LÝ THUYẾT NỀN"
tf_th.paragraphs[0].alignment = PP_ALIGN.CENTER
tf_th.paragraphs[0].runs[0].font.name = 'Times New Roman'
tf_th.paragraphs[0].runs[0].font.size = Pt(12)
tf_th.paragraphs[0].runs[0].font.bold = True

add_box(slide, "Lý thuyết quản lý nhà nước\n&\nLý thuyết quản trị công", Inches(0.8), Inches(1.6), Inches(2.5), Inches(0.8), font_size=11, bold=True)
add_box(slide, "Lý thuyết thể chế", Inches(3.5), Inches(1.6), Inches(1.5), Inches(0.8), font_size=11, bold=True)
add_box(slide, "Lý thuyết kinh tế số\n&\nLý thuyết đổi mới sáng tạo", Inches(5.2), Inches(1.6), Inches(2.2), Inches(0.8), font_size=11, bold=True)
add_box(slide, "Các nghiên cứu\ntrước đây", Inches(7.6), Inches(1.6), Inches(1.6), Inches(0.8), font_size=11, bold=True)

add_arrow(slide, Inches(4.8), Inches(2.8), Inches(0.4), Inches(0.4), 'DOWN')


# --- BOTTOM SECTION: VARIABLES ---
# Dashed box for Conceptual Model
model_container = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.2), Inches(3.3), Inches(9.6), Inches(3.8))
model_container.fill.background()
model_container.line.color.rgb = RGBColor(0, 0, 0)
model_container.line.width = Pt(1.5)
model_container.line.dash_style = 4
slide.shapes._spTree.remove(model_container._element)
slide.shapes._spTree.insert(2, model_container._element)

t_md = slide.shapes.add_textbox(Inches(0.3), Inches(3.4), Inches(4.0), Inches(0.3))
t_md.text_frame.text = "MÔ HÌNH BIẾN NGHIÊN CỨU"
t_md.text_frame.paragraphs[0].runs[0].font.name = 'Times New Roman'
t_md.text_frame.paragraphs[0].runs[0].font.size = Pt(11)
t_md.text_frame.paragraphs[0].runs[0].font.bold = True
t_md.text_frame.paragraphs[0].runs[0].font.italic = True

# Moderating Variable
text_yt = "CÁC YẾU TỐ ẢNH HƯỞNG\n(Biến điều tiết)\n\n- Nhóm yếu tố bên trong\n- Nhóm yếu tố bên ngoài"
add_box(slide, text_yt, Inches(3.5), Inches(3.6), Inches(3.0), Inches(1.2), font_size=11, bold=True)

# Main Variables
text_nd = "NỘI DUNG QUẢN LÝ NHÀ NƯỚC\n(Biến độc lập)\n\n1. Xây dựng chiến lược, chính sách, pháp luật\n2. Tổ chức thực hiện chiến lược, chính sách, pháp luật\n3. Hỗ trợ xác lập, khai thác, chuyển giao, thương mại hóa\n4. Thanh tra, kiểm tra, giám sát và thực thi quyền"
add_box(slide, text_nd, Inches(0.5), Inches(4.8), Inches(3.8), Inches(2.0), font_size=11, bold=True)

# Arrow from Independent to Dependent
add_arrow(slide, Inches(4.4), Inches(5.6), Inches(1.2), Inches(0.4), 'RIGHT')

# Arrow from Moderating to the Main Arrow
add_arrow(slide, Inches(4.8), Inches(4.8), Inches(0.4), Inches(0.7), 'DOWN')


text_cl = "CHẤT LƯỢNG QUẢN LÝ NHÀ NƯỚC\n(Biến phụ thuộc)\n\n1. Hiệu lực\n2. Hiệu quả\n3. Tính phù hợp\n4. Tính đồng bộ và thống nhất\n5. Tính minh bạch"
add_box(slide, text_cl, Inches(5.7), Inches(4.8), Inches(3.8), Inches(2.0), font_size=11, bold=True)

prs.save("KhungLyThuyetNghienCuu_ThuanTuy.pptx")
print("Saved pure theoretical framework.")
