from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_AUTO_SIZE

prs = Presentation()
prs.slide_width = Inches(11.69) # A4 Landscape
prs.slide_height = Inches(8.27)
slide = prs.slides.add_slide(prs.slide_layouts[6])

def add_box(slide, text, left, top, width, height, font_size=10):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    # ALL BLACK AND WHITE
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(255, 255, 255)
    shape.line.color.rgb = RGBColor(0, 0, 0)
    shape.line.width = Pt(1.5)
    
    tf = shape.text_frame
    tf.text = text
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    
    for i, p in enumerate(tf.paragraphs):
        # We manually align and do not use p.level = 1 because it adds uncontrollable large margins
        if "•" in p.text or p.text.startswith("1.") or p.text.startswith("2.") or p.text.startswith("3.") or p.text.startswith("4."):
            p.alignment = PP_ALIGN.LEFT
        else:
            p.alignment = PP_ALIGN.CENTER
            
        for run in p.runs:
            run.font.name = 'Times New Roman'
            run.font.size = Pt(font_size)
            run.font.color.rgb = RGBColor(0, 0, 0)
            # Bold for headings
            if i == 0 or (i == 1 and p.text.isupper()):
                run.font.bold = True
    return shape

def add_block_arrow(slide, left, top, width, height, direction='DOWN'):
    shape_type = MSO_SHAPE.DOWN_ARROW
    if direction == 'RIGHT': shape_type = MSO_SHAPE.RIGHT_ARROW
    elif direction == 'LEFT': shape_type = MSO_SHAPE.LEFT_ARROW
    elif direction == 'UP': shape_type = MSO_SHAPE.UP_ARROW
    
    shape = slide.shapes.add_shape(shape_type, left, top, width, height)
    # ALL BLACK AND WHITE
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(255, 255, 255)
    shape.line.color.rgb = RGBColor(0, 0, 0)
    shape.line.width = Pt(1.0)
    return shape

def add_dashed_line(slide, x, y, cx, cy):
    line = slide.shapes.add_shape(MSO_SHAPE.LINE_INVERSE, x, y, cx, cy)
    line.line.color.rgb = RGBColor(0,0,0)
    line.line.width = Pt(1.5)
    line.line.dash_style = 4
    return line

def add_arrowhead(slide, tip_x, tip_y, direction='LEFT'):
    size = Inches(0.12)
    
    # Calculate top/left based on direction and tip
    if direction == 'LEFT':
        left = tip_x
        top = tip_y - size/2
        rot = 270
    elif direction == 'RIGHT':
        left = tip_x - size
        top = tip_y - size/2
        rot = 90
    elif direction == 'DOWN':
        left = tip_x - size/2
        top = tip_y - size
        rot = 180
    else: # UP
        left = tip_x - size/2
        top = tip_y
        rot = 0
        
    tri = slide.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE, left, top, size, size)
    tri.fill.solid()
    tri.fill.fore_color.rgb = RGBColor(0,0,0) # Solid black head
    tri.line.color.rgb = RGBColor(0,0,0)
    tri.rotation = rot

# Title
t_title = slide.shapes.add_textbox(Inches(0.5), Inches(0.05), Inches(10.69), Inches(0.4))
t_title.text_frame.text = "KHUNG LÝ THUYẾT NGHIÊN CỨU"
t_title.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
t_title.text_frame.paragraphs[0].runs[0].font.name = 'Times New Roman'
t_title.text_frame.paragraphs[0].runs[0].font.size = Pt(16)
t_title.text_frame.paragraphs[0].runs[0].font.bold = True

# --- LÝ THUYẾT NỀN TẢNG ---
t1 = "LÝ THUYẾT NỀN TẢNG\nLý thuyết quản lý công  •  Lý thuyết quản lý công mới  •  Lý thuyết quản trị số  •  Lý thuyết quản trị thích ứng  •  Lý thuyết thể chế  •  Lý thuyết về quyền sở hữu trí tuệ"
add_box(slide, t1, Inches(1.5), Inches(0.4), Inches(8.69), Inches(0.6), font_size=10)

add_block_arrow(slide, Inches(5.75), Inches(1.0), Inches(0.2), Inches(0.2))

# --- KHÁI NIỆM CỐT LÕI ---
t2 = "KHÁI NIỆM CỐT LÕI\nQuyền sở hữu trí tuệ  •  Quản lý nhà nước  •  Kinh tế số  •  Quản trị số  •  Quản lý nhà nước về quyền SHTT"
add_box(slide, t2, Inches(1.5), Inches(1.2), Inches(8.69), Inches(0.5), font_size=10)

add_block_arrow(slide, Inches(5.75), Inches(1.7), Inches(0.2), Inches(0.25))

# --- BIẾN ĐỘC LẬP ---
t3 = "BIẾN ĐỘC LẬP\nCÁC YẾU TỐ ĐẦU VÀO\n\n• Nhóm thể chế: Pháp luật, Chính sách, Cơ chế phối hợp\n• Nhóm nguồn lực: Nguồn nhân lực, Tài chính, Hạ tầng số\n• Nhóm công nghệ: Dữ liệu số, AI, Nền tảng số"
add_box(slide, t3, Inches(0.5), Inches(1.95), Inches(3.2), Inches(1.8), font_size=10)

# Mũi tên từ Biến độc lập sang Biến trung tâm
add_block_arrow(slide, Inches(3.7), Inches(2.75), Inches(0.5), Inches(0.2), 'RIGHT')

# --- BIẾN TRUNG TÂM ---
t4 = "BIẾN TRUNG TÂM NGHIÊN CỨU\nQUẢN LÝ NHÀ NƯỚC VỀ QUYỀN SHTT TRONG NỀN KINH TẾ SỐ\n\n1. Xây dựng chiến lược, chính sách, pháp luật\n2. Tổ chức thực hiện\n3. Hỗ trợ xác lập, khai thác, chuyển giao, thương mại hóa tài sản trí tuệ\n4. Thanh tra, kiểm tra, giám sát, thực thi quyền sở hữu trí tuệ"
add_box(slide, t4, Inches(4.2), Inches(1.95), Inches(3.7), Inches(1.6), font_size=10)

add_block_arrow(slide, Inches(5.95), Inches(3.55), Inches(0.2), Inches(0.25))

# --- BIẾN TRUNG GIAN ---
t5 = "BIẾN TRUNG GIAN\nNĂNG LỰC QLNN VỀ QUYỀN SHTT\n\n• Năng lực hoạch định chính sách\n• Năng lực tổ chức thực hiện\n• Năng lực phối hợp liên ngành\n• Năng lực chuyển đổi số\n• Năng lực thanh tra và thực thi"
add_box(slide, t5, Inches(4.4), Inches(3.8), Inches(3.3), Inches(1.4), font_size=10)

add_block_arrow(slide, Inches(5.95), Inches(5.2), Inches(0.2), Inches(0.25))

# --- BIẾN PHỤ THUỘC ---
t6 = "BIẾN PHỤ THUỘC\nCHẤT LƯỢNG QLNN VỀ QUYỀN SỞ HỮU TRÍ TUỆ\n\n• Hiệu lực   • Hiệu quả   • Tính phù hợp   • Tính đồng bộ   • Tính minh bạch"
add_box(slide, t6, Inches(4.4), Inches(5.45), Inches(3.3), Inches(0.8), font_size=10)

add_block_arrow(slide, Inches(5.95), Inches(6.25), Inches(0.2), Inches(0.25))

# --- KẾT QUẢ/TÁC ĐỘNG ---
t7 = "KẾT QUẢ VÀ TÁC ĐỘNG\n\n• Bảo vệ hiệu quả quyền SHTT       • Thúc đẩy đổi mới sáng tạo\n• Phát triển tài sản trí tuệ         • Nâng cao năng lực cạnh tranh quốc gia\n• Thúc đẩy phát triển kinh tế số"
add_box(slide, t7, Inches(4.2), Inches(6.5), Inches(3.7), Inches(0.9), font_size=10)


# --- YẾU TỐ ĐIỀU TIẾT ---
tdt = "YẾU TỐ ĐIỀU TIẾT\n(Các yếu tố ảnh hưởng)\n\n• Mức độ phát triển kinh tế số\n• Năng lực đổi mới sáng tạo\n• Hội nhập quốc tế\n• Nhận thức xã hội\n• Mức độ cạnh tranh thị trường"
add_box(slide, tdt, Inches(8.3), Inches(2.9), Inches(2.9), Inches(2.0), font_size=10)

# Mũi tên điều tiết nét đứt chỉ vào mũi tên Trung tâm -> Trung gian (Y=3.675)
add_dashed_line(slide, Inches(8.3), Inches(3.675), Inches(-2.05), Inches(0))
add_arrowhead(slide, Inches(6.25), Inches(3.675), 'LEFT')

# Mũi tên điều tiết nét đứt chỉ vào mũi tên Trung gian -> Phụ thuộc (Y=5.325)
add_dashed_line(slide, Inches(8.3), Inches(5.325), Inches(-2.05), Inches(0))
add_arrowhead(slide, Inches(6.25), Inches(5.325), 'LEFT')

# Vertical line connecting the two dashed arrows
add_dashed_line(slide, Inches(8.3), Inches(3.675), Inches(0), Inches(1.65))


# --- FEEDBACK LOOP ---
tfb = "PHẢN HỒI CHÍNH SÁCH\nKết quả thực hiện  ➔  Đánh giá  ➔  Điều chỉnh chính sách, thể chế, công cụ  ➔  Hoàn thiện QLNN"
add_box(slide, tfb, Inches(1.5), Inches(7.6), Inches(8.69), Inches(0.5), font_size=10)

# Từ Kết quả xuống Feedback
add_dashed_line(slide, Inches(6.05), Inches(7.4), Inches(0), Inches(0.2))
add_arrowhead(slide, Inches(6.05), Inches(7.6), 'DOWN')

# Từ khối Feedback vòng lên Biến trung tâm
# Ra từ trái Feedback
add_dashed_line(slide, Inches(1.5), Inches(7.85), Inches(-1.3), Inches(0))
# Đi lên qua Biến độc lập
add_dashed_line(slide, Inches(0.2), Inches(7.85), Inches(0), Inches(-6.1))
# Đi ngang sang phải phía trên Biến độc lập
add_dashed_line(slide, Inches(0.2), Inches(1.75), Inches(4.7), Inches(0))
# Đi xuống cắm vào Biến trung tâm
add_dashed_line(slide, Inches(4.9), Inches(1.75), Inches(0), Inches(0.2))
add_arrowhead(slide, Inches(4.9), Inches(1.95), 'DOWN')

prs.save("KhungLyThuyet_GS_BW_Final.pptx")
print("Saved BW, no overflow, properly formatted PPTX.")
