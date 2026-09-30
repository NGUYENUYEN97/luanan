from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_AUTO_SIZE
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls

prs = Presentation()
prs.slide_width = Inches(11.69)
prs.slide_height = Inches(8.27)
slide = prs.slides.add_slide(prs.slide_layouts[6])

def add_box(slide, text, left, top, width, height, fill_color=(255, 255, 255), border_color=(0,0,0), border_width=1.5):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(*fill_color)
    shape.line.color.rgb = RGBColor(*border_color)
    shape.line.width = Pt(border_width)
    
    tf = shape.text_frame
    tf.text = text
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    
    for i, p in enumerate(tf.paragraphs):
        if i == 0:
            p.alignment = PP_ALIGN.CENTER
        elif p.text.startswith("Hiệu quả QLNN") or p.text.startswith("Điều chỉnh"):
            p.alignment = PP_ALIGN.CENTER
        elif p.text.startswith("•"):
            p.alignment = PP_ALIGN.LEFT
            p.level = 0
        else:
            p.alignment = PP_ALIGN.CENTER
            
        for run in p.runs:
            run.font.name = 'Times New Roman'
            run.font.size = Pt(10) if i == 0 else Pt(9)
            run.font.color.rgb = RGBColor(0, 0, 0)
            
            if i == 0:
                run.font.bold = True
            elif p.text.startswith("Hiệu quả QLNN"):
                run.font.italic = True
    return shape

def add_line_arrow(slide, begin_x, begin_y, end_x, end_y, dashed=False, arrow_end=True, arrow_start=False, thickness=1.5):
    connector = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, begin_x, begin_y, end_x, end_y)
    connector.line.color.rgb = RGBColor(0,0,0)
    connector.line.width = Pt(thickness)
    if dashed:
        connector.line.dash_style = 4
        
    ln = connector.element.spPr.ln
    if ln is not None:
        if arrow_end:
            tailEnd = parse_xml(r'<a:tailEnd type="triangle" w="med" len="med" %s/>' % nsdecls('a'))
            ln.append(tailEnd)
        if arrow_start:
            headEnd = parse_xml(r'<a:headEnd type="triangle" w="med" len="med" %s/>' % nsdecls('a'))
            ln.append(headEnd)
    return connector


# --- TITLE ---
t_title = slide.shapes.add_textbox(Inches(0.5), Inches(0.1), Inches(10.69), Inches(0.6))
t_title.text_frame.text = "KHUNG LÝ THUYẾT: QUẢN LÝ NHÀ NƯỚC VỀ QUYỀN SHTT TRONG NỀN KINH TẾ SỐ\n(Tiếp cận theo hệ thống IPOOIF)"
t_title.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
t_title.text_frame.paragraphs[0].runs[0].font.name = 'Times New Roman'
t_title.text_frame.paragraphs[0].runs[0].font.size = Pt(14)
t_title.text_frame.paragraphs[0].runs[0].font.bold = True
t_title.text_frame.paragraphs[1].alignment = PP_ALIGN.CENTER
t_title.text_frame.paragraphs[1].runs[0].font.name = 'Times New Roman'
t_title.text_frame.paragraphs[1].runs[0].font.size = Pt(12)
t_title.text_frame.paragraphs[1].runs[0].font.italic = True


# --- CONTEXT BOX ---
ctx_shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.15), Inches(0.9), Inches(11.39), Inches(7.1))
ctx_shape.fill.solid()
ctx_shape.fill.fore_color.rgb = RGBColor(255, 255, 255)
ctx_shape.line.color.rgb = RGBColor(0, 0, 0)
ctx_shape.line.width = Pt(2.0)
ctx_shape.line.dash_style = 4

lbl = slide.shapes.add_textbox(Inches(0.3), Inches(1.0), Inches(5.0), Inches(0.4))
lbl.text_frame.text = "BỐI CẢNH: NỀN KINH TẾ SỐ (Môi trường nghiên cứu)"
lbl.text_frame.paragraphs[0].runs[0].font.name = 'Times New Roman'
lbl.text_frame.paragraphs[0].runs[0].font.size = Pt(11)
lbl.text_frame.paragraphs[0].runs[0].font.bold = True
lbl.text_frame.paragraphs[0].runs[0].font.italic = True


# --- X Coordinates and Widths ---
# Box 1: 0.4 (1.8) -> 2.2
# Gap 1: 0.2
# Box 2: 2.4 (2.4) -> 4.8
# Gap 2: 0.2
# Box 3: 5.0 (2.1) -> 7.1
# Gap 3: 0.2
# Box 4: 7.3 (2.1) -> 9.4
# Gap 4: 0.2
# Box 5: 9.6 (1.8) -> 11.4


# --- 1. INPUT ---
t1 = "1. INPUT (Đầu vào)\n\n• Khung pháp lý\n• Nguồn nhân lực\n• Nguồn lực tài chính\n• Hạ tầng số\n• Dữ liệu số\n• Phối hợp thể chế"
add_box(slide, t1, Inches(0.4), Inches(2.3), Inches(1.8), Inches(2.4), fill_color=(245, 245, 245))

# --- 2. PROCESS ---
t2 = "2. PROCESS (Quá trình)\n\n• Hoạch định chính sách\n• Ban hành & thực thi pháp luật\n• Đăng ký & xác lập quyền\n• Thanh tra, kiểm tra\n• Xử lý vi phạm\n• Giải quyết tranh chấp\n• Quản lý dữ liệu SHTT\n• Giám sát môi trường số\n• Hợp tác quốc tế\n• Quản trị nền tảng số"
add_box(slide, t2, Inches(2.4), Inches(1.6), Inches(2.4), Inches(3.8), fill_color=(245, 245, 245))

# --- 3. OUTPUT ---
t3 = "3. OUTPUT (Đầu ra)\n\n• Cơ sở dữ liệu SHTT số\n• Dịch vụ công trực tuyến\n• Hồ sơ điện tử\n• Kết nối & chia sẻ dữ liệu\n• Nâng cao phát hiện vi phạm"
add_box(slide, t3, Inches(5.0), Inches(2.3), Inches(2.1), Inches(2.4), fill_color=(245, 245, 245))

# --- 4. OUTCOME ---
t4 = "4. OUTCOME (Kết quả)\n\nHiệu quả QLNN về quyền SHTT:\n• Hiệu lực\n• Hiệu quả\n• Tính minh bạch\n• Khả năng thích ứng số\n• Mức độ hài lòng"
add_box(slide, t4, Inches(7.3), Inches(2.3), Inches(2.1), Inches(2.4), fill_color=(245, 245, 245))

# --- 5. IMPACT ---
t5 = "5. IMPACT (Tác động)\n\n• Hiệu quả bảo vệ quyền SHTT\n• Thúc đẩy Đổi mới sáng tạo\n• Phát triển kinh tế số"
add_box(slide, t5, Inches(9.6), Inches(2.4), Inches(1.8), Inches(2.2), fill_color=(245, 245, 245))


# --- HORIZONTAL ARROWS ---
add_line_arrow(slide, Inches(2.2), Inches(3.5), Inches(2.4), Inches(3.5), thickness=2.0)
add_line_arrow(slide, Inches(4.8), Inches(3.5), Inches(5.0), Inches(3.5), thickness=2.0)
add_line_arrow(slide, Inches(7.1), Inches(3.5), Inches(7.3), Inches(3.5), thickness=2.0)
add_line_arrow(slide, Inches(9.4), Inches(3.5), Inches(9.6), Inches(3.5), thickness=2.0)


# --- 6. FEEDBACK ---
tfb = "6. FEEDBACK (Phản hồi chính sách)\nĐiều chỉnh chính sách   •   Hoàn thiện pháp luật   •   Nâng cấp nguồn lực   •   Đổi mới mô hình quản lý"
# X = center of Box 1 (0.4 + 0.9 = 1.3) to center of Box 5 (9.6 + 0.9 = 10.5)
# Width = 10.5 - 1.3 = 9.2
add_box(slide, tfb, Inches(1.3), Inches(6.0), Inches(9.2), Inches(0.7), fill_color=(255, 255, 255), border_width=1.5)

# Drop arrow from Box 5 (Impact)
add_line_arrow(slide, Inches(10.5), Inches(4.6), Inches(10.5), Inches(6.0), dashed=True, thickness=1.5)

# Up arrow from Feedback to Box 1 (Input)
add_line_arrow(slide, Inches(1.3), Inches(6.0), Inches(1.3), Inches(4.7), dashed=True, thickness=1.5)


prs.save("KhungLyThuyet_IPOOIF_Chuan.pptx")
print("Saved PPTX for IPOOIF framework.")
