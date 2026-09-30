from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

prs = Presentation()
prs.slide_width = Inches(8.27)
prs.slide_height = Inches(11.69)
slide = prs.slides.add_slide(prs.slide_layouts[6])

def add_box(slide, text, left, top, width, height, font_size=10, bold=False):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(255, 255, 255)
    shape.line.color.rgb = RGBColor(0, 0, 0)
    shape.line.width = Pt(1.5)
    text_frame = shape.text_frame
    text_frame.text = text
    text_frame.word_wrap = True
    for paragraph in text_frame.paragraphs:
        paragraph.alignment = PP_ALIGN.CENTER
        for run in paragraph.runs:
            run.font.name = 'Times New Roman'
            run.font.size = Pt(font_size)
            run.font.color.rgb = RGBColor(0, 0, 0)
            run.font.bold = bold
    return shape

def add_arrow(slide, left, top, width, height, direction='DOWN'):
    shape_type = MSO_SHAPE.DOWN_ARROW
    if direction == 'RIGHT': shape_type = MSO_SHAPE.RIGHT_ARROW
    elif direction == 'UP': shape_type = MSO_SHAPE.UP_ARROW
    
    shape = slide.shapes.add_shape(shape_type, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(0, 0, 0)
    shape.line.color.rgb = RGBColor(0, 0, 0)
    return shape

def add_separator(slide, y_pos, text):
    # Line
    line = slide.shapes.add_shape(MSO_SHAPE.LINE_INVERSE, Inches(0.2), Inches(y_pos), Inches(7.87), Inches(0))
    line.line.color.rgb = RGBColor(0, 0, 0)
    line.line.width = Pt(2)
    # Text
    tb = slide.shapes.add_textbox(Inches(0.2), Inches(y_pos - 0.1), Inches(7.87), Inches(0.3))
    tb.text_frame.text = text
    tb.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    tb.text_frame.paragraphs[0].runs[0].font.name = 'Times New Roman'
    tb.text_frame.paragraphs[0].runs[0].font.size = Pt(12)
    tb.text_frame.paragraphs[0].runs[0].font.bold = True

# --- TẦNG 1 ---
add_separator(slide, 0.3, "TẦNG 1: THEORETICAL FRAMEWORK (Nghiên cứu dựa trên lý thuyết nào?)")

t1_theories = "Lý thuyết Quản lý công                 Lý thuyết Quản trị số                 Lý thuyết Đổi mới sáng tạo\nLý thuyết Quản trị công mới            Lý thuyết Thể chế                  Lý thuyết Quyền SHTT"
add_box(slide, t1_theories, Inches(0.5), Inches(0.6), Inches(7.27), Inches(0.8), font_size=11, bold=True)

add_arrow(slide, Inches(4.0), Inches(1.4), Inches(0.27), Inches(0.2))

t1_concepts = "Khái niệm cốt lõi:\nQuyền SHTT  -  Nhà nước  -  Quản lý nhà nước  -  Kinh tế số  -  Quản trị số"
add_box(slide, t1_concepts, Inches(1.0), Inches(1.6), Inches(6.27), Inches(0.5), font_size=11, bold=True)

# --- TẦNG 2 ---
add_separator(slide, 2.3, "TẦNG 2: ANALYTICAL / CONCEPTUAL FRAMEWORK (Phân tích cái gì?)")

add_box(slide, "QUẢN LÝ NHÀ NƯỚC VỀ QUYỀN SHTT TRONG NỀN KINH TẾ SỐ", Inches(1.5), Inches(2.6), Inches(5.27), Inches(0.5), font_size=12, bold=True)

# Branching lines from center down to 3 columns
line_vert1 = slide.shapes.add_shape(MSO_SHAPE.LINE_INVERSE, Inches(4.135), Inches(3.1), Inches(0), Inches(0.2))
line_vert1.line.color.rgb = RGBColor(0,0,0)
line_vert1.line.width = Pt(1.5)

line_horz = slide.shapes.add_shape(MSO_SHAPE.LINE_INVERSE, Inches(1.6), Inches(3.3), Inches(5.07), Inches(0))
line_horz.line.color.rgb = RGBColor(0,0,0)
line_horz.line.width = Pt(1.5)

for x in [1.6, 4.135, 6.67]:
    arr = slide.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(x - 0.1), Inches(3.3), Inches(0.2), Inches(0.2))
    arr.fill.solid()
    arr.fill.fore_color.rgb = RGBColor(0,0,0)
    arr.line.color.rgb = RGBColor(0,0,0)

# Columns
col1 = "NỘI DUNG QLNN\n\n• Chiến lược\n• Chính sách\n• Pháp luật\n• Tổ chức thực hiện\n• Hỗ trợ thương mại hóa\n• Thanh tra, thực thi"
add_box(slide, col1, Inches(0.5), Inches(3.5), Inches(2.2), Inches(2.0), font_size=11)

col2 = "CÔNG CỤ QLNN\n\n• Công cụ Pháp luật\n• Công cụ Hành chính\n• Công cụ Kinh tế\n• Công cụ Công nghệ"
add_box(slide, col2, Inches(3.035), Inches(3.5), Inches(2.2), Inches(2.0), font_size=11)

col3 = "YẾU TỐ ẢNH HƯỞNG\n\n• Thể chế\n• Công nghệ số\n• Trí tuệ nhân tạo (AI)\n• Dữ liệu\n• Nguồn nhân lực\n• Phối hợp liên ngành\n• Hợp tác, Nhận thức"
add_box(slide, col3, Inches(5.57), Inches(3.5), Inches(2.2), Inches(2.0), font_size=11)

# Converging lines down to Chat Luong
line_horz2 = slide.shapes.add_shape(MSO_SHAPE.LINE_INVERSE, Inches(1.6), Inches(5.7), Inches(5.07), Inches(0))
line_horz2.line.color.rgb = RGBColor(0,0,0)
line_horz2.line.width = Pt(1.5)
for x in [1.6, 4.135, 6.67]:
    l_v = slide.shapes.add_shape(MSO_SHAPE.LINE_INVERSE, Inches(x), Inches(5.5), Inches(0), Inches(0.2))
    l_v.line.color.rgb = RGBColor(0,0,0)
    l_v.line.width = Pt(1.5)

arr_down = slide.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(4.035), Inches(5.7), Inches(0.2), Inches(0.2))
arr_down.fill.solid()
arr_down.fill.fore_color.rgb = RGBColor(0,0,0)
arr_down.line.color.rgb = RGBColor(0,0,0)

add_box(slide, "CHẤT LƯỢNG QLNN\nHiệu lực - Hiệu quả - Phù hợp - Đồng bộ - Minh bạch", Inches(1.5), Inches(5.9), Inches(5.27), Inches(0.5), font_size=11, bold=True)

add_arrow(slide, Inches(4.0), Inches(6.4), Inches(0.27), Inches(0.2))

add_box(slide, "TÁC ĐỘNG\nBảo vệ quyền SHTT - Đổi mới sáng tạo - Phát triển kinh tế số", Inches(1.5), Inches(6.6), Inches(5.27), Inches(0.5), font_size=11, bold=True)

# --- TẦNG 3 ---
add_separator(slide, 7.3, "TẦNG 3: RESEARCH FRAMEWORK (Luận án triển khai nghiên cứu như thế nào?)")

# Vertical boxes
y_start = 7.7
h_box = 0.5
spacing = 0.25

steps = [
    "Tổng quan nghiên cứu",
    "Cơ sở lý luận và khung phân tích",
    "Phân tích thực trạng QLNN về quyền SHTT trong nền kinh tế số",
    "Đánh giá kết quả, hạn chế, nguyên nhân",
    "Quan điểm, định hướng và giải pháp"
]

for i, text in enumerate(steps):
    y = y_start + i * (h_box + spacing)
    add_box(slide, text, Inches(2.0), Inches(y), Inches(4.27), Inches(h_box), font_size=11, bold=True)
    if i < len(steps) - 1:
        add_arrow(slide, Inches(4.0), Inches(y + h_box), Inches(0.27), Inches(spacing))

prs.save("TongThe_3_Tang.pptx")
print("Saved 3-tier master map.")
