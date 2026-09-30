from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_AUTO_SIZE
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls

prs = Presentation()
prs.slide_width = Inches(11.69)
prs.slide_height = Inches(8.27)
slide = prs.slides.add_slide(prs.slide_layouts[6])

def add_box(slide, text, left, top, width, height, font_size=10, fill_color=(255, 255, 255), border_color=(0,0,0), border_width=1.5, bold_header=True, dash_style=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    if fill_color:
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(*fill_color)
    else:
        shape.fill.background()
        
    shape.line.color.rgb = RGBColor(*border_color)
    shape.line.width = Pt(border_width)
    if dash_style:
        shape.line.dash_style = dash_style
    
    tf = shape.text_frame
    tf.text = text
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    
    for i, p in enumerate(tf.paragraphs):
        if p.text.startswith("•"):
            p.alignment = PP_ALIGN.LEFT
            p.level = 0
        else:
            p.alignment = PP_ALIGN.CENTER
            
        for run in p.runs:
            run.font.name = 'Times New Roman'
            run.font.size = Pt(font_size)
            run.font.color.rgb = RGBColor(0, 0, 0)
            if bold_header and i == 0:
                run.font.bold = True
    return shape

def add_line_arrow(slide, begin_x, begin_y, end_x, end_y, dashed=False, arrow_end=True, arrow_start=False, thickness=1.5):
    connector = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, begin_x, begin_y, end_x, end_y)
    connector.line.color.rgb = RGBColor(0,0,0)
    connector.line.width = Pt(thickness)
    if dashed:
        connector.line.dash_style = 4
        
    ln = connector.element.spPr.ln
    if ln is not None:
        if arrow_end:
            tailEnd = parse_xml(r'<a:tailEnd type="triangle" w="med" len="med" %s/>' % nsdecls('a'))
            ln.append(tailEnd)
        if arrow_start:
            headEnd = parse_xml(r'<a:headEnd type="triangle" w="med" len="med" %s/>' % nsdecls('a'))
            ln.append(headEnd)
    return connector

# TITLE
t_title = slide.shapes.add_textbox(Inches(0.5), Inches(0.05), Inches(10.69), Inches(0.5))
t_title.text_frame.text = "KHUNG NGHIÊN CỨU TỔNG THỂ CÓ CẤU TRÚC ĐO LƯỜNG\n(Chi tiết hóa các biến quan sát)"
t_title.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
t_title.text_frame.paragraphs[0].runs[0].font.name = 'Times New Roman'
t_title.text_frame.paragraphs[0].runs[0].font.size = Pt(15)
t_title.text_frame.paragraphs[0].runs[0].font.bold = True
t_title.text_frame.paragraphs[1].alignment = PP_ALIGN.CENTER
t_title.text_frame.paragraphs[1].runs[0].font.name = 'Times New Roman'
t_title.text_frame.paragraphs[1].runs[0].font.size = Pt(12)
t_title.text_frame.paragraphs[1].runs[0].font.italic = True

# THEORY BOX
t_theory = "LÝ THUYẾT NỀN TẢNG\nLý thuyết quản lý công   •   Lý thuyết thể chế   •   Lý thuyết quyền SHTT   •   Lý thuyết KT số"
add_box(slide, t_theory, Inches(1.5), Inches(0.6), Inches(8.69), Inches(0.5), font_size=11, fill_color=(245, 245, 245))

# ---------------------------------------------------------
# IV CONTAINER
add_box(slide, "CÁC BIẾN ĐỘC LẬP", Inches(0.2), Inches(1.2), Inches(2.5), Inches(5.0), font_size=11, fill_color=None, dash_style=4)

ivs = [
    "Khung pháp lý",
    "Năng lực cán bộ",
    "Nguồn lực tài chính",
    "Hạ tầng số",
    "Phối hợp thể chế",
    "Nhận thức xã hội"
]

iv_y_centers = []
start_y = 1.55
for i, iv in enumerate(ivs):
    y = start_y + i * 0.75
    shape = add_box(slide, iv, Inches(0.35), Inches(y), Inches(2.2), Inches(0.55), font_size=11, bold_header=False, fill_color=(250, 250, 250))
    iv_y_centers.append(y + 0.275)
# ---------------------------------------------------------

# MV (X = 3.1)
t_mv = "BIẾN TRUNG GIAN\nMức độ CĐS trong quản lý\n\n• Số hóa & liên thông CSDL\n• Tự động hóa quy trình nội bộ\n• Cung cấp dịch vụ công trực tuyến"
add_box(slide, t_mv, Inches(3.1), Inches(2.7), Inches(2.4), Inches(1.8), font_size=10, fill_color=(245, 245, 245))

# DV (X = 6.0)
t_dv = "BIẾN PHỤ THUỘC\nHiệu quả QLNN về SHTT\n\n• Hiệu lực bảo vệ pháp lý\n• Tính minh bạch thông tin\n• Tốc độ & năng lực phản ứng số\n• Mức độ hài lòng của công chúng"
add_box(slide, t_dv, Inches(6.0), Inches(2.7), Inches(2.4), Inches(1.9), font_size=10, fill_color=(245, 245, 245))

# IMPACT (X = 8.9)
t_imp = "TÁC ĐỘNG VĨ MÔ\n\n• Thúc đẩy Đổi mới sáng tạo\n• Gia tăng tài sản trí tuệ quốc gia\n• Đóng góp Phát triển KT số"
add_box(slide, t_imp, Inches(8.9), Inches(2.8), Inches(2.4), Inches(1.6), font_size=10, fill_color=(245, 245, 245))


# ARROWS IV -> MV (Converging to Center Y = 3.6)
for y in iv_y_centers:
    add_line_arrow(slide, Inches(2.55), Inches(y), Inches(3.1), Inches(3.6), thickness=1.5)

# ARROWS MV -> DV -> IMPACT
add_line_arrow(slide, Inches(5.5), Inches(3.6), Inches(6.0), Inches(3.6), thickness=2.0)
add_line_arrow(slide, Inches(8.4), Inches(3.6), Inches(8.9), Inches(3.6), thickness=2.0)


# MODERATING VARIABLE
t_mod = "BIẾN ĐIỀU TIẾT\nMức độ phát triển KT số\n\n• Hạ tầng số của địa phương\n• Mức độ ứng dụng TMĐT\n• Tính phổ cập công nghệ"
add_box(slide, t_mod, Inches(4.5), Inches(1.2), Inches(2.6), Inches(1.3), font_size=10, fill_color=(245, 245, 245))
# Drop Arrow
add_line_arrow(slide, Inches(5.8), Inches(2.5), Inches(5.8), Inches(3.6), dashed=True, thickness=1.5)


# CONTROL VARIABLE
t_ctrl = "BIẾN KIỂM SOÁT\n\n• Loại địa phương (Tỉnh/Thành)\n• Quy mô kinh tế (GRDP)\n• Đặc thù ngành nghề"
add_box(slide, t_ctrl, Inches(6.0), Inches(4.9), Inches(2.4), Inches(1.3), font_size=10, fill_color=(245, 245, 245))
# Up Arrow
add_line_arrow(slide, Inches(7.2), Inches(4.9), Inches(7.2), Inches(4.6), dashed=True, thickness=1.5)


# FEEDBACK LOOP
t_fb = "PHẢN HỒI CHÍNH SÁCH\nĐiều chỉnh chính sách   •   Hoàn thiện pháp luật   •   Nâng cấp nguồn lực   •   Đổi mới mô hình quản lý"
add_box(slide, t_fb, Inches(1.4), Inches(6.5), Inches(8.7), Inches(0.55), font_size=10)

# Drop arrow from Impact (X = 10.1, Bottom = 4.4)
add_line_arrow(slide, Inches(10.1), Inches(4.4), Inches(10.1), Inches(6.5), dashed=True, thickness=1.5)

# Up arrow from Feedback to IV Container Bottom (X = 1.45, Container Bottom = 6.2)
add_line_arrow(slide, Inches(1.45), Inches(6.5), Inches(1.45), Inches(6.2), dashed=True, thickness=1.5)


# ARROWS FROM THEORY
add_line_arrow(slide, Inches(2.5), Inches(1.1), Inches(2.5), Inches(1.2), dashed=True)
add_line_arrow(slide, Inches(5.8), Inches(1.1), Inches(5.8), Inches(1.2), dashed=True)
add_line_arrow(slide, Inches(7.2), Inches(1.1), Inches(7.2), Inches(2.7), dashed=True)


prs.save("MoHinhNghienCuu_DinhLuong_ChiTiet.pptx")
print("Saved detailed SEM model with items and fixed arrows.")
