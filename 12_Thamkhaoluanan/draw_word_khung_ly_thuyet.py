import docx
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import nsdecls
from docx.oxml import parse_xml

def set_cell_background(cell, color_hex):
    # color_hex should be string without '#' like 'E2F0D9'
    shading_elm = parse_xml(r'<w:shd {} w:fill="{}"/>'.format(nsdecls('w'), color_hex))
    cell._tc.get_or_add_tcPr().append(shading_elm)

def add_arrow_paragraph(doc, direction='down'):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run('⬇' if direction=='down' else '⬆')
    run.font.size = Pt(24)
    run.font.color.rgb = RGBColor(89, 89, 89)

doc = Document()
# Set A4 margins
sections = doc.sections
for section in sections:
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

# Title
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run("KHUNG LÝ THUYẾT NGHIÊN CỨU")
run.font.name = 'Times New Roman'
run.font.size = Pt(16)
run.font.bold = True

def create_full_width_box(doc, text, bg_color):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'
    cell = table.cell(0,0)
    set_cell_background(cell, bg_color)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(11)
    if ":" in text or "LÝ THUYẾT" in text or "KHÁI NIỆM" in text:
        # bold the first line roughly
        pass # simplifying
    return table

# 1. Lý thuyết
t1 = "LÝ THUYẾT NỀN TẢNG\n\nLý thuyết quản lý công   |   Lý thuyết quản lý công mới   |   Lý thuyết quản trị số\nLý thuyết quản trị thích ứng   |   Lý thuyết thể chế   |   Lý thuyết về quyền SHTT"
create_full_width_box(doc, t1, 'D9E1F2')

add_arrow_paragraph(doc)

# 2. Khái niệm
t2 = "KHÁI NIỆM CỐT LÕI\n\nQuyền sở hữu trí tuệ  -  Quản lý nhà nước  -  Kinh tế số  -  Quản trị số\nQuản lý nhà nước về quyền sở hữu trí tuệ trong nền kinh tế số"
create_full_width_box(doc, t2, 'FFF2CC')

add_arrow_paragraph(doc)

# 3. Mô hình nhân quả
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run("MÔ HÌNH NHÂN QUẢ GIỮA CÁC BIẾN")
run.font.name = 'Times New Roman'
run.font.bold = True
run.font.italic = True

table = doc.add_table(rows=3, cols=3)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
# Don't use Table Grid for the outer, we want specific cell borders but python-docx is tricky with selective borders.
# Let's just use Table Grid and remove borders where empty, or just leave it. Table Grid is fine, we can just hide borders or keep them.
# Actually, let's just make a 3x3 layout.
for row in table.rows:
    for cell in row.cells:
        cell.width = Inches(2.0)

# Merge col 2 (arrows) and col 3 (Điều tiết) vertically?
# Biến độc lập [0,0]
cell_dl = table.cell(0,0)
set_cell_background(cell_dl, 'E2F0D9')
cell_dl.text = "BIẾN ĐỘC LẬP:\nNỘI DUNG QUẢN LÝ NHÀ NƯỚC\n\nX1. Xây dựng CL, CS, pháp luật\nX2. Tổ chức thực hiện\nX3. Hỗ trợ thương mại hóa\nX4. Thanh tra, thực thi"

# Arrow down [1,0]
# We'll just put a down arrow in a text block, but wait, Biến trung gian needs to be in [1,0] if we stack them.
# Stacking them in Col 0 means 3 rows.
# But we need arrows BETWEEN them. So 5 rows!
# Let's redo table structure.
table._element.getparent().remove(table._element)

table = doc.add_table(rows=5, cols=3)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
widths = [Inches(3.0), Inches(0.5), Inches(2.5)]
for row in table.rows:
    row.cells[0].width = widths[0]
    row.cells[1].width = widths[1]
    row.cells[2].width = widths[2]

def format_cell(cell, text, bg_color):
    set_cell_background(cell, bg_color)
    cell.text = text
    for p in cell.paragraphs:
        for r in p.runs:
            r.font.name = 'Times New Roman'
            r.font.size = Pt(11)

# Biến độc lập
format_cell(table.cell(0,0), "BIẾN ĐỘC LẬP: NỘI DUNG QLNN\n\nX1. Xây dựng chiến lược, chính sách, pháp luật\nX2. Tổ chức thực hiện\nX3. Hỗ trợ xác lập, khai thác, thương mại hóa\nX4. Thanh tra, kiểm tra, giám sát, thực thi", 'E2F0D9')

# Arrow 1
table.cell(1,0).text = "⬇"
table.cell(1,0).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

# Biến trung gian
format_cell(table.cell(2,0), "BIẾN TRUNG GIAN: NĂNG LỰC QLNN\n\nM1. Năng lực hoạch định chính sách\nM2. Năng lực tổ chức thực hiện\nM3. Năng lực phối hợp liên ngành\nM4. Năng lực ứng dụng công nghệ số\nM5. Năng lực thanh tra, giám sát, thực thi", 'FCE4D6')

# Arrow 2
table.cell(3,0).text = "⬇"
table.cell(3,0).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

# Biến phụ thuộc
format_cell(table.cell(4,0), "BIẾN PHỤ THUỘC: CHẤT LƯỢNG QLNN\n\nY1. Hiệu lực\nY2. Hiệu quả\nY3. Tính phù hợp\nY4. Tính đồng bộ và thống nhất\nY5. Tính minh bạch", 'E2F0D9')

# Merge Col 2 (Biến điều tiết) across all 5 rows
a = table.cell(0,2)
b = table.cell(4,2)
a.merge(b)
format_cell(a, "BIẾN ĐIỀU TIẾT\n\nNHÓM THỂ CHẾ:\n- Mức độ hoàn thiện pháp luật\n- Chính sách phát triển KT số\n\nNHÓM CÔNG NGHỆ:\n- Hạ tầng số, AI, Dữ liệu, Nền tảng\n\nNHÓM MÔI TRƯỜNG:\n- Trình độ phát triển kinh tế số\n- Năng lực đổi mới sáng tạo\n- Hội nhập, Nhận thức xã hội", 'EDE2F6')
a.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

# Merge Col 1 (Arrows pointing left) across all 5 rows
c = table.cell(0,1)
d = table.cell(4,1)
c.merge(d)
c.text = "⬅"
c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

add_arrow_paragraph(doc)

# 4. Tác động
t4 = "KẾT QUẢ / TÁC ĐỘNG\n\nO1. Bảo vệ hiệu quả quyền SHTT   |   O2. Thúc đẩy đổi mới sáng tạo   |   O3. Phát triển tài sản trí tuệ\nO4. Nâng cao năng lực cạnh tranh quốc gia   |   O5. Thúc đẩy phát triển kinh tế số"
create_full_width_box(doc, t4, 'FFE699')

add_arrow_paragraph(doc)

# 5. Feedback
t5 = "VÒNG PHẢN HỒI CHÍNH SÁCH (FEEDBACK)\n\nKết quả thực hiện ➔ Đánh giá ➔ Điều chỉnh (Chính sách, Thể chế, Công cụ) ➔ Hoàn thiện QLNN\n\n(Lưu ý: Quá trình này sẽ quay ngược trở lại tác động vào Nội dung QLNN ở trên)"
create_full_width_box(doc, t5, 'F2F2F2')

doc.save('KhungLyThuyet_GS.docx')
print("Saved DOCX layout.")
