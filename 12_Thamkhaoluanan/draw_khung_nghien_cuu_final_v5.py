from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_AUTO_SIZE
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls

prs = Presentation()
prs.slide_width = Inches(11.69)
prs.slide_height = Inches(8.27)
slide = prs.slides.add_slide(prs.slide_layouts[6])

def add_box(slide, text, left, top, width, height, font_size=10, fill_color=(255, 255, 255), border_color=(0,0,0), border_width=1.5, bold_header=True, dash_style=None, italic_subheader=False):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    if fill_color:
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(*fill_color)
    else:
        shape.fill.background()
        
    shape.line.color.rgb = RGBColor(*border_color)
    shape.line.width = Pt(border_width)
    if dash_style:
        shape.line.dash_style = dash_style
    
    tf = shape.text_frame
    tf.text = text
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    
    for i, p in enumerate(tf.paragraphs):
        if p.text.startswith("•") or p.text.startswith(" ⬇"):
            p.alignment = PP_ALIGN.CENTER if p.text.startswith(" ⬇") else PP_ALIGN.LEFT
            if p.text.startswith("•"): p.level = 0
        else:
            p.alignment = PP_ALIGN.CENTER
            
        for run in p.runs:
            run.font.name = 'Times New Roman'
            run.font.size = Pt(font_size)
            run.font.color.rgb = RGBColor(0, 0, 0)
            if bold_header and i == 0:
                run.font.bold = True
            if italic_subheader and i == 1 and p.text != "":
                run.font.italic = True
                run.font.bold = True
    return shape

def add_line_arrow(slide, begin_x, begin_y, end_x, end_y, dashed=False, arrow_end=True, arrow_start=False, thickness=1.5):
    connector = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, begin_x, begin_y, end_x, end_y)
    connector.line.color.rgb = RGBColor(0,0,0)
    connector.line.width = Pt(thickness)
    if dashed:
        connector.line.dash_style = 4
        
    ln = connector.element.spPr.ln
    if ln is not None:
        if arrow_end:
            tailEnd = parse_xml(r'<a:tailEnd type="triangle" w="med" len="med" %s/>' % nsdecls('a'))
            ln.append(tailEnd)
        if arrow_start:
            headEnd = parse_xml(r'<a:headEnd type="triangle" w="med" len="med" %s/>' % nsdecls('a'))
            ln.append(headEnd)
    return connector


# CONTEXT BOX
ctx_shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.1), Inches(0.4), Inches(11.49), Inches(7.5))
ctx_shape.fill.solid()
ctx_shape.fill.fore_color.rgb = RGBColor(248, 252, 255)
ctx_shape.line.color.rgb = RGBColor(0, 102, 204)
ctx_shape.line.width = Pt(2.0)
ctx_shape.line.dash_style = 4

lbl_ctx = slide.shapes.add_textbox(Inches(0.1), Inches(0.5), Inches(11.49), Inches(0.5))
lbl_ctx.text_frame.text = "NỀN KINH TẾ SỐ (Bối cảnh nghiên cứu bao trùm toàn bộ hệ thống)"
lbl_ctx.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
lbl_ctx.text_frame.paragraphs[0].runs[0].font.name = 'Times New Roman'
lbl_ctx.text_frame.paragraphs[0].runs[0].font.size = Pt(14)
lbl_ctx.text_frame.paragraphs[0].runs[0].font.bold = True
lbl_ctx.text_frame.paragraphs[0].runs[0].font.color.rgb = RGBColor(0, 51, 153)


# ---------------------------------------------------------
# COL 1: INPUT CONTAINER
add_box(slide, "INPUT\n(Các biến độc lập)", Inches(0.2), Inches(1.2), Inches(2.1), Inches(5.0), font_size=11, fill_color=None, dash_style=4)

ivs = [
    "Khung pháp lý",
    "Nguồn nhân lực",
    "Nguồn lực tài chính",
    "Hạ tầng số",
    "Dữ liệu số",
    "Phối hợp thể chế"
]

iv_y_centers = []
start_y = 1.7
for i, iv in enumerate(ivs):
    y = start_y + i * 0.7
    shape = add_box(slide, iv, Inches(0.3), Inches(y), Inches(1.9), Inches(0.55), font_size=11, bold_header=False, fill_color=(255, 255, 255))
    iv_y_centers.append(y + 0.275)
# ---------------------------------------------------------

# COL 2: PROCESS (X = 2.55)
t_mv = "PROCESS\nNăng lực quản trị số\n(Biến trung gian)\n\n• Hoạch định chính sách số\n• Năng lực thực thi & vận hành\n• Quản trị nền tảng SHTT"
add_box(slide, t_mv, Inches(2.55), Inches(2.3), Inches(1.9), Inches(2.6), font_size=10, fill_color=(255, 255, 255), italic_subheader=True)

# COL 3: OUTPUT (X = 4.8)
t_out = "OUTPUT\n(Kết quả đầu ra)\n\n• CSDL SHTT số\n• Dịch vụ công trực tuyến\n• Hồ sơ điện tử\n• Kết nối & chia sẻ dữ liệu\n• Phát hiện vi phạm AI"
add_box(slide, t_out, Inches(4.8), Inches(2.3), Inches(1.9), Inches(2.6), font_size=10, fill_color=(255, 255, 255))

# COL 4: OUTCOME (X = 7.05)
t_dv = "OUTCOME\nHiệu quả QLNN về SHTT\n(Biến phụ thuộc)\n\n• Hiệu lực quản lý\n• Hiệu quả quản lý\n• Tính minh bạch\n• Khả năng thích ứng số\n• Mức độ hài lòng"
add_box(slide, t_dv, Inches(7.05), Inches(2.3), Inches(1.9), Inches(2.6), font_size=10, fill_color=(255, 255, 255), italic_subheader=True)

# COL 5: IMPACT (X = 9.3)
t_imp = "IMPACT\n(Tác động vĩ mô)\n\nHiệu quả bảo vệ quyền SHTT\n ⬇ \nĐổi mới sáng tạo\n ⬇ \nGia tăng giá trị TSTT\n ⬇ \nNâng cao năng lực cạnh tranh\n ⬇ \nPhát triển kinh tế số"
shape_imp = add_box(slide, t_imp, Inches(9.3), Inches(1.7), Inches(2.1), Inches(3.8), font_size=9, fill_color=(255, 255, 255))
# Center align all paragraphs in Impact box for the cascade effect
for p in shape_imp.text_frame.paragraphs:
    p.alignment = PP_ALIGN.CENTER


# MAIN ARROWS (Center Y = 3.6)
for y in iv_y_centers:
    add_line_arrow(slide, Inches(2.2), Inches(y), Inches(2.55), Inches(3.6), thickness=1.5)

add_line_arrow(slide, Inches(4.45), Inches(3.6), Inches(4.8), Inches(3.6), thickness=2.0)
add_line_arrow(slide, Inches(6.7), Inches(3.6), Inches(7.05), Inches(3.6), thickness=2.0)
add_line_arrow(slide, Inches(8.95), Inches(3.6), Inches(9.3), Inches(3.6), thickness=2.0)


# MODERATING VARIABLE
t_mod = "BIẾN ĐIỀU TIẾT\nMức độ phát triển KT số"
add_box(slide, t_mod, Inches(5.8), Inches(1.1), Inches(2.2), Inches(0.8), font_size=10, fill_color=(255, 255, 255))
# Drop Arrow pointing to the transition between OUTPUT and OUTCOME
add_line_arrow(slide, Inches(6.9), Inches(1.9), Inches(6.9), Inches(3.6), dashed=True, thickness=1.5)


# FEEDBACK LOOP
t_fb = "POLICY FEEDBACK (Phản hồi chính sách)\nĐiều chỉnh chính sách   •   Hoàn thiện pháp luật   •   Nâng cấp nguồn lực   •   Đổi mới mô hình quản lý"
add_box(slide, t_fb, Inches(1.2), Inches(6.8), Inches(9.1), Inches(0.6), font_size=10)

# Drop arrow from IMPACT (X = 10.35, Bottom = 5.5)
add_line_arrow(slide, Inches(10.35), Inches(5.5), Inches(10.35), Inches(7.1), dashed=True, thickness=1.5)
add_line_arrow(slide, Inches(10.35), Inches(7.1), Inches(10.3), Inches(7.1), dashed=True, thickness=1.5, arrow_end=False) # Corner connection

# Up arrow from FEEDBACK to INPUT
add_line_arrow(slide, Inches(1.2), Inches(7.1), Inches(1.25), Inches(6.2), dashed=True, thickness=1.5)


prs.save("KhungNghienCuu_DinhLuong_Va_LyThuyet_SieuCap.pptx")
print("Saved the ultimate integrated framework.")
