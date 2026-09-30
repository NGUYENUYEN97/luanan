from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_AUTO_SIZE

prs = Presentation()
prs.slide_width = Inches(11.69) # A4 Landscape
prs.slide_height = Inches(8.27)
slide = prs.slides.add_slide(prs.slide_layouts[6])

def add_box(slide, text, left, top, width, height, font_size=11, bold=False):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(255, 255, 255)
    shape.line.color.rgb = RGBColor(0, 0, 0)
    shape.line.width = Pt(1.5)
    
    text_frame = shape.text_frame
    text_frame.text = text
    text_frame.word_wrap = True
    text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    
    for i, paragraph in enumerate(text_frame.paragraphs):
        if "•" in paragraph.text or paragraph.text.startswith("-"):
            paragraph.alignment = PP_ALIGN.LEFT
        else:
            paragraph.alignment = PP_ALIGN.CENTER
            
        for run in paragraph.runs:
            run.font.name = 'Times New Roman'
            run.font.size = Pt(font_size)
            run.font.color.rgb = RGBColor(0, 0, 0)
            if i == 0:
                run.font.bold = True
            else:
                run.font.bold = bold
    return shape

def add_arrow(slide, left, top, width, height, direction='DOWN'):
    shape_type = MSO_SHAPE.DOWN_ARROW
    if direction == 'RIGHT': shape_type = MSO_SHAPE.RIGHT_ARROW
    elif direction == 'UP': shape_type = MSO_SHAPE.UP_ARROW
    elif direction == 'LEFT': shape_type = MSO_SHAPE.LEFT_ARROW
    
    shape = slide.shapes.add_shape(shape_type, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(0, 0, 0)
    shape.line.color.rgb = RGBColor(0, 0, 0)
    return shape

# Title
t_title = slide.shapes.add_textbox(Inches(0.5), Inches(0.1), Inches(10.69), Inches(0.4))
t_title.text_frame.text = "KHUNG LÝ THUYẾT VÀ MÔ HÌNH PHÂN TÍCH"
t_title.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
t_title.text_frame.paragraphs[0].runs[0].font.name = 'Times New Roman'
t_title.text_frame.paragraphs[0].runs[0].font.size = Pt(16)
t_title.text_frame.paragraphs[0].runs[0].font.bold = True

# --- TOP AREA (Theories & Concepts) ---
t1_text = "LÝ THUYẾT NỀN TẢNG\n\nLý thuyết quản lý công   •   Lý thuyết quản lý công mới   •   Lý thuyết quản trị số   •   Lý thuyết quản trị thích ứng   •   Lý thuyết thể chế   •   Lý thuyết về quyền sở hữu trí tuệ"
add_box(slide, t1_text, Inches(0.5), Inches(0.6), Inches(10.69), Inches(0.8), font_size=11)

t2_text = "KHÁI NIỆM CỐT LÕI\n\nQuyền sở hữu trí tuệ   •   Quản lý nhà nước   •   Kinh tế số   •   Quản trị số   •   Quản lý nhà nước về quyền sở hữu trí tuệ trong nền kinh tế số"
add_box(slide, t2_text, Inches(0.5), Inches(1.5), Inches(10.69), Inches(0.7), font_size=11)


# --- MODERATOR AREA (Biến điều tiết) ---
# Dashed container for Biến điều tiết
mod_container = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(2.2), Inches(2.4), Inches(7.2), Inches(1.5))
mod_container.fill.background()
mod_container.line.color.rgb = RGBColor(0, 0, 0)
mod_container.line.width = Pt(1.5)
mod_container.line.dash_style = 4
slide.shapes._spTree.remove(mod_container._element)
slide.shapes._spTree.insert(2, mod_container._element)

tb_mod = slide.shapes.add_textbox(Inches(2.3), Inches(2.45), Inches(4.0), Inches(0.3))
tb_mod.text_frame.text = "BIẾN ĐIỀU TIẾT (Các yếu tố ảnh hưởng)"
tb_mod.text_frame.paragraphs[0].runs[0].font.name = 'Times New Roman'
tb_mod.text_frame.paragraphs[0].runs[0].font.bold = True
tb_mod.text_frame.paragraphs[0].runs[0].font.size = Pt(11)

# Inner boxes for Moderator
m1 = "NHÓM THỂ CHẾ\n• Mức độ hoàn thiện pháp luật\n• Chính sách kinh tế số"
add_box(slide, m1, Inches(2.3), Inches(2.8), Inches(2.2), Inches(1.0), font_size=11)

m2 = "NHÓM CÔNG NGHỆ\n• Hạ tầng số, nền tảng\n• AI, dữ liệu lớn"
add_box(slide, m2, Inches(4.7), Inches(2.8), Inches(2.2), Inches(1.0), font_size=11)

m3 = "NHÓM MÔI TRƯỜNG\n• Trình độ KT số, hội nhập\n• Năng lực đổi mới sáng tạo"
add_box(slide, m3, Inches(7.1), Inches(2.8), Inches(2.2), Inches(1.0), font_size=11)


# --- CAUSAL FLOW AREA (Middle) ---
b_dl = "BIẾN ĐỘC LẬP\n(Nội dung quản lý nhà nước)\n\n• Xây dựng chiến lược, chính sách, pháp luật\n• Tổ chức thực hiện\n• Hỗ trợ xác lập, khai thác, thương mại hóa\n• Thanh tra, kiểm tra, giám sát, thực thi"
add_box(slide, b_dl, Inches(0.3), Inches(4.7), Inches(2.6), Inches(1.8), font_size=11)

# Arrow 1
add_arrow(slide, Inches(2.95), Inches(5.5), Inches(0.35), Inches(0.2), 'RIGHT')
# Arrow pointing down from Mod to Arrow 1
add_arrow(slide, Inches(3.125), Inches(3.9), Inches(0.15), Inches(1.5), 'DOWN')


b_tg = "BIẾN TRUNG GIAN\n(Năng lực quản lý nhà nước)\n\n• Năng lực hoạch định chính sách\n• Năng lực tổ chức thực hiện\n• Năng lực phối hợp liên ngành\n• Năng lực ứng dụng công nghệ số\n• Năng lực thanh tra, giám sát, thực thi"
add_box(slide, b_tg, Inches(3.35), Inches(4.7), Inches(2.5), Inches(1.8), font_size=11)

# Arrow 2
add_arrow(slide, Inches(5.9), Inches(5.5), Inches(0.35), Inches(0.2), 'RIGHT')
# Arrow pointing down from Mod to Arrow 2
add_arrow(slide, Inches(6.075), Inches(3.9), Inches(0.15), Inches(1.5), 'DOWN')


b_pt = "BIẾN PHỤ THUỘC\n(Chất lượng quản lý nhà nước)\n\n• Tính hiệu lực\n• Tính hiệu quả\n• Tính phù hợp\n• Tính đồng bộ và thống nhất\n• Tính minh bạch"
add_box(slide, b_pt, Inches(6.3), Inches(4.7), Inches(2.5), Inches(1.8), font_size=11)

# Arrow 3
add_arrow(slide, Inches(8.85), Inches(5.5), Inches(0.35), Inches(0.2), 'RIGHT')
# Arrow pointing down from Mod to Arrow 3 (Optional, maybe not needed if it only affects Content->Quality, but it looks balanced. Let's just point to Arrow 1 and 2).


b_kq = "KẾT QUẢ VÀ TÁC ĐỘNG\n\n• Bảo vệ hiệu quả quyền SHTT\n• Thúc đẩy đổi mới sáng tạo\n• Phát triển tài sản trí tuệ\n• Nâng cao năng lực cạnh tranh\n• Phát triển kinh tế số"
add_box(slide, b_kq, Inches(9.25), Inches(4.7), Inches(2.2), Inches(1.8), font_size=11)


# --- FEEDBACK LOOP ---
b_fb = "VÒNG PHẢN HỒI CHÍNH SÁCH\n\nĐánh giá hiệu quả   ➔   Điều chỉnh (chính sách, thể chế, công cụ)   ➔   Hoàn thiện hệ thống quản lý"
add_box(slide, b_fb, Inches(2.0), Inches(7.0), Inches(7.69), Inches(0.8), font_size=11)

# Lines for feedback
# Line down from Tác động
l_down = slide.shapes.add_shape(MSO_SHAPE.LINE_INVERSE, Inches(10.35), Inches(6.5), Inches(0), Inches(0.9))
l_down.line.color.rgb = RGBColor(0,0,0)
l_down.line.width = Pt(1.5)
l_down.line.dash_style = 4 

# Line left to feedback box
l_left = slide.shapes.add_shape(MSO_SHAPE.LINE_INVERSE, Inches(10.35), Inches(7.4), Inches(-0.66), Inches(0))
l_left.line.color.rgb = RGBColor(0,0,0)
l_left.line.width = Pt(1.5)
l_left.line.dash_style = 4

# Line from feedback box left to start
l_left2 = slide.shapes.add_shape(MSO_SHAPE.LINE_INVERSE, Inches(2.0), Inches(7.4), Inches(-0.4), Inches(0))
l_left2.line.color.rgb = RGBColor(0,0,0)
l_left2.line.width = Pt(1.5)
l_left2.line.dash_style = 4

# Line up to Biến độc lập
l_up = slide.shapes.add_shape(MSO_SHAPE.LINE_INVERSE, Inches(1.6), Inches(7.4), Inches(0), Inches(-0.9))
l_up.line.color.rgb = RGBColor(0,0,0)
l_up.line.width = Pt(1.5)
l_up.line.dash_style = 4

# Arrow pointing into Biến độc lập (UP)
arr_fb = slide.shapes.add_shape(MSO_SHAPE.UP_ARROW, Inches(1.5), Inches(6.5), Inches(0.2), Inches(0.2))
arr_fb.fill.solid()
arr_fb.fill.fore_color.rgb = RGBColor(0,0,0)
arr_fb.line.color.rgb = RGBColor(0,0,0)

prs.save("KhungLyThuyet_Landscape_ToiUu.pptx")
print("Saved Optimal Landscape PPTX layout.")
