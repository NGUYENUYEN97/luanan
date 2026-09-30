from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

prs = Presentation()
blank_slide_layout = prs.slide_layouts[6]
slide = prs.slides.add_slide(blank_slide_layout)

def add_box(slide, text, left, top, width, height, font_size=12, bold=False):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, left, top, width, height
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(255, 255, 255)
    shape.line.color.rgb = RGBColor(0, 0, 0)
    shape.line.width = Pt(1.5)
    
    text_frame = shape.text_frame
    text_frame.text = text
    text_frame.word_wrap = True
    
    for paragraph in text_frame.paragraphs:
        paragraph.alignment = PP_ALIGN.CENTER
        for run in paragraph.runs:
            run.font.name = 'Times New Roman'
            run.font.size = Pt(font_size)
            run.font.color.rgb = RGBColor(0, 0, 0)
            run.font.bold = bold
    return shape

def add_block_arrow(slide, left, top, width, height, direction='RIGHT'):
    shape_type = MSO_SHAPE.RIGHT_ARROW
    if direction == 'DOWN':
        shape_type = MSO_SHAPE.DOWN_ARROW
    elif direction == 'UP':
        shape_type = MSO_SHAPE.UP_ARROW
        
    shape = slide.shapes.add_shape(
        shape_type, left, top, width, height
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(0, 0, 0)
    shape.line.color.rgb = RGBColor(0, 0, 0)
    return shape

# Title
title_shape = slide.shapes.add_textbox(Inches(1), Inches(0.1), Inches(8), Inches(0.5))
title_tf = title_shape.text_frame
title_tf.text = "Khung Phân Tích Hoàn Chỉnh: Quản Lý Nhà Nước Về Sở Hữu Trí Tuệ Trong Kinh Tế Số"
title_tf.paragraphs[0].alignment = PP_ALIGN.CENTER
title_tf.paragraphs[0].runs[0].font.name = 'Times New Roman'
title_tf.paragraphs[0].runs[0].font.size = Pt(16)
title_tf.paragraphs[0].runs[0].font.bold = True

# 1. Khối Nhân Tố Ảnh Hưởng
text_factors = "CÁC NHÂN TỐ ẢNH HƯỞNG\n\n1. Đặc trưng của kinh tế số\n2. Hạ tầng công nghệ số quốc gia\n3. Bối cảnh hội nhập quốc tế\n4. Năng lực cơ quan nhà nước\n5. Nhận thức của xã hội"
box1 = add_box(slide, text_factors, Inches(0.2), Inches(1.8), Inches(2.3), Inches(3.5), font_size=12)

# 2. Khối Quản lý nhà nước
shape2 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(3.0), Inches(0.8), Inches(4.5), Inches(6.0))
shape2.fill.background()
shape2.line.color.rgb = RGBColor(0, 0, 0)
shape2.line.width = Pt(1.5)
shape2.line.dash_style = 4

# Inner elements
text_chu_the = "CHỦ THỂ QUẢN LÝ\n(Chính phủ, Bộ KH&CN, Bộ TT&TT, Cơ quan Tư pháp)"
box_ct = add_box(slide, text_chu_the, Inches(3.3), Inches(1.0), Inches(3.9), Inches(0.7), font_size=11, bold=True)

text_cong_cu = "CÔNG CỤ QUẢN LÝ\n(Pháp luật, Hành chính, Kinh tế, Công nghệ)"
box_cc = add_box(slide, text_cong_cu, Inches(3.3), Inches(2.0), Inches(3.9), Inches(0.6), font_size=11)

text_noi_dung = "NỘI DUNG QUẢN LÝ NHÀ NƯỚC\n\n1. Hoạch định và ban hành chính sách, pháp luật\n2. Tuyên truyền, phổ biến, giáo dục pháp luật\n3. Quản lý việc xác lập quyền sở hữu trí tuệ trực tuyến\n4. Thanh tra, kiểm tra và xử lý vi phạm trên mạng\n5. Hỗ trợ phát triển và Hợp tác quốc tế về sở hữu trí tuệ"
box_nd = add_box(slide, text_noi_dung, Inches(3.3), Inches(3.0), Inches(3.9), Inches(1.9), font_size=11)

text_khach_the = "KHÁCH THỂ QUẢN LÝ\n(Người sáng tạo số, Doanh nghiệp, Nền tảng số trung gian)"
box_kt = add_box(slide, text_khach_the, Inches(3.3), Inches(5.3), Inches(3.9), Inches(0.7), font_size=11, bold=True)

# 3. Khối Mục Tiêu
text_goals = "MỤC TIÊU ĐẠT ĐƯỢC\n\n1. Bảo vệ hiệu quả quyền lợi chủ thể\n2. Thúc đẩy hệ sinh thái đổi mới sáng tạo\n3. Phát triển an toàn, bền vững nền kinh tế số"
box3 = add_box(slide, text_goals, Inches(8.0), Inches(1.8), Inches(1.8), Inches(3.5), font_size=12)

# Arrows
add_block_arrow(slide, Inches(2.55), Inches(3.3), Inches(0.4), Inches(0.4), 'RIGHT') # Factors to QLNN
add_block_arrow(slide, Inches(5.1), Inches(1.75), Inches(0.3), Inches(0.2), 'DOWN') # CT to CC
add_block_arrow(slide, Inches(5.1), Inches(2.65), Inches(0.3), Inches(0.3), 'DOWN') # CC to ND
add_block_arrow(slide, Inches(5.1), Inches(4.95), Inches(0.3), Inches(0.3), 'DOWN') # ND to KT
add_block_arrow(slide, Inches(7.55), Inches(3.3), Inches(0.4), Inches(0.4), 'RIGHT') # QLNN to Goals

prs.save("KhungPhanTich_HoanChinh.pptx")
print("PowerPoint generated successfully.")
