"""
KHUNG NGHIÊN CỨU - POWERPOINT
Bố cục dọc + ngang:
  [LÝ THUYẾT NỀN TẢNG          ← full width]
              ↓
  [YẾU TỐ] → [NỘI DUNG QLNN] → [HIỆU QUẢ QLNN]   ← 3 hộp ngang
                                        ↓
                         [KẾT QUẢ VÀ TÁC ĐỘNG]     ← full width
                                        ↓
                         [PHẢN HỒI CHÍNH SÁCH]      ← full width
              ↑___________________________________|
"""
from pptx import Presentation
from pptx.util import Cm, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_CONNECTOR_TYPE
from pptx.oxml.ns import qn
from lxml import etree
import os

def rgb(h): return RGBColor.from_string(h)

prs = Presentation()
prs.slide_width  = Cm(33.87)
prs.slide_height = Cm(19.05)
slide = prs.slides.add_slide(prs.slide_layouts[6])

# ─── HELPERS ───────────────────────────────────────────────
def add_rect(x, y, w, h, fill='FFFFFF', border='222222',
             bpt=1.5, dashed=False):
    shp = slide.shapes.add_shape(1, Cm(x), Cm(y), Cm(w), Cm(h))
    if fill.upper() == 'NONE':
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = rgb(fill)
    shp.line.color.rgb = rgb(border)
    shp.line.width = Pt(bpt)
    if dashed:
        ln = shp.line._ln
        pd = etree.SubElement(ln, qn('a:prstDash'))
        pd.set('val', 'sysDash')
    return shp

def tf_write(shp, lines, anchor='middle'):
    tf = shp.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE if anchor == 'middle' else MSO_ANCHOR.TOP
    for i, ln in enumerate(lines):
        txt  = ln[0]; sz = ln[1]
        bold = ln[2] if len(ln) > 2 else False
        ital = ln[3] if len(ln) > 3 else False
        algn = ln[4] if len(ln) > 4 else PP_ALIGN.CENTER
        co   = ln[5] if len(ln) > 5 else '000000'
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = algn
        p.space_before = Pt(1); p.space_after = Pt(1)
        if txt == '':
            p.space_before = Pt(sz); continue
        run = p.add_run(); run.text = txt
        run.font.name  = 'Times New Roman'
        run.font.size  = Pt(sz)
        run.font.bold  = bold; run.font.italic = ital
        run.font.color.rgb = rgb(co)

def label(x, y, w, h, text, sz=9, bold=False, italic=False,
          align=PP_ALIGN.CENTER, color='000000'):
    txb = slide.shapes.add_textbox(Cm(x), Cm(y), Cm(w), Cm(h))
    tf  = txb.text_frame; tf.word_wrap = True
    p   = tf.paragraphs[0]; p.alignment = align
    run = p.add_run(); run.text = text
    run.font.name = 'Times New Roman'; run.font.size = Pt(sz)
    run.font.bold = bold; run.font.italic = italic
    run.font.color.rgb = rgb(color)

def arrow(x1, y1, x2, y2, color='444444', wpt=1.5,
          dashed=False, tail=True):
    con = slide.shapes.add_connector(
        MSO_CONNECTOR_TYPE.STRAIGHT,
        Cm(x1), Cm(y1), Cm(x2), Cm(y2))
    con.line.color.rgb = rgb(color)
    con.line.width = Pt(wpt)
    ln = con.line._ln
    if dashed:
        pd = etree.SubElement(ln, qn('a:prstDash'))
        pd.set('val', 'sysDash')
    he = etree.SubElement(ln, qn('a:headEnd'))
    he.set('type', 'none')
    if tail:
        te = etree.SubElement(ln, qn('a:tailEnd'))
        te.set('type', 'triangle')
        te.set('w', 'med'); te.set('len', 'med')

# ─── TỌA ĐỘ ────────────────────────────────────────────────
SL = 1.0    # left margin
ST = 0.5    # top margin
SW = 31.87  # usable width

# ─── 0. KHUNG BỐI CẢNH ─────────────────────────────────────
add_rect(SL-0.3, ST-0.1, SW+0.6, 18.2,
         fill='F9F9F9', border='999999', bpt=0.8, dashed=True)
label(SL, ST, 14, 0.45,
      'PHẠM VI NGHIÊN CỨU: QLNN VỀ QUYỀN SỞ HỮU TRÍ TUỆ TRONG NỀN KINH TẾ SỐ',
      sz=9, bold=True, italic=True, color='555555', align=PP_ALIGN.LEFT)

# ─── 1. LÝ THUYẾT NỀN TẢNG ────────────────────────────────
LT_Y = ST + 0.5; LT_H = 1.9
s1 = add_rect(SL, LT_Y, SW, LT_H, fill='E8E8E8', border='444444', bpt=1.5)
tf_write(s1, [
    ('LÝ THUYẾT NỀN TẢNG', 13, True),
    ('', 4),
    ('Lý thuyết quản lý công mới  •  Lý thuyết thể chế  •  Lý thuyết quản trị số  •  Lý thuyết quyền sở hữu trí tuệ  •  Lý thuyết về SHTT trong KTS',
     10, False, False, PP_ALIGN.CENTER, '333333'),
])

# ─── 2. MŨI TÊN ↓ (LT → 3 hộp) ────────────────────────────
BOX_Y = LT_Y + LT_H + 0.5
arrow(SL + SW/2, LT_Y + LT_H, SL + SW/2, BOX_Y, wpt=1.5, tail=True)

# ─── 3. BA HỘP NGANG ────────────────────────────────────────
BOX_H = 8.5
ARR   = 0.7   # khoảng mũi tên
GAP   = 0.3
BOX_W = (SW - 2*(ARR + GAP)) / 3

x_yt = SL
x_nd = x_yt + BOX_W + ARR + GAP
x_hq = x_nd + BOX_W + ARR + GAP

# Hộp YẾU TỐ
s_yt = add_rect(x_yt, BOX_Y, BOX_W, BOX_H,
                fill='FFFFFF', border='111111', bpt=2.0)
tf_write(s_yt, [
    ('CÁC YẾU TỐ', 12, True),
    ('ẢNH HƯỞNG', 12, True),
    ('', 6),
    ('(1) Khung pháp lý và chính sách', 10, False, False, PP_ALIGN.LEFT),
    ('(2) Năng lực tổ chức và cán bộ', 10, False, False, PP_ALIGN.LEFT),
    ('(3) Hạ tầng công nghệ số', 10, False, False, PP_ALIGN.LEFT),
    ('(4) Nguồn lực tài chính', 10, False, False, PP_ALIGN.LEFT),
    ('(5) Phối hợp thể chế', 10, False, False, PP_ALIGN.LEFT),
    ('(6) Nhận thức xã hội về quyền SHTT', 10, False, False, PP_ALIGN.LEFT),
], anchor='top')

arrow(x_yt + BOX_W, BOX_Y + BOX_H/2, x_nd, BOX_Y + BOX_H/2, wpt=2.0, tail=True)

# Hộp NỘI DUNG QLNN
s_nd = add_rect(x_nd, BOX_Y, BOX_W, BOX_H,
                fill='FFFFFF', border='111111', bpt=2.0)
tf_write(s_nd, [
    ('NỘI DUNG QLNN', 12, True),
    ('VỀ QUYỀN SHTT', 12, True),
    ('', 6),
    ('(1) Hoạch định chính sách,', 10, False, False, PP_ALIGN.LEFT),
    ('      chiến lược quyền SHTT', 10, False, False, PP_ALIGN.LEFT),
    ('(2) Xây dựng và hoàn thiện', 10, False, False, PP_ALIGN.LEFT),
    ('      hệ thống pháp luật', 10, False, False, PP_ALIGN.LEFT),
    ('(3) Tổ chức bộ máy quản lý', 10, False, False, PP_ALIGN.LEFT),
    ('(4) Thanh tra, kiểm tra,', 10, False, False, PP_ALIGN.LEFT),
    ('      xử lý vi phạm', 10, False, False, PP_ALIGN.LEFT),
    ('(5) Ứng dụng công nghệ số', 10, False, False, PP_ALIGN.LEFT),
    ('      trong quản lý quyền SHTT', 10, False, False, PP_ALIGN.LEFT),
], anchor='top')

arrow(x_nd + BOX_W, BOX_Y + BOX_H/2, x_hq, BOX_Y + BOX_H/2, wpt=2.0, tail=True)

# Hộp HIỆU QUẢ QLNN
s_hq = add_rect(x_hq, BOX_Y, BOX_W, BOX_H,
                fill='FFFFFF', border='111111', bpt=2.0)
tf_write(s_hq, [
    ('HIỆU QUẢ QLNN', 12, True),
    ('VỀ QUYỀN SHTT', 12, True),
    ('', 6),
    ('(1) Hiệu lực thực thi', 10, False, False, PP_ALIGN.LEFT),
    ('      pháp luật về quyền SHTT', 10, False, False, PP_ALIGN.LEFT),
    ('(2) Bảo vệ quyền SHTT', 10, False, False, PP_ALIGN.LEFT),
    ('      trên không gian số', 10, False, False, PP_ALIGN.LEFT),
    ('(3) Tính minh bạch,', 10, False, False, PP_ALIGN.LEFT),
    ('      công khai', 10, False, False, PP_ALIGN.LEFT),
    ('(4) Mức độ hài lòng', 10, False, False, PP_ALIGN.LEFT),
    ('      của chủ thể quyền SHTT', 10, False, False, PP_ALIGN.LEFT),
    ('(5) Khả năng thích ứng', 10, False, False, PP_ALIGN.LEFT),
    ('      với công nghệ mới', 10, False, False, PP_ALIGN.LEFT),
], anchor='top')

# ─── 4. KẾT QUẢ VÀ TÁC ĐỘNG (full width, dưới 3 hộp) ──────
KQ_Y = BOX_Y + BOX_H + 0.5
KQ_H = 2.8

# ↓ Mũi tên từ HIỆU QUẢ xuống KẾT QUẢ
arrow(x_hq + BOX_W/2, BOX_Y + BOX_H, x_hq + BOX_W/2, KQ_Y, wpt=1.8, tail=True)

s_kq = add_rect(SL, KQ_Y, SW, KQ_H, fill='FFFFFF', border='111111', bpt=2.0)
tf_write(s_kq, [
    ('KẾT QUẢ VÀ TÁC ĐỘNG', 12, True),
    ('', 5),
    ('(1) Thúc đẩy đổi mới sáng tạo, phát triển kinh tế tri thức          '
     '(2) Phát triển kinh tế số bền vững, đảm bảo công bằng số',
     10, False, False, PP_ALIGN.LEFT),
    ('(3) Bảo vệ quyền lợi hợp pháp của chủ thể sáng tạo          '
     '(4) Nâng cao năng lực cạnh tranh quốc gia',
     10, False, False, PP_ALIGN.LEFT),
    ('(5) Thu hút đầu tư nước ngoài và thúc đẩy chuyển giao công nghệ',
     10, False, False, PP_ALIGN.LEFT),
])

# ─── 5. PHẢN HỒI CHÍNH SÁCH (full width, dưới KẾT QUẢ) ────
PH_Y = KQ_Y + KQ_H + 0.4
PH_H = 2.0

# ↓ Mũi tên từ KẾT QUẢ xuống PHẢN HỒI
arrow(SL + SW/2, KQ_Y + KQ_H, SL + SW/2, PH_Y, wpt=1.8, tail=True)

s_ph = add_rect(SL, PH_Y, SW, PH_H,
                fill='F0F0F0', border='555555', bpt=1.5, dashed=True)
tf_write(s_ph, [
    ('PHẢN HỒI CHÍNH SÁCH', 12, True),
    ('', 4),
    ('Thực tiễn thực thi  →  Đánh giá, tổng kết kết quả  →  '
     'Điều chỉnh chính sách  →  Hoàn thiện QLNN về quyền SHTT',
     10, False, True, PP_ALIGN.CENTER, '444444'),
])

# ─── 6. VÒNG PHẢN HỒI hình U bên TRÁI ──────────────────────
# PHẢN HỒI → trái → lên → vào YẾU TỐ
OUT_X = SL - 0.7
# Ngang trái
arrow(SL, PH_Y + PH_H/2, OUT_X, PH_Y + PH_H/2,
      color='666666', wpt=1.2, dashed=True, tail=False)
# Dọc lên
arrow(OUT_X, PH_Y + PH_H/2, OUT_X, BOX_Y + BOX_H/2,
      color='666666', wpt=1.2, dashed=True, tail=False)
# Ngang vào YẾU TỐ
arrow(OUT_X, BOX_Y + BOX_H/2, x_yt, BOX_Y + BOX_H/2,
      color='666666', wpt=1.2, dashed=True, tail=True)
# Nhãn
label(OUT_X - 2.2, (PH_Y + BOX_Y + BOX_H) / 2 - 0.5, 2.0, 1.2,
      'phản hồi\nchính sách', sz=8, italic=True, color='777777',
      align=PP_ALIGN.RIGHT)

# ─── NGUỒN ──────────────────────────────────────────────────
label(SL, PH_Y + PH_H + 0.2, SW, 0.5,
      'Nguồn: Tác giả tự xây dựng trên cơ sở tổng quan lý thuyết',
      sz=9, italic=True, color='555555', align=PP_ALIGN.CENTER)

# ─── LƯU ────────────────────────────────────────────────────
out = r"g:\My Drive\Luan an uyen 26\LuanAn_TS_SHTT_KinhTeSo\12_Thamkhaoluanan\KhungNghienCuu_LuanAn_v2.pptx"
prs.save(out)
print(f"[OK] {out}")
print("  Layout: LT -> 3 hop ngang -> KET QUA -> PHAN HOI (doc)")
print("  Vong phan hoi: PHAN HOI -> trai -> len -> YEU TO")
