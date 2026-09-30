from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_AUTO_SIZE

prs = Presentation()
prs.slide_width = Inches(11.69)
prs.slide_height = Inches(8.27)
slide = prs.slides.add_slide(prs.slide_layouts[6])

def add_box(slide, text, left, top, width, height, fill_rgb, text_rgb=(0,0,0), font_size=11):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(*fill_rgb)
    shape.line.color.rgb = RGBColor(0, 0, 0)
    shape.line.width = Pt(1.0)
    
    tf = shape.text_frame
    tf.text = text
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    
    for i, p in enumerate(tf.paragraphs):
        if "•" in p.text or p.text.startswith("1.") or p.text.startswith("2.") or p.text.startswith("3.") or p.text.startswith("4."):
            p.alignment = PP_ALIGN.LEFT
        else:
            p.alignment = PP_ALIGN.CENTER
            
        for run in p.runs:
            run.font.name = 'Times New Roman'
            run.font.size = Pt(font_size)
            run.font.color.rgb = RGBColor(*text_rgb)
            if i == 0 or p.text.isupper():
                run.font.bold = True
    return shape

def add_arrow(slide, left, top, width, height, direction='DOWN', dashed=False):
    shape_type = MSO_SHAPE.DOWN_ARROW
    if direction == 'RIGHT': shape_type = MSO_SHAPE.RIGHT_ARROW
    elif direction == 'UP': shape_type = MSO_SHAPE.UP_ARROW
    elif direction == 'LEFT': shape_type = MSO_SHAPE.LEFT_ARROW
    
    shape = slide.shapes.add_shape(shape_type, left, top, width, height)
    
    if dashed:
        shape.fill.background() # transparent
        shape.line.color.rgb = RGBColor(0, 0, 0)
        shape.line.dash_style = 4
        shape.line.width = Pt(1.5)
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(89, 89, 89)
        shape.line.color.rgb = RGBColor(0, 0, 0)
    return shape

# Title
t_title = slide.shapes.add_textbox(Inches(0.5), Inches(0.05), Inches(10.69), Inches(0.4))
t_title.text_frame.text = "KHUNG LÝ THUYẾT NGHIÊN CỨU"
t_title.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
t_title.text_frame.paragraphs[0].runs[0].font.name = 'Times New Roman'
t_title.text_frame.paragraphs[0].runs[0].font.size = Pt(16)
t_title.text_frame.paragraphs[0].runs[0].font.bold = True

# --- Tầng 1 & 2 ---
t1 = "TẦNG 1: LÝ THUYẾT NỀN TẢNG\nPublic Management  •  New Public Management  •  Digital Governance  •  Adaptive Governance  •  Institutional Theory  •  Intellectual Property Theory"
add_box(slide, t1, Inches(1.5), Inches(0.4), Inches(8.69), Inches(0.6), (31, 73, 125), (255,255,255), 11)

add_arrow(slide, Inches(5.7), Inches(1.0), Inches(0.25), Inches(0.2))

t2 = "TẦNG 2: KHÁI NIỆM CỐT LÕI\nQuyền sở hữu trí tuệ  •  Quản lý nhà nước  •  Kinh tế số  •  Quản trị số  •  Quản lý nhà nước về quyền SHTT"
add_box(slide, t2, Inches(1.5), Inches(1.2), Inches(8.69), Inches(0.5), (230, 238, 248), (0,0,0), 11)

add_arrow(slide, Inches(5.7), Inches(1.7), Inches(0.25), Inches(0.2))

# --- CÁC KHỐI CHÍNH ---
# Tầng 3 (Trái)
t3 = "TẦNG 3: BIẾN ĐỘC LẬP\nCÁC YẾU TỐ ĐẦU VÀO\n\n• Nhóm thể chế: Pháp luật, Chính sách, Cơ chế phối hợp\n• Nhóm nguồn lực: Nguồn nhân lực, Tài chính, Hạ tầng số\n• Nhóm công nghệ: Dữ liệu số, AI, Nền tảng số"
add_box(slide, t3, Inches(0.3), Inches(1.9), Inches(3.2), Inches(1.8), (169, 208, 142))

# Mũi tên từ Tầng 3 sang Tầng 4
add_arrow(slide, Inches(3.5), Inches(2.6), Inches(0.3), Inches(0.25), 'RIGHT')


# Tầng 4 (Giữa)
t4 = "TẦNG 4: BIẾN TRUNG TÂM NGHIÊN CỨU\nQUẢN LÝ NHÀ NƯỚC VỀ QUYỀN SHTT TRONG NỀN KINH TẾ SỐ\n\n1. Xây dựng chiến lược, chính sách và pháp luật\n2. Tổ chức thực hiện\n3. Hỗ trợ xác lập, khai thác, chuyển giao và thương mại hóa TSTT\n4. Thanh tra, kiểm tra, giám sát và thực thi quyền SHTT"
add_box(slide, t4, Inches(3.8), Inches(1.9), Inches(4.0), Inches(1.6), (157, 195, 230))

add_arrow(slide, Inches(5.7), Inches(3.5), Inches(0.25), Inches(0.3))


# Tầng 5 (Giữa)
t5 = "TẦNG 5: BIẾN TRUNG GIAN\nNĂNG LỰC QLNN VỀ QUYỀN SHTT\n\n• Năng lực hoạch định chính sách\n• Năng lực tổ chức thực hiện\n• Năng lực phối hợp liên ngành\n• Năng lực chuyển đổi số\n• Năng lực thanh tra và thực thi"
add_box(slide, t5, Inches(4.0), Inches(3.8), Inches(3.6), Inches(1.3), (244, 177, 131))

add_arrow(slide, Inches(5.7), Inches(5.1), Inches(0.25), Inches(0.3))


# Tầng 6 (Giữa)
t6 = "TẦNG 6: BIẾN PHỤ THUỘC\nCHẤT LƯỢNG QLNN VỀ QUYỀN SỞ HỮU TRÍ TUỆ\n\n• Hiệu lực   • Hiệu quả   • Tính phù hợp   • Tính đồng bộ   • Tính minh bạch"
add_box(slide, t6, Inches(4.0), Inches(5.4), Inches(3.6), Inches(0.8), (255, 217, 102))

add_arrow(slide, Inches(5.7), Inches(6.2), Inches(0.25), Inches(0.3))


# Tầng 7 (Giữa)
t7 = "TẦNG 7: KẾT QUẢ VÀ TÁC ĐỘNG\n\n• Bảo vệ hiệu quả quyền SHTT       • Thúc đẩy đổi mới sáng tạo\n• Phát triển tài sản trí tuệ         • Nâng cao năng lực cạnh tranh quốc gia\n• Thúc đẩy phát triển kinh tế số"
add_box(slide, t7, Inches(3.8), Inches(6.5), Inches(4.0), Inches(0.9), (118, 204, 196))


# YẾU TỐ ĐIỀU TIẾT (Phải)
tdt = "YẾU TỐ ĐIỀU TIẾT\n\n• Mức độ phát triển kinh tế số\n• Năng lực đổi mới sáng tạo\n• Hội nhập quốc tế\n• Nhận thức xã hội\n• Mức độ cạnh tranh thị trường"
add_box(slide, tdt, Inches(8.3), Inches(2.9), Inches(3.0), Inches(2.0), (177, 160, 199))

# Mũi tên điều tiết nét đứt chỉ vào mũi tên Tầng 4->5
add_arrow(slide, Inches(7.8), Inches(3.6), Inches(0.5), Inches(0.25), 'LEFT', dashed=True)
l1 = slide.shapes.add_shape(MSO_SHAPE.LINE_INVERSE, Inches(7.8), Inches(3.725), Inches(-1.85), Inches(0))
l1.line.color.rgb = RGBColor(0,0,0); l1.line.width = Pt(1.5); l1.line.dash_style = 4

# Mũi tên điều tiết nét đứt chỉ vào mũi tên Tầng 5->6
add_arrow(slide, Inches(7.8), Inches(5.2), Inches(0.5), Inches(0.25), 'LEFT', dashed=True)
l2 = slide.shapes.add_shape(MSO_SHAPE.LINE_INVERSE, Inches(7.8), Inches(5.325), Inches(-1.85), Inches(0))
l2.line.color.rgb = RGBColor(0,0,0); l2.line.width = Pt(1.5); l2.line.dash_style = 4

# Vertical line connecting the two dashed arrows from the moderator box
l3 = slide.shapes.add_shape(MSO_SHAPE.LINE_INVERSE, Inches(8.3), Inches(3.725), Inches(0), Inches(1.6))
l3.line.color.rgb = RGBColor(0,0,0); l3.line.width = Pt(1.5); l3.line.dash_style = 4


# --- FEEDBACK LOOP ---
tfb = "PHẢN HỒI CHÍNH SÁCH\nKết quả thực hiện  ➔  Đánh giá  ➔  Điều chỉnh chính sách, thể chế, công cụ  ➔  Hoàn thiện QLNN"
add_box(slide, tfb, Inches(1.5), Inches(7.6), Inches(8.69), Inches(0.5), (217, 217, 217))

# Đường nét đứt phản hồi
# Từ trái khối Feedback vòng lên Tầng 4
fb_l1 = slide.shapes.add_shape(MSO_SHAPE.LINE_INVERSE, Inches(1.5), Inches(7.85), Inches(-1.3), Inches(0))
fb_l1.line.color.rgb = RGBColor(0,0,0); fb_l1.line.dash_style = 4; fb_l1.line.width = Pt(1.5)

fb_l2 = slide.shapes.add_shape(MSO_SHAPE.LINE_INVERSE, Inches(0.2), Inches(7.85), Inches(0), Inches(-5.15))
fb_l2.line.color.rgb = RGBColor(0,0,0); fb_l2.line.dash_style = 4; fb_l2.line.width = Pt(1.5)

fb_l3 = slide.shapes.add_shape(MSO_SHAPE.LINE_INVERSE, Inches(0.2), Inches(2.7), Inches(3.6), Inches(0))
fb_l3.line.color.rgb = RGBColor(0,0,0); fb_l3.line.dash_style = 4; fb_l3.line.width = Pt(1.5)

add_arrow(slide, Inches(3.5), Inches(2.58), Inches(0.3), Inches(0.25), 'RIGHT', dashed=True)
# Override color for the arrow
fb_arr = slide.shapes[-1]
fb_arr.line.color.rgb = RGBColor(0,0,0)


prs.save("KhungLyThuyet_7Tang_GS.pptx")
print("Saved 7-layer PPTX layout.")
