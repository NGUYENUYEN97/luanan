import docx
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import nsdecls
from docx.oxml import parse_xml

def set_cell_background(cell, color_hex):
    shading_elm = parse_xml(r'<w:shd {} w:fill="{}"/>'.format(nsdecls('w'), color_hex))
    cell._tc.get_or_add_tcPr().append(shading_elm)

def add_arrow_paragraph(doc, direction='down'):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run('⬇' if direction=='down' else '⬆')
    run.font.size = Pt(20)
    run.font.color.rgb = RGBColor(0, 0, 0)

doc = Document()
sections = doc.sections
for section in sections:
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

# Title
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run("KHUNG LÝ THUYẾT NGHIÊN CỨU")
run.font.name = 'Times New Roman'
run.font.size = Pt(14)
run.font.bold = True

def create_full_width_box(doc, text):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'
    cell = table.cell(0,0)
    set_cell_background(cell, 'FFFFFF')
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    
    lines = text.split('\n')
    for i, line in enumerate(lines):
        p = cell.paragraphs[i] if i < len(cell.paragraphs) else cell.add_paragraph()
        if "•" in line or "-" in line:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            # Add a small left indent for bullet points
            p.paragraph_format.left_indent = Inches(0.2)
        else:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            
        run = p.add_run(line)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(11)
        # Bold first line usually
        if i == 0 and line.isupper():
            run.font.bold = True
    return table

# 1. Lý thuyết
t1 = "LÝ THUYẾT NỀN TẢNG\n\n• Lý thuyết quản lý công   |   • Lý thuyết quản lý công mới   |   • Lý thuyết quản trị số\n• Lý thuyết quản trị thích ứng   |   • Lý thuyết thể chế   |   • Lý thuyết về quyền sở hữu trí tuệ"
create_full_width_box(doc, t1)

add_arrow_paragraph(doc)

# 2. Khái niệm
t2 = "KHÁI NIỆM CỐT LÕI\n\n• Quyền sở hữu trí tuệ   |   • Quản lý nhà nước   |   • Kinh tế số   |   • Quản trị số\n• Quản lý nhà nước về quyền sở hữu trí tuệ trong nền kinh tế số"
create_full_width_box(doc, t2)

add_arrow_paragraph(doc)

# 3. Mô hình nhân quả
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run("MÔ HÌNH NHÂN QUẢ GIỮA CÁC BIẾN")
run.font.name = 'Times New Roman'
run.font.bold = True
run.font.italic = True

table = doc.add_table(rows=5, cols=3)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.style = 'Table Grid'
widths = [Inches(3.3), Inches(0.4), Inches(2.57)]
for row in table.rows:
    row.cells[0].width = widths[0]
    row.cells[1].width = widths[1]
    row.cells[2].width = widths[2]

def format_cell(cell, text):
    set_cell_background(cell, 'FFFFFF')
    # Clear existing paragraph
    p = cell.paragraphs[0]
    p.text = ""
    lines = text.split('\n')
    for i, line in enumerate(lines):
        p = cell.paragraphs[i] if i < len(cell.paragraphs) else cell.add_paragraph()
        if "•" in line or line.startswith("-"):
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.left_indent = Inches(0.1)
        else:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            
        run = p.add_run(line)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(11)
        if i == 0 and line.isupper():
            run.font.bold = True

# Biến độc lập
format_cell(table.cell(0,0), "BIẾN ĐỘC LẬP\n(Nội dung quản lý nhà nước)\n\n• Xây dựng chiến lược, chính sách, pháp luật\n• Tổ chức thực hiện\n• Hỗ trợ xác lập, khai thác, thương mại hóa\n• Thanh tra, kiểm tra, giám sát, thực thi")

# Arrow 1
table.cell(1,0).text = "⬇"
table.cell(1,0).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
table.cell(1,0).paragraphs[0].runs[0].font.size = Pt(16)

# Biến trung gian
format_cell(table.cell(2,0), "BIẾN TRUNG GIAN\n(Năng lực quản lý nhà nước)\n\n• Năng lực hoạch định chính sách\n• Năng lực tổ chức thực hiện\n• Năng lực phối hợp liên ngành\n• Năng lực ứng dụng công nghệ số\n• Năng lực thanh tra, giám sát, thực thi")

# Arrow 2
table.cell(3,0).text = "⬇"
table.cell(3,0).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
table.cell(3,0).paragraphs[0].runs[0].font.size = Pt(16)

# Biến phụ thuộc
format_cell(table.cell(4,0), "BIẾN PHỤ THUỘC\n(Chất lượng quản lý nhà nước)\n\n• Tính hiệu lực\n• Tính hiệu quả\n• Tính phù hợp\n• Tính đồng bộ và thống nhất\n• Tính minh bạch")

# Merge Col 2 (Biến điều tiết) across all 5 rows
a = table.cell(0,2)
b = table.cell(4,2)
a.merge(b)
format_cell(a, "BIẾN ĐIỀU TIẾT\n(Các yếu tố ảnh hưởng)\n\nNhóm thể chế:\n• Mức độ hoàn thiện pháp luật\n• Chính sách phát triển kinh tế số\n\nNhóm công nghệ:\n• Hạ tầng số, AI, dữ liệu, nền tảng\n\nNhóm môi trường:\n• Trình độ phát triển kinh tế số\n• Năng lực đổi mới sáng tạo\n• Hội nhập, nhận thức xã hội")
a.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

# Merge Col 1 (Arrows pointing left) across all 5 rows
c = table.cell(0,1)
d = table.cell(4,1)
c.merge(d)
c.text = "⬅"
c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
c.paragraphs[0].runs[0].font.size = Pt(16)
c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

# Remove borders for empty cells containing arrows (1,0) and (3,0) if possible, but python-docx Table Grid adds all borders. It's fine for academic word drafts.

add_arrow_paragraph(doc)

# 4. Tác động
t4 = "KẾT QUẢ VÀ TÁC ĐỘNG\n\n• Bảo vệ hiệu quả quyền SHTT   |   • Thúc đẩy đổi mới sáng tạo   |   • Phát triển tài sản trí tuệ\n• Nâng cao năng lực cạnh tranh quốc gia   |   • Thúc đẩy phát triển kinh tế số"
create_full_width_box(doc, t4)

add_arrow_paragraph(doc)

# 5. Feedback
t5 = "VÒNG PHẢN HỒI CHÍNH SÁCH\n\nKết quả thực hiện  ➔  Đánh giá  ➔  Điều chỉnh (chính sách, thể chế, công cụ)  ➔  Hoàn thiện\n\n(Lưu ý: Quá trình này quay ngược trở lại tác động vào Nội dung quản lý nhà nước)"
create_full_width_box(doc, t5)

doc.save('KhungLyThuyet_BW.docx')
print("Saved BW DOCX layout.")
