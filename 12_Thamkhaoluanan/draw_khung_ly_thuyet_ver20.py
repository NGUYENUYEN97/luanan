from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_AUTO_SIZE
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls

prs = Presentation()
prs.slide_width = Inches(11.69)
prs.slide_height = Inches(8.27)
slide = prs.slides.add_slide(prs.slide_layouts[6])

def add_box(slide, text, left, top, width, height, font_size=11, fill_color=(250, 250, 250), border_color=(0,0,0)):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(*fill_color)
    shape.line.color.rgb = RGBColor(*border_color)
    shape.line.width = Pt(1.5)
    
    tf = shape.text_frame
    tf.text = text
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    
    for i, p in enumerate(tf.paragraphs):
        # Header is the first line
        if i == 0:
            p.alignment = PP_ALIGN.CENTER
        else:
            p.alignment = PP_ALIGN.LEFT
            p.level = 0
            
        for run in p.runs:
            run.font.name = 'Times New Roman'
            run.font.size = Pt(font_size)
            run.font.color.rgb = RGBColor(0, 0, 0)
            
            if i == 0:
                run.font.bold = True
    return shape

def add_line_arrow(slide, begin_x, begin_y, end_x, end_y, dashed=False, arrow_end=True, arrow_start=False, thickness=1.5):
    connector = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, begin_x, begin_y, end_x, end_y)
    connector.line.color.rgb = RGBColor(0,0,0)
    connector.line.width = Pt(thickness)
    if dashed:
        connector.line.dash_style = 4
        
    ln = connector.element.spPr.ln
    if ln is not None:
        if arrow_end:
            tailEnd = parse_xml(r'<a:tailEnd type="triangle" w="med" len="med" %s/>' % nsdecls('a'))
            ln.append(tailEnd)
        if arrow_start:
            headEnd = parse_xml(r'<a:headEnd type="triangle" w="med" len="med" %s/>' % nsdecls('a'))
            ln.append(headEnd)
    return connector


# --- TITLE ---
t_title = slide.shapes.add_textbox(Inches(0.5), Inches(0.2), Inches(10.69), Inches(0.5))
t_title.text_frame.text = "KHUNG NGHIÊN CỨU ĐỀ TÀI\nQuản lý nhà nước về quyền sở hữu trí tuệ trong nền kinh tế số"
t_title.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
t_title.text_frame.paragraphs[0].runs[0].font.name = 'Times New Roman'
t_title.text_frame.paragraphs[0].runs[0].font.size = Pt(16)
t_title.text_frame.paragraphs[0].runs[0].font.bold = True
t_title.text_frame.paragraphs[1].alignment = PP_ALIGN.CENTER
t_title.text_frame.paragraphs[1].runs[0].font.name = 'Times New Roman'
t_title.text_frame.paragraphs[1].runs[0].font.size = Pt(13)
t_title.text_frame.paragraphs[1].runs[0].font.italic = True


# --- BIẾN ĐỘC LẬP (Left) ---
t1 = "CÁC BIẾN ĐỘC LẬP\n\n• Khung pháp lý\n• Năng lực cán bộ\n• Hạ tầng số\n• Nguồn lực tài chính\n• Phối hợp thể chế\n• Nhận thức xã hội"
add_box(slide, t1, Inches(0.8), Inches(2.4), Inches(3.0), Inches(2.7), font_size=11)


# --- BIẾN TRUNG GIAN (Center) ---
t2 = "BIẾN TRUNG GIAN\n\n• Mức độ chuyển đổi số trong quản lý"
add_box(slide, t2, Inches(4.35), Inches(3.0), Inches(3.0), Inches(1.5), font_size=11)


# --- BIẾN PHỤ THUỘC (Right) ---
t3 = "BIẾN PHỤ THUỘC\n\n• Hiệu quả quản lý nhà nước về quyền SHTT"
add_box(slide, t3, Inches(7.9), Inches(3.0), Inches(3.0), Inches(1.5), font_size=11)


# --- ARROWS (Horizontal) ---
# Y Center is 3.75
add_line_arrow(slide, Inches(3.8), Inches(3.75), Inches(4.35), Inches(3.75), thickness=2.0) # Độc lập -> Trung gian
add_line_arrow(slide, Inches(7.35), Inches(3.75), Inches(7.9), Inches(3.75), thickness=2.0) # Trung gian -> Phụ thuộc


# --- BIẾN ĐIỀU TIẾT (Top Center) ---
tdt = "BIẾN ĐIỀU TIẾT\n\n• Mức độ phát triển kinh tế số"
add_box(slide, tdt, Inches(4.35), Inches(1.1), Inches(3.0), Inches(1.2), font_size=11)

# Arrow from Điều tiết down to Trung gian
add_line_arrow(slide, Inches(5.85), Inches(2.3), Inches(5.85), Inches(3.0), dashed=True, thickness=1.5)


# --- BIẾN KIỂM SOÁT (Bottom Right) ---
tks = "BIẾN KIỂM SOÁT\n\n• Loại địa phương\n• Quy mô nền kinh tế\n• Ngành nghề"
add_box(slide, tks, Inches(7.9), Inches(5.2), Inches(3.0), Inches(1.6), font_size=11)

# Arrow from Kiểm soát up to Phụ thuộc
add_line_arrow(slide, Inches(9.4), Inches(5.2), Inches(9.4), Inches(4.5), dashed=True, thickness=1.5)


prs.save("KhungNghienCuu_ChinhThuc_Final.pptx")
print("Saved final exact structural model.")
