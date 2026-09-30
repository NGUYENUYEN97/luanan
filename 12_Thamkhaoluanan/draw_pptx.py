from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_CONNECTOR

prs = Presentation()
# Use a blank slide layout
blank_slide_layout = prs.slide_layouts[6]
slide = prs.slides.add_slide(blank_slide_layout)

def add_box(slide, text, left, top, width, height, font_size=12, bold=False):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, left, top, width, height
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(255, 255, 255) # White fill
    shape.line.color.rgb = RGBColor(0, 0, 0) # Black border
    shape.line.width = Pt(1.5)
    
    text_frame = shape.text_frame
    text_frame.text = text
    text_frame.word_wrap = True
    
    for paragraph in text_frame.paragraphs:
        paragraph.alignment = PP_ALIGN.CENTER
        for run in paragraph.runs:
            run.font.name = 'Times New Roman'
            run.font.size = Pt(font_size)
            run.font.color.rgb = RGBColor(0, 0, 0)
            run.font.bold = bold
    return shape

def add_arrow(slide, start_shape, end_shape, start_side='right', end_side='left', dashed=False):
    # Calculate connection points based on sides
    # start_shape bounding box
    sx, sy, sw, sh = start_shape.left, start_shape.top, start_shape.width, start_shape.height
    ex, ey, ew, eh = end_shape.left, end_shape.top, end_shape.width, end_shape.height
    
    pts = {
        'left': (sx, sy + sh/2),
        'right': (sx + sw, sy + sh/2),
        'top': (sx + sw/2, sy),
        'bottom': (sx + sw/2, sy + sh)
    }
    pte = {
        'left': (ex, ey + eh/2),
        'right': (ex + ew, ey + eh/2),
        'top': (ex + ew/2, ey),
        'bottom': (ex + ew/2, ey + eh)
    }
    
    start_x, start_y = pts[start_side]
    end_x, end_y = pte[end_side]
    
    connector = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, start_x, start_y, end_x, end_y
    )
    connector.line.color.rgb = RGBColor(0, 0, 0)
    connector.line.width = Pt(1.5)
    # Add arrow head to the end
    # PowerPoint requires manipulating XML directly for arrowheads in python-pptx in some versions, 
    # but we can set line formatting if supported. Let's try to add arrowhead using standard line properties if possible.
    # Actually python-pptx doesn't have a direct API for arrow heads on standard connectors.
    # We will use an auto shape block arrow instead, or just a simple line.
    pass # Wait, let's draw MSO_SHAPE.RIGHT_ARROW instead of lines to make it easy and visible.

def add_block_arrow(slide, left, top, width, height, direction='RIGHT'):
    shape_type = MSO_SHAPE.RIGHT_ARROW
    if direction == 'DOWN':
        shape_type = MSO_SHAPE.DOWN_ARROW
    elif direction == 'UP':
        shape_type = MSO_SHAPE.UP_ARROW
        
    shape = slide.shapes.add_shape(
        shape_type, left, top, width, height
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(0, 0, 0)
    shape.line.color.rgb = RGBColor(0, 0, 0)
    return shape

# Title
title_shape = slide.shapes.add_textbox(Inches(1), Inches(0.2), Inches(8), Inches(0.5))
title_tf = title_shape.text_frame
title_tf.text = "Khung Phân Tích: Quản Lý Nhà Nước Về Sở Hữu Trí Tuệ Trong Kinh Tế Số"
title_tf.paragraphs[0].alignment = PP_ALIGN.CENTER
title_tf.paragraphs[0].runs[0].font.name = 'Times New Roman'
title_tf.paragraphs[0].runs[0].font.size = Pt(18)
title_tf.paragraphs[0].runs[0].font.bold = True

# 1. Khối Nhân Tố Ảnh Hưởng
text_factors = "CÁC NHÂN TỐ ẢNH HƯỞNG\n\n1. Đặc điểm của kinh tế số\n2. Bối cảnh hội nhập quốc tế\n3. Năng lực cơ quan quản lý\n4. Nhận thức của xã hội"
box1 = add_box(slide, text_factors, Inches(0.5), Inches(2.0), Inches(2.5), Inches(3.0), font_size=13)

# 2. Khối Quản lý nhà nước (Outer dashed box representation, we will just use a large box)
shape2 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(3.5), Inches(1.0), Inches(3.8), Inches(5.0))
shape2.fill.background()
shape2.line.color.rgb = RGBColor(0, 0, 0)
shape2.line.width = Pt(1.5)
shape2.line.dash_style = 4 # Dashed

# Inner elements of QLNN
text_chu_the = "CHỦ THỂ QUẢN LÝ\n(Chính phủ, Bộ KH&CN, Cục SHTT)"
box_ct = add_box(slide, text_chu_the, Inches(3.8), Inches(1.2), Inches(3.2), Inches(0.8), bold=True)

text_noi_dung = "NỘI DUNG QUẢN LÝ\n\n1. Ban hành chính sách, pháp luật\n2. Tổ chức bộ máy, xác lập quyền\n3. Thanh tra, kiểm tra, xử lý vi phạm\n4. Hỗ trợ thương mại hóa tài sản trí tuệ"
box_nd = add_box(slide, text_noi_dung, Inches(3.8), Inches(2.5), Inches(3.2), Inches(1.8))

text_khach_the = "KHÁCH THỂ QUẢN LÝ\n(Doanh nghiệp, Cá nhân, Nền tảng số trung gian)"
box_kt = add_box(slide, text_khach_the, Inches(3.8), Inches(4.8), Inches(3.2), Inches(0.8), bold=True)

# 3. Khối Mục Tiêu
text_goals = "MỤC TIÊU ĐẠT ĐƯỢC\n\n1. Bảo hộ hiệu quả\n2. Thúc đẩy đổi mới sáng tạo\n3. Phát triển nền kinh tế số"
box3 = add_box(slide, text_goals, Inches(7.8), Inches(2.0), Inches(1.8), Inches(3.0), font_size=13)

# Arrows
add_block_arrow(slide, Inches(3.05), Inches(3.3), Inches(0.4), Inches(0.4), 'RIGHT') # Factors to QLNN
add_block_arrow(slide, Inches(5.2), Inches(2.05), Inches(0.4), Inches(0.4), 'DOWN') # CT to ND
add_block_arrow(slide, Inches(5.2), Inches(4.35), Inches(0.4), Inches(0.4), 'DOWN') # ND to KT
add_block_arrow(slide, Inches(7.35), Inches(3.3), Inches(0.4), Inches(0.4), 'RIGHT') # QLNN to Goals

prs.save("KhungPhanTich.pptx")
print("PowerPoint generated successfully.")
