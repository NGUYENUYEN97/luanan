from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_AUTO_SIZE
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls

prs = Presentation()
prs.slide_width = Inches(11.69)
prs.slide_height = Inches(8.27)
slide = prs.slides.add_slide(prs.slide_layouts[6])

def add_box(slide, text, left, top, width, height, font_size=10, bold_first_line=True):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(255, 255, 255)
    shape.line.color.rgb = RGBColor(0, 0, 0)
    shape.line.width = Pt(1.5)
    
    tf = shape.text_frame
    tf.text = text
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    
    for i, p in enumerate(tf.paragraphs):
        if "•" in p.text or p.text.startswith("1.") or p.text.startswith("2.") or p.text.startswith("3.") or p.text.startswith("4."):
            p.alignment = PP_ALIGN.LEFT
        else:
            p.alignment = PP_ALIGN.CENTER
            
        for run in p.runs:
            run.font.name = 'Times New Roman'
            run.font.size = Pt(font_size)
            run.font.color.rgb = RGBColor(0, 0, 0)
            
            if bold_first_line and (i == 0 or (i == 1 and p.text.isupper())):
                run.font.bold = True
    return shape

def add_line_arrow(slide, begin_x, begin_y, end_x, end_y, dashed=False, arrow_end=True, arrow_start=False):
    connector = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, begin_x, begin_y, end_x, end_y)
    connector.line.color.rgb = RGBColor(0,0,0)
    connector.line.width = Pt(1.5)
    if dashed:
        connector.line.dash_style = 4
        
    ln = connector.element.spPr.ln
    if ln is not None:
        if arrow_end:
            tailEnd = parse_xml(r'<a:tailEnd type="triangle" w="med" len="med" %s/>' % nsdecls('a'))
            ln.append(tailEnd)
        if arrow_start:
            headEnd = parse_xml(r'<a:headEnd type="triangle" w="med" len="med" %s/>' % nsdecls('a'))
            ln.append(headEnd)
    return connector

# Title
t_title = slide.shapes.add_textbox(Inches(0.5), Inches(0.05), Inches(10.69), Inches(0.4))
t_title.text_frame.text = "KHUNG LÝ THUYẾT NGHIÊN CỨU"
t_title.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
t_title.text_frame.paragraphs[0].runs[0].font.name = 'Times New Roman'
t_title.text_frame.paragraphs[0].runs[0].font.size = Pt(16)
t_title.text_frame.paragraphs[0].runs[0].font.bold = True

# --- LÝ THUYẾT NỀN TẢNG ---
# Đưa thêm Lý thuyết tăng trưởng nội sinh (đặc trưng của kinh tế học về ĐMST)
t1 = "LÝ THUYẾT NỀN TẢNG\nLý thuyết quản lý công  •  Lý thuyết thể chế  •  Lý thuyết tăng trưởng nội sinh  •  Lý thuyết về quyền sở hữu trí tuệ"
add_box(slide, t1, Inches(1.5), Inches(0.5), Inches(8.69), Inches(0.6), font_size=10)
add_line_arrow(slide, Inches(6.0), Inches(1.1), Inches(6.0), Inches(1.5))


# --- BIẾN ĐỘC LẬP (Left) ---
t2 = "BIẾN ĐỘC LẬP\n(QUẢN LÝ NHÀ NƯỚC VỀ QUYỀN SHTT)\n\n1. Xây dựng chiến lược, chính sách, pháp luật\n2. Tổ chức thực hiện\n3. Quản lý hoạt động xác lập, bảo vệ và hỗ trợ khai thác, thương mại hóa TSTT\n4. Thanh tra, kiểm tra, giám sát, thực thi quyền"
add_box(slide, t2, Inches(0.3), Inches(1.9), Inches(3.6), Inches(2.0), font_size=10)
# Mũi tên từ Độc lập sang Trung gian
add_line_arrow(slide, Inches(3.9), Inches(2.9), Inches(4.3), Inches(2.9))


# --- BIẾN TRUNG GIAN (Center) ---
t3 = "BIẾN TRUNG GIAN\n(MÔI TRƯỜNG BẢO HỘ SHTT)\n\n• Mức độ tuân thủ và thực thi pháp luật\n• Tính minh bạch, thuận lợi của thủ tục hành chính\n• Giảm thiểu rủi ro và chi phí giao dịch"
add_box(slide, t3, Inches(4.3), Inches(1.5), Inches(3.4), Inches(1.3), font_size=10)
add_line_arrow(slide, Inches(6.0), Inches(2.8), Inches(6.0), Inches(3.2))


# --- BIẾN PHỤ THUỘC (Center) ---
t4 = "BIẾN PHỤ THUỘC\n(SỰ GIA TĂNG TÀI SẢN TRÍ TUỆ)\n\n• Tăng trưởng số lượng văn bằng bảo hộ\n• Tăng tỷ lệ thương mại hóa tài sản trí tuệ\n• Gia tăng giá trị tài sản vô hình của doanh nghiệp"
add_box(slide, t4, Inches(4.2), Inches(3.2), Inches(3.6), Inches(1.3), font_size=10)
add_line_arrow(slide, Inches(6.0), Inches(4.5), Inches(6.0), Inches(5.0))


# --- KẾT QUẢ/TÁC ĐỘNG (Center) ---
t5 = "KẾT QUẢ VÀ TÁC ĐỘNG\n(PHÁT TRIỂN KINH TẾ SỐ)\n\n• Thúc đẩy đổi mới sáng tạo toàn diện\n• Tăng trưởng kinh tế dựa trên tri thức\n• Nâng cao năng lực cạnh tranh quốc gia"
add_box(slide, t5, Inches(4.1), Inches(5.0), Inches(3.8), Inches(1.3), font_size=10)


# --- BIẾN ĐIỀU TIẾT (Right) ---
tdt = "BIẾN ĐIỀU TIẾT\n(CÁC YẾU TỐ MÔI TRƯỜNG VĨ MÔ)\n\n• Trình độ công nghệ (AI, Big Data)\n• Mức độ hội nhập quốc tế\n• Nhận thức của doanh nghiệp và xã hội\n• Cạnh tranh thị trường"
add_box(slide, tdt, Inches(8.3), Inches(2.2), Inches(3.0), Inches(1.8), font_size=10)

# Mũi tên điều tiết nét đứt chỉ vào mũi tên Trung gian -> Phụ thuộc (Y=3.0)
# Biến điều tiết ảnh hưởng đến toàn bộ quá trình chuyển hóa từ QLNN sang Kết quả
add_line_arrow(slide, Inches(8.3), Inches(3.0), Inches(6.0), Inches(3.0), dashed=True)
add_line_arrow(slide, Inches(8.3), Inches(4.75), Inches(6.0), Inches(4.75), dashed=True)
add_line_arrow(slide, Inches(8.3), Inches(3.0), Inches(8.3), Inches(4.75), dashed=True, arrow_end=False)


# --- FEEDBACK LOOP ---
tfb = "PHẢN HỒI CHÍNH SÁCH\nYêu cầu từ thực tiễn phát triển kinh tế ➔ Đánh giá hiệu quả quản lý ➔ Hoàn thiện chiến lược và chính sách QLNN"
add_box(slide, tfb, Inches(1.5), Inches(7.2), Inches(8.69), Inches(0.6), font_size=10)

# Từ Kết quả xuống Feedback
add_line_arrow(slide, Inches(6.0), Inches(6.3), Inches(6.0), Inches(7.2), dashed=True)

# Từ khối Feedback vòng lên Biến Độc lập (QLNN)
add_line_arrow(slide, Inches(1.5), Inches(7.5), Inches(0.15), Inches(7.5), dashed=True, arrow_end=False)
add_line_arrow(slide, Inches(0.15), Inches(7.5), Inches(0.15), Inches(1.2), dashed=True, arrow_end=False)
add_line_arrow(slide, Inches(0.15), Inches(1.2), Inches(2.1), Inches(1.2), dashed=True, arrow_end=False)
add_line_arrow(slide, Inches(2.1), Inches(1.2), Inches(2.1), Inches(1.9), dashed=True)

prs.save("KhungLyThuyet_KinhTe_Chuan.pptx")
print("Saved PPTX with Economic impact model.")
