"""
KHUNG PHÂN TÍCH - VẼ MỚI HOÀN TOÀN
Áp dụng đủ 6 điểm góp ý hội đồng phản biện.
Bố cục A4 ngang (29.7 x 21 cm), font Times New Roman.

Layout:
  Tiêu đề
  [LÝ THUYẾT NỀN TẢNG]                          ← full width
            ↓
  [NHÓM BÊN TRONG] → [NỘI DUNG QLNN] → [CHẤT LƯỢNG QLNN]
  [NHÓM BÊN NGOÀI]                              ↓
            ↑                            [KẾT QUẢ VÀ TÁC ĐỘNG]
            |                                    ↓
            └─────── [PHẢN HỒI CHÍNH SÁCH] ─────┘
"""
import sys; sys.stdout.reconfigure(encoding='utf-8')
from pptx import Presentation
from pptx.util import Cm, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree

def rgb(h): return RGBColor.from_string(h)

prs = Presentation()
prs.slide_width  = Cm(29.7)
prs.slide_height = Cm(21.0)
slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank

# ─── HELPERS ───────────────────────────────────────────────
def add_box(x, y, w, h, fill='FFFFFF', border='000000', bpt=1.2, dashed=False):
    shp = slide.shapes.add_shape(1, Cm(x), Cm(y), Cm(w), Cm(h))
    if fill.upper() == 'NONE':
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = rgb(fill)
    shp.line.color.rgb = rgb(border)
    shp.line.width = Pt(bpt)
    if dashed:
        ln = shp.line._ln
        pd = etree.SubElement(ln, qn('a:prstDash'))
        pd.set('val', 'sysDash')
    return shp

def write_text(shp, lines, anchor='top'):
    tf = shp.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP if anchor == 'top' else MSO_ANCHOR.MIDDLE
    for i, ln in enumerate(lines):
        txt   = ln[0]
        sz    = ln[1]
        bold  = ln[2] if len(ln) > 2 else False
        ital  = ln[3] if len(ln) > 3 else False
        algn  = ln[4] if len(ln) > 4 else PP_ALIGN.CENTER
        color = ln[5] if len(ln) > 5 else '000000'
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = algn
        p.space_before = Pt(0.5)
        p.space_after  = Pt(0.5)
        if txt == '':
            p.space_before = Pt(sz)
            continue
        run = p.add_run()
        run.text = txt
        run.font.name  = 'Times New Roman'
        run.font.size  = Pt(sz)
        run.font.bold  = bold
        run.font.italic = ital
        run.font.color.rgb = rgb(color)

def add_arrow(x1, y1, x2, y2, wpt=1.5, color='000000'):
    """Mũi tên từ (x1,y1) đến (x2,y2) cm"""
    cx = slide.shapes.add_connector(
        1, Cm(x1), Cm(y1), Cm(x2), Cm(y2))
    cx.line.color.rgb = rgb(color)
    cx.line.width = Pt(wpt)
    # Đầu mũi tên
    ln = cx.line._ln
    tail = etree.SubElement(ln, qn('a:tailEnd'))
    tail.set('type', 'triangle')
    tail.set('w', 'med')
    tail.set('len', 'med')
    return cx

def add_line(x1, y1, x2, y2, wpt=1.2, color='000000', dashed=False):
    """Đường kẻ không đầu mũi tên"""
    cx = slide.shapes.add_connector(
        1, Cm(x1), Cm(y1), Cm(x2), Cm(y2))
    cx.line.color.rgb = rgb(color)
    cx.line.width = Pt(wpt)
    if dashed:
        ln = cx.line._ln
        pd = etree.SubElement(ln, qn('a:prstDash'))
        pd.set('val', 'sysDash')
    return cx

def add_label(x, y, w, h, text, sz=9, bold=False, italic=False, color='000000', align=PP_ALIGN.CENTER):
    tb = slide.shapes.add_textbox(Cm(x), Cm(y), Cm(w), Cm(h))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name  = 'Times New Roman'
    run.font.size  = Pt(sz)
    run.font.bold  = bold
    run.font.italic = italic
    run.font.color.rgb = rgb(color)
    return tb

# ─── KÍCH THƯỚC CHUNG ─────────────────────────────────────
SL = 1.0   # slide left margin
SW = 27.7  # usable width
MID = SL + SW/2  # center

# ═══════════════════════════════════════════════════════════
# TIÊU ĐỀ
# ═══════════════════════════════════════════════════════════
add_label(SL, 0.2, SW, 0.7,
    'KHUNG PHÂN TÍCH: QUẢN LÝ NHÀ NƯỚC VỀ QUYỀN SỞ HỮU TRÍ TUỆ TRONG NỀN KINH TẾ SỐ',
    sz=13, bold=True, color='000000')

# ═══════════════════════════════════════════════════════════
# 1. LÝ THUYẾT NỀN TẢNG
# ═══════════════════════════════════════════════════════════
LT_Y = 1.2; LT_H = 1.6
s_lt = add_box(SL, LT_Y, SW, LT_H, fill='F0F0F0', border='333333')
write_text(s_lt, [
    ('LÝ THUYẾT NỀN TẢNG', 11, True),
    ('', 3),
    ('Lý thuyết quản lý công mới   •   Lý thuyết thể chế   •   Lý thuyết quản trị số   •   Lý thuyết quyền sở hữu trí tuệ',
     9.5, False, False, PP_ALIGN.CENTER, '333333'),
], anchor='middle')

# Mũi tên LT → xuống
add_arrow(MID, LT_Y + LT_H, MID, LT_Y + LT_H + 0.5)

# ═══════════════════════════════════════════════════════════
# 2. BA CỘT: YẾU TỐ | NỘI DUNG | CHẤT LƯỢNG
# ═══════════════════════════════════════════════════════════
ROW_Y = 3.3
COL_GAP = 0.8
COL_W = (SW - 2 * COL_GAP) / 3  # ~8.7 cm mỗi cột

x_yt = SL                      # cột trái
x_nd = SL + COL_W + COL_GAP    # cột giữa
x_cl = x_nd + COL_W + COL_GAP  # cột phải

YT_H = 11.5  # chiều cao cột yếu tố
ND_H = 11.5
CL_H = 7.0   # chất lượng thấp hơn

# ─── 2.1 CÁC YẾU TỐ ẢNH HƯỞNG (outer box) ────────────────
s_yt_outer = add_box(x_yt, ROW_Y, COL_W, YT_H, fill='NONE', border='333333', bpt=1.5)
write_text(s_yt_outer, [
    ('CÁC YẾU TỐ ẢNH HƯỞNG', 10.5, True),
])

# Nhóm bên trong
bt_y = ROW_Y + 1.3
bt_h = 4.5
s_bt = add_box(x_yt + 0.3, bt_y, COL_W - 0.6, bt_h, fill='E8F0FE', border='3367A0', bpt=1.0)
write_text(s_bt, [
    ('Nhóm yếu tố bên trong', 9, True, False, PP_ALIGN.CENTER, '1A3C6E'),
    ('', 4),
    ('(1) Khung pháp lý và chính sách', 9, False, False, PP_ALIGN.LEFT),
    ('(2) Năng lực tổ chức và cán bộ', 9, False, False, PP_ALIGN.LEFT),
    ('(3) Nguồn lực tài chính', 9, False, False, PP_ALIGN.LEFT),
    ('(4) Phối hợp thể chế', 9, False, False, PP_ALIGN.LEFT),
], anchor='top')

# Nhóm bên ngoài
bn_y = bt_y + bt_h + 0.5
bn_h = 5.0
s_bn = add_box(x_yt + 0.3, bn_y, COL_W - 0.6, bn_h, fill='FFF3E0', border='A06633', bpt=1.0)
write_text(s_bn, [
    ('Nhóm yếu tố bên ngoài', 9, True, False, PP_ALIGN.CENTER, '6E3C1A'),
    ('', 4),
    ('(5) Hạ tầng công nghệ số', 9, False, False, PP_ALIGN.LEFT),
    ('(6) Đặc trưng nền kinh tế số', 9, False, False, PP_ALIGN.LEFT),
    ('(7) Nhận thức xã hội', 9, False, False, PP_ALIGN.LEFT),
    ('     về quyền SHTT', 9, False, False, PP_ALIGN.LEFT),
    ('(8) Hội nhập quốc tế', 9, False, False, PP_ALIGN.LEFT),
], anchor='top')

# ─── 2.2 NỘI DUNG QLNN ───────────────────────────────────
s_nd = add_box(x_nd, ROW_Y, COL_W, ND_H, fill='FFFFFF', border='333333', bpt=1.5)
write_text(s_nd, [
    ('NỘI DUNG', 10.5, True),
    ('QUẢN LÝ NHÀ NƯỚC', 10.5, True),
    ('VỀ QUYỀN SHTT', 10.5, True),
    ('', 6),
    ('(1) Xây dựng chiến lược,', 9, False, False, PP_ALIGN.LEFT),
    ('     chính sách và pháp luật', 9, False, False, PP_ALIGN.LEFT),
    ('', 4),
    ('(2) Tổ chức thực hiện', 9, False, False, PP_ALIGN.LEFT),
    ('     chiến lược, chính sách', 9, False, False, PP_ALIGN.LEFT),
    ('', 4),
    ('(3) Hỗ trợ xác lập, khai thác,', 9, False, False, PP_ALIGN.LEFT),
    ('     chuyển giao và thương mại', 9, False, False, PP_ALIGN.LEFT),
    ('     hóa tài sản trí tuệ', 9, False, False, PP_ALIGN.LEFT),
    ('', 4),
    ('(4) Thanh tra, kiểm tra,', 9, False, False, PP_ALIGN.LEFT),
    ('     giám sát và thực thi', 9, False, False, PP_ALIGN.LEFT),
    ('     quyền SHTT', 9, False, False, PP_ALIGN.LEFT),
])

# ─── 2.3 CHẤT LƯỢNG QLNN ─────────────────────────────────
s_cl = add_box(x_cl, ROW_Y, COL_W, CL_H, fill='FFFFFF', border='333333', bpt=1.5)
write_text(s_cl, [
    ('CHẤT LƯỢNG', 10.5, True),
    ('QUẢN LÝ NHÀ NƯỚC', 10.5, True),
    ('VỀ QUYỀN SHTT', 10.5, True),
    ('', 6),
    ('Tiêu chí đánh giá:', 9, True, True, PP_ALIGN.LEFT, '333333'),
    ('', 3),
    ('(1) Hiệu lực', 9.5, False, False, PP_ALIGN.LEFT),
    ('(2) Hiệu quả', 9.5, False, False, PP_ALIGN.LEFT),
    ('(3) Tính phù hợp', 9.5, False, False, PP_ALIGN.LEFT),
    ('(4) Tính đồng bộ', 9.5, False, False, PP_ALIGN.LEFT),
    ('     và thống nhất', 9.5, False, False, PP_ALIGN.LEFT),
    ('(5) Tính minh bạch', 9.5, False, False, PP_ALIGN.LEFT),
])

# ─── MŨI TÊN NGANG: YẾU TỐ → NỘI DUNG → CHẤT LƯỢNG ────
mid_row = ROW_Y + 4.5  # y giữa hàng
add_arrow(x_yt + COL_W, mid_row, x_nd, mid_row, wpt=1.8)
add_arrow(x_nd + COL_W, mid_row, x_cl, mid_row, wpt=1.8)

# ═══════════════════════════════════════════════════════════
# 3. KẾT QUẢ VÀ TÁC ĐỘNG (dưới CHẤT LƯỢNG)
# ═══════════════════════════════════════════════════════════
KQ_Y = ROW_Y + CL_H + 0.5
KQ_H = 4.0
s_kq = add_box(x_cl, KQ_Y, COL_W, KQ_H, fill='FFFFFF', border='333333', bpt=1.5)
write_text(s_kq, [
    ('KẾT QUẢ VÀ TÁC ĐỘNG', 10.5, True),
    ('', 4),
    ('• Thúc đẩy đổi mới sáng tạo', 9, False, False, PP_ALIGN.LEFT),
    ('• Phát triển KT số bền vững', 9, False, False, PP_ALIGN.LEFT),
    ('• Bảo vệ quyền lợi', 9, False, False, PP_ALIGN.LEFT),
    ('  chủ thể sáng tạo', 9, False, False, PP_ALIGN.LEFT),
    ('• Nâng cao năng lực', 9, False, False, PP_ALIGN.LEFT),
    ('  cạnh tranh quốc gia', 9, False, False, PP_ALIGN.LEFT),
    ('• Thu hút đầu tư và CGCN', 9, False, False, PP_ALIGN.LEFT),
])

# Mũi tên CHẤT LƯỢNG → KẾT QUẢ
add_arrow(x_cl + COL_W/2, ROW_Y + CL_H, x_cl + COL_W/2, KQ_Y, wpt=1.8)

# ═══════════════════════════════════════════════════════════
# 4. PHẢN HỒI CHÍNH SÁCH (full width phía dưới)
# ═══════════════════════════════════════════════════════════
PH_Y = 15.8; PH_H = 1.6
s_ph = add_box(SL, PH_Y, SW, PH_H, fill='F5F5F5', border='333333', dashed=True)
write_text(s_ph, [
    ('PHẢN HỒI CHÍNH SÁCH', 10.5, True),
    ('', 2),
    ('Thực tiễn thực thi  →  Đánh giá, tổng kết  →  Điều chỉnh chính sách  →  Hoàn thiện QLNN về quyền SHTT',
     9, False, True, PP_ALIGN.CENTER, '444444'),
], anchor='middle')

# Mũi tên KẾT QUẢ → PHẢN HỒI
add_arrow(x_cl + COL_W/2, KQ_Y + KQ_H, x_cl + COL_W/2, PH_Y, wpt=1.5)

# ═══════════════════════════════════════════════════════════
# 5. VÒNG PHẢN HỒI: PHẢN HỒI → trái → lên → YẾU TỐ
# ═══════════════════════════════════════════════════════════
fb_x = SL - 0.1  # bên trái ngoài khung
fb_y_top = ROW_Y + YT_H/2
fb_y_bot = PH_Y + PH_H/2

# Đường ngang từ PHẢN HỒI sang trái
add_line(SL, fb_y_bot, fb_x, fb_y_bot, wpt=1.2, color='555555')
# Đường dọc lên
add_line(fb_x, fb_y_bot, fb_x, fb_y_top, wpt=1.2, color='555555')
# Mũi tên vào YẾU TỐ
add_arrow(fb_x, fb_y_top, SL, fb_y_top, wpt=1.2, color='555555')
# Nhãn vòng phản hồi
add_label(fb_x - 1.2, fb_y_top + 1.5, 1.5, 3.0,
    'Vòng\nphản\nhồi', sz=7.5, italic=True, color='666666')

# ═══════════════════════════════════════════════════════════
# NGUỒN
# ═══════════════════════════════════════════════════════════
add_label(SL, 17.8, SW, 0.5,
    'Nguồn: Tác giả tổng hợp và đề xuất',
    sz=8, italic=True, color='777777')

# ─── LƯU ─────────────────────────────────────────────────
out = r"g:\My Drive\Luan an uyen 26\LuanAn_TS_SHTT_KinhTeSo\12_Thamkhaoluanan\KhungPhanTich_HoiDong_v2.pptx"
prs.save(out)
print(f'[OK] {out}')
print()
print('=== BỐ CỤC ===')
print(' Tiêu đề: KHUNG PHÂN TÍCH QLNN VỀ QUYỀN SHTT TRONG NỀN KTS')
print(' [LÝ THUYẾT NỀN TẢNG]  ← full width')
print('      ↓')
print(' [YẾU TỐ ẢNH HƯỞNG] → [NỘI DUNG QLNN] → [CHẤT LƯỢNG QLNN]')
print('  • Bên trong (4)       4 nội dung         5 tiêu chí')
print('  • Bên ngoài (4)            ↓')
print('      ↑               [KẾT QUẢ VÀ TÁC ĐỘNG]')
print('      |                       ↓')
print('      └───── [PHẢN HỒI CHÍNH SÁCH] ──────┘')
