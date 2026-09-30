from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_AUTO_SIZE
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls

prs = Presentation()
prs.slide_width = Inches(11.69)
prs.slide_height = Inches(8.27)
slide = prs.slides.add_slide(prs.slide_layouts[6])

def add_box(slide, text, left, top, width, height, font_size=10, bold_first_line=True):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(255, 255, 255)
    shape.line.color.rgb = RGBColor(0, 0, 0)
    shape.line.width = Pt(1.5)
    
    tf = shape.text_frame
    tf.text = text
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    
    for i, p in enumerate(tf.paragraphs):
        if "•" in p.text or p.text.startswith("1.") or p.text.startswith("2.") or p.text.startswith("3.") or p.text.startswith("4."):
            p.alignment = PP_ALIGN.LEFT
        else:
            p.alignment = PP_ALIGN.CENTER
            
        for run in p.runs:
            run.font.name = 'Times New Roman'
            run.font.size = Pt(font_size)
            run.font.color.rgb = RGBColor(0, 0, 0)
            
            if bold_first_line and (i == 0 or (i == 1 and p.text.isupper())):
                run.font.bold = True
    return shape

def add_line_arrow(slide, begin_x, begin_y, end_x, end_y, dashed=False, arrow_end=True, arrow_start=False):
    connector = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, begin_x, begin_y, end_x, end_y)
    connector.line.color.rgb = RGBColor(0,0,0)
    connector.line.width = Pt(1.5)
    if dashed:
        connector.line.dash_style = 4
        
    ln = connector.element.spPr.ln
    if ln is not None:
        if arrow_end:
            tailEnd = parse_xml(r'<a:tailEnd type="triangle" w="med" len="med" %s/>' % nsdecls('a'))
            ln.append(tailEnd)
        if arrow_start:
            headEnd = parse_xml(r'<a:headEnd type="triangle" w="med" len="med" %s/>' % nsdecls('a'))
            ln.append(headEnd)
    return connector

# Title
t_title = slide.shapes.add_textbox(Inches(0.5), Inches(0.05), Inches(10.69), Inches(0.4))
t_title.text_frame.text = "KHUNG LÝ THUYẾT NGHIÊN CỨU"
t_title.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
t_title.text_frame.paragraphs[0].runs[0].font.name = 'Times New Roman'
t_title.text_frame.paragraphs[0].runs[0].font.size = Pt(16)
t_title.text_frame.paragraphs[0].runs[0].font.bold = True


# --- LÝ THUYẾT NỀN TẢNG ---
t1 = "LÝ THUYẾT NỀN TẢNG\nLý thuyết quản lý công  •  Lý thuyết thể chế  •  Lý thuyết tăng trưởng nội sinh  •  Lý thuyết về quyền sở hữu trí tuệ"
add_box(slide, t1, Inches(1.5), Inches(0.4), Inches(8.69), Inches(0.5), font_size=10)
add_line_arrow(slide, Inches(6.0), Inches(0.9), Inches(6.0), Inches(1.3))


# --- BIẾN ĐỘC LẬP (Left) ---
t2 = "BIẾN ĐỘC LẬP\n(CÁC YẾU TỐ ĐẦU VÀO)\n\n• Nhóm thể chế: Pháp luật, Chính sách\n• Nhóm nguồn lực: Nhân lực, Tài chính, Hạ tầng\n• Nhóm công nghệ: Dữ liệu số, Nền tảng số, AI"
add_box(slide, t2, Inches(0.2), Inches(1.8), Inches(3.4), Inches(1.5), font_size=10)
# Mũi tên từ Độc lập sang Trung tâm
add_line_arrow(slide, Inches(3.6), Inches(2.55), Inches(4.0), Inches(2.55))


# --- BIẾN TRUNG TÂM (Center) ---
t3 = "BIẾN TRUNG TÂM\n(QUẢN LÝ NHÀ NƯỚC VỀ QUYỀN SHTT)\n\n1. Xây dựng chiến lược, chính sách, pháp luật\n2. Tổ chức thực hiện\n3. Quản lý hoạt động xác lập, bảo vệ và hỗ trợ khai thác, thương mại hóa TSTT\n4. Thanh tra, kiểm tra, giám sát, thực thi quyền"
add_box(slide, t3, Inches(4.0), Inches(1.3), Inches(4.0), Inches(2.1), font_size=10)
add_line_arrow(slide, Inches(6.0), Inches(3.4), Inches(6.0), Inches(3.9))


# --- BIẾN PHỤ THUỘC (Center) ---
t4 = "BIẾN PHỤ THUỘC\n(SỰ GIA TĂNG TÀI SẢN TRÍ TUỆ)\n\n• Tăng trưởng số lượng văn bằng bảo hộ\n• Tăng tỷ lệ thương mại hóa tài sản trí tuệ\n• Gia tăng giá trị tài sản vô hình của doanh nghiệp"
add_box(slide, t4, Inches(4.2), Inches(3.9), Inches(3.6), Inches(1.3), font_size=10)
add_line_arrow(slide, Inches(6.0), Inches(5.2), Inches(6.0), Inches(5.7))


# --- KẾT QUẢ/TÁC ĐỘNG (Center) ---
t5 = "KẾT QUẢ VÀ TÁC ĐỘNG\n(PHÁT TRIỂN KINH TẾ SỐ)\n\n• Thúc đẩy đổi mới sáng tạo toàn diện\n• Tăng trưởng kinh tế dựa trên tri thức\n• Nâng cao năng lực cạnh tranh quốc gia"
add_box(slide, t5, Inches(4.1), Inches(5.7), Inches(3.8), Inches(1.2), font_size=10)


# --- BIẾN ĐIỀU TIẾT (Right) ---
tdt = "BIẾN ĐIỀU TIẾT\n(CÁC YẾU TỐ MÔI TRƯỜNG VĨ MÔ)\n\n• Trình độ công nghệ (AI, Big Data)\n• Mức độ hội nhập quốc tế\n• Nhận thức của doanh nghiệp và xã hội\n• Cạnh tranh thị trường"
add_box(slide, tdt, Inches(8.5), Inches(2.3), Inches(2.9), Inches(1.8), font_size=10)

# Mũi tên điều tiết nét đứt chỉ vào mũi tên Trung tâm -> Phụ thuộc (Y=3.65)
add_line_arrow(slide, Inches(8.5), Inches(3.65), Inches(6.0), Inches(3.65), dashed=True)


# --- FEEDBACK LOOP ---
tfb = "PHẢN HỒI CHÍNH SÁCH\nThực tiễn phát triển kinh tế ➔ Đánh giá ➔ Điều chỉnh chính sách, thể chế ➔ Hoàn thiện QLNN"
add_box(slide, tfb, Inches(1.5), Inches(7.3), Inches(8.69), Inches(0.5), font_size=10)

# Từ Kết quả xuống Feedback
add_line_arrow(slide, Inches(6.0), Inches(6.9), Inches(6.0), Inches(7.3), dashed=True)

# Từ khối Feedback vòng lên Biến Độc lập (Đầu vào) để làm mới chu trình
add_line_arrow(slide, Inches(1.5), Inches(7.55), Inches(0.15), Inches(7.55), dashed=True, arrow_end=False)
add_line_arrow(slide, Inches(0.15), Inches(7.55), Inches(0.15), Inches(1.1), dashed=True, arrow_end=False)
add_line_arrow(slide, Inches(0.15), Inches(1.1), Inches(1.9), Inches(1.1), dashed=True, arrow_end=False)
add_line_arrow(slide, Inches(1.9), Inches(1.1), Inches(1.9), Inches(1.8), dashed=True)

prs.save("KhungLyThuyet_KinhTe_BienTrungTam.pptx")
print("Saved PPTX with Quản lý nhà nước as Biến trung tâm.")
