import os
import pandas as pd
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

BASE_DIR = r"g:\My Drive\Luan an uyen 26\LuanAn_TS_SHTT_KinhTeSo\07_Detailed_Outline"
os.makedirs(BASE_DIR, exist_ok=True)

# ==========================================
# 1. GENERATE EXCEL MATRICES
# ==========================================

# 1. PRODUCT EXTRACTION MAP
product_columns = ["Sản phẩm/Mục đích", "Nội dung trích xuất từ luận án", "Chương/Mục tương ứng", "Yêu cầu chỉnh sửa/phát triển"]
product_data = [
    ["Chuyên đề Tổng quan", "Tổng hợp lý thuyết nền tảng, kinh nghiệm quốc tế và tình hình nghiên cứu SHTT.", "Chương 1, 2", ""],
    ["Chuyên đề nghiên cứu 1", "Khung pháp lý về QLNN đối với SHTT trong nền kinh tế số.", "Chương 2", ""],
    ["Chuyên đề nghiên cứu 2", "Đánh giá thực trạng QLNN về SHTT tại VN (Tập trung vào vi phạm trên Internet).", "Chương 3", ""],
    ["Chuyên đề nghiên cứu 3", "Giải pháp nâng cao hiệu quả QLNN (Mô hình quản lý mới, ứng dụng AI).", "Chương 4", ""],
    ["Bài báo trong nước 1", "Phân tích khoảng trống pháp luật về xử lý vi phạm bản quyền trên nền tảng xuyên biên giới.", "Tiểu mục 2.2, 3.1", ""],
    ["Bài báo trong nước 2", "Đánh giá thực trạng ứng dụng công nghệ (AI, Big Data) trong QLNN về SHTT.", "Tiểu mục 3.3", ""],
    ["Bài báo trong nước 3", "Chính sách khuyến khích và bảo vệ đổi mới sáng tạo trong doanh nghiệp số.", "Tiểu mục 3.4, 4.2", ""],
    ["Bài báo quốc tế", "Nghiên cứu định lượng các yếu tố ảnh hưởng đến hiệu lực thực thi SHTT trong kinh tế số tại nước đang phát triển: Trường hợp Việt Nam.", "Chương 2, 3", "Cần dịch sang tiếng Anh, format theo chuẩn ISI/Scopus"],
    ["Đề tài cấp cơ sở", "Giải pháp chuyển đổi số trong công tác quản lý SHTT tại địa phương/Bộ ban ngành cụ thể.", "Chương 3, 4", "Cần giới hạn phạm vi địa bàn hẹp hơn"]
]
pd.DataFrame(product_data, columns=product_columns).to_excel(os.path.join(BASE_DIR, "PRODUCT_EXTRACTION_MAP.xlsx"), index=False)

# 2. NOVELTY CLAIM MATRIX
novelty_columns = ["Loại đóng góp mới", "Tuyên bố cụ thể (Novelty Claim)", "Cơ sở chứng minh", "Dữ liệu/Tài liệu hỗ trợ", "Rủi ro phản biện", "Cách bảo vệ trước hội đồng"]
novelty_data = [
    ["Lý luận", "Xây dựng khung phân tích mới về QLNN đối với SHTT bằng cách tích hợp mô hình thể chế mạng (Networked Governance) thay cho mô hình quản lý tập trung truyền thống.", "Phê phán lý thuyết cũ không đáp ứng được tính phi biên giới của nền tảng số.", "Literature Review, Các mô hình quốc tế (TQ, HQ).", "Hội đồng cho rằng mô hình mạng khó áp dụng ở VN.", "Chứng minh qua số liệu thực tiễn và xu hướng toàn cầu hóa."],
    ["Thực tiễn", "Định lượng mức độ ảnh hưởng của hạ tầng công nghệ và sự phối hợp liên ngành đến hiệu quả thực thi quyền SHTT trực tuyến tại VN.", "Kết quả phân tích hồi quy/PLS-SEM lần đầu tiên được áp dụng vào chủ đề này tại VN.", "Survey 100 mẫu, Báo cáo kiểm định độ tin cậy.", "Cỡ mẫu nhỏ (100) có thể bị bắt bẻ.", "Khẳng định đây là nghiên cứu tiên phong (Exploratory) và có sự bù đắp từ phỏng vấn sâu chuyên gia (Triangulation)."],
    ["Chính sách/QLNN", "Đề xuất bộ cơ chế tự động hóa kết hợp AI và quy định chi tiết trách nhiệm nhà cung cấp dịch vụ trung gian.", "Dựa trên bất cập thực tế (Chương 3) và học hỏi chuẩn mực quốc tế.", "Interview Transcripts, Thống kê vi phạm kéo dài.", "Luật SHTT 2022 đã có quy định tương tự.", "Chỉ rõ Luật 2022 mới chỉ nêu khung, còn thiếu quy trình thực thi kỹ thuật và cơ chế giám sát."]
]
pd.DataFrame(novelty_data, columns=novelty_columns).to_excel(os.path.join(BASE_DIR, "NOVELTY_CLAIM_MATRIX.xlsx"), index=False)

# 3. OUTLINE WITH EVIDENCE PLAN
evidence_columns = ["Chương", "Mục", "Luận điểm chính", "Evidence type (Tài liệu/Pháp lý/Data)", "Nguồn cụ thể"]
pd.DataFrame(columns=evidence_columns).to_excel(os.path.join(BASE_DIR, "OUTLINE_WITH_EVIDENCE_PLAN.xlsx"), index=False)

# 4. DATA USE MAP
data_map_columns = ["Biến khảo sát (Variable Code)", "Sử dụng ở Chương/Mục", "Mục đích sử dụng (Thống kê/Kiểm định)"]
pd.DataFrame(columns=data_map_columns).to_excel(os.path.join(BASE_DIR, "DATA_USE_MAP.xlsx"), index=False)

# 5. FIGURE TABLE REGISTER
fig_columns = ["Mã số", "Loại (Hình/Bảng)", "Tên", "Nguồn", "Vị trí trong luận án"]
pd.DataFrame(columns=fig_columns).to_excel(os.path.join(BASE_DIR, "FIGURE_TABLE_REGISTER.xlsx"), index=False)

# 6. LEGAL ARGUMENT MATRIX
legal_columns = ["Luận điểm trong luận án", "Văn bản pháp luật liên quan", "Điều khoản cụ thể", "Ghi chú/Đánh giá"]
pd.DataFrame(columns=legal_columns).to_excel(os.path.join(BASE_DIR, "LEGAL_ARGUMENT_MATRIX.xlsx"), index=False)

# 7. THESIS COHERENCE MAP
coherence_columns = ["Chương 1 (Vấn đề nghiên cứu)", "Chương 2 (Khung lý thuyết giải quyết)", "Chương 3 (Bằng chứng thực tiễn)", "Chương 4 (Giải pháp tương ứng)", "Kết luận Logic"]
pd.DataFrame(columns=coherence_columns).to_excel(os.path.join(BASE_DIR, "THESIS_COHERENCE_MAP.xlsx"), index=False)

# ==========================================
# 2. GENERATE DETAILED WORD OUTLINE
# ==========================================

def add_outline_section(doc, level_text, title, objective, argument, content, refs, laws, survey, interview, tables, table_src, citation_pos, products, priority, risk, self_write):
    doc.add_heading(f"{level_text} {title}", level=2 if len(level_text.split('.'))==2 else 3)
    p = doc.add_paragraph()
    p.add_run("1. Tên mục: ").bold = True; p.add_run(title + "\n")
    p.add_run("2. Mục tiêu: ").bold = True; p.add_run(objective + "\n")
    p.add_run("3. Luận điểm chính: ").bold = True; p.add_run(argument + "\n")
    p.add_run("4. Nội dung cần viết: ").bold = True; p.add_run(content + "\n")
    p.add_run("5. Nguồn tài liệu tham khảo dự kiến: ").bold = True; p.add_run(refs + "\n")
    p.add_run("6. Văn bản pháp quy dự kiến: ").bold = True; p.add_run(laws + "\n")
    p.add_run("7. Dữ liệu khảo sát dự kiến: ").bold = True; p.add_run(survey + "\n")
    p.add_run("8. Dữ liệu phỏng vấn dự kiến: ").bold = True; p.add_run(interview + "\n")
    p.add_run("9. Bảng/hình/biểu đồ dự kiến: ").bold = True; p.add_run(tables + "\n")
    p.add_run("10. Nguồn của bảng/hình: ").bold = True; p.add_run(table_src + "\n")
    p.add_run("11. Vị trí trích dẫn: ").bold = True; p.add_run(citation_pos + "\n")
    p.add_run("12. Sản phẩm liên quan: ").bold = True; p.add_run(products + "\n")
    p.add_run("13. Mức độ ưu tiên: ").bold = True; p.add_run(priority + "\n")
    p.add_run("14. Rủi ro thiếu nguồn/dữ liệu: ").bold = True; p.add_run(risk + "\n")
    p.add_run("15. Nội dung học viên tự viết/xác nhận: ").bold = True; p.add_run(self_write + "\n")
    doc.add_paragraph("-" * 60)

doc = Document()
doc.add_heading('BẢN THIẾT KẾ THI CÔNG LUẬN ÁN (BLUEPRINT OUTLINE)', 0)

# Chương 1
doc.add_heading('CHƯƠNG 1: TỔNG QUAN TÌNH HÌNH NGHIÊN CỨU', level=1)
add_outline_section(doc, "1.1.", "Tình hình nghiên cứu trên thế giới", 
                    "Chỉ ra dòng chảy lý thuyết quốc tế về SHTT trong kỷ nguyên số.", 
                    "Thế giới đang chuyển từ quản lý cứng nhắc sang quản trị mạng lưới và áp dụng công nghệ tự động.", 
                    "- Điểm lại các nghiên cứu về bản quyền trên Internet.\n- Các nghiên cứu về nền tảng xuyên biên giới.",
                    "LITERATURE_SCREENING_MATRIX, NotebookLM Extracts.", "Các Hiệp ước của WIPO", "Không", "Không", 
                    "Bảng 1.1: Tổng hợp các khoảng trống nghiên cứu thế giới", "Tác giả tổng hợp", "Trải đều trong mục", 
                    "Chuyên đề Tổng quan", "Cao", "Thiếu tài liệu cập nhật (2024-2025)", "Phần tóm tắt và bình luận đánh giá cá nhân.")

# Chương 2
doc.add_heading('CHƯƠNG 2: CƠ SỞ LÝ LUẬN VÀ PHÁP LÝ', level=1)
add_outline_section(doc, "2.1.", "Khái niệm và đặc điểm của quyền SHTT trong kinh tế số", 
                    "Xây dựng định nghĩa làm nền tảng cho toàn bộ luận án.", 
                    "SHTT trong kinh tế số mang tính vô hình cao, dễ sao chép, lan truyền xuyên biên giới.", 
                    "- Định nghĩa kinh tế số.\n- Phân tích sự thay đổi bản chất của tài sản trí tuệ.",
                    "Giáo trình SHTT, WIPO Guidelines", "Luật SHTT 2022", "Không", "Không", 
                    "Hình 2.1: Sơ đồ vòng đời tài sản trí tuệ số", "Tác giả xây dựng", "Đầu mục 2.1", 
                    "Bài báo trong nước 1, Chuyên đề 1", "Cao", "Bị trùng lặp lý thuyết sách giáo khoa", "Xây dựng định nghĩa mới mang dấu ấn cá nhân.")

# Chương 3
doc.add_heading('CHƯƠNG 3: THỰC TRẠNG QUẢN LÝ NHÀ NƯỚC VỀ SHTT', level=1)
add_outline_section(doc, "3.1.", "Thực trạng khung pháp lý và chính sách", 
                    "Chứng minh sự bất cập của pháp luật hiện hành qua số liệu.", 
                    "Pháp luật Việt Nam đã có cải thiện nhưng vẫn đi sau thực tiễn công nghệ.", 
                    "- Phân tích các điều khoản Luật SHTT 2022.\n- Phân tích số liệu vi phạm không bị xử lý.",
                    "Báo cáo của Thanh tra Bộ KHCN, Cục SHTT", "Nghị định xử phạt vi phạm hành chính SHTT", 
                    "Biến LF1, LF2, LF3", "Mã CG01, CG02", "Biểu đồ 3.1: Số lượng vụ vi phạm bị xử lý 2020-2025", "Bộ KHCN & Data Analysis Report", "Sau khi nêu luật", 
                    "Bài báo quốc tế, Chuyên đề 2", "Rất cao", "Khó xin số liệu xử phạt từ cơ quan công an", "Nhận định về nguyên nhân của sự bất cập.")

# Chương 4
doc.add_heading('CHƯƠNG 4: GIẢI PHÁP VÀ KIẾN NGHỊ', level=1)
add_outline_section(doc, "4.1.", "Mô hình quản lý nhà nước ứng dụng công nghệ số", 
                    "Đề xuất mô hình đột phá mang tính đóng góp mới.", 
                    "Phải áp dụng AI và liên thông cơ sở dữ liệu quốc gia để xử lý vi phạm tức thời.", 
                    "- Kiến trúc hệ thống liên thông.\n- Đề xuất thay đổi quy trình hành chính.",
                    "Kinh nghiệm HQ, TQ từ Chương 1", "Dự thảo Luật dữ liệu (nếu có)", "Không", "Mã CG03, CG05", 
                    "Hình 4.1: Mô hình mạng lưới quản lý SHTT ứng dụng AI", "Tác giả đề xuất", "Cuối mục 4.1", 
                    "Bài báo trong nước 3, Đề tài cấp cơ sở", "Rất cao", "Mô hình thiếu tính khả thi kỹ thuật", "Lập luận bảo vệ tính khả thi của mô hình.")

doc.save(os.path.join(BASE_DIR, "DETAILED_THESIS_OUTLINE.docx"))
print("Detailed thesis outline and matrices created successfully.")
