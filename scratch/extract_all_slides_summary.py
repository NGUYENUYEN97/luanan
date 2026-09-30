import os
import sys
import pypdf
import pptx

sys.stdout.reconfigure(encoding='utf-8')

wipo_dir = r"g:\My Drive\Luan an uyen 26\LuanAn_TS_SHTT_KinhTeSo\bài báo sáng chế\WIPO 26-27.8.2026-20260826T025931Z-1-001\WIPO 26-27.8.2026"

out_dir = r"C:\Users\Admin\.gemini\antigravity-ide\brain\2664ec95-b734-4052-ac21-ea79e5aed358\scratch"

# Parse Slide Trinh Bay Ve CNCL.pdf
print("=== SLIDE TRINH BAY VE CNCL.pdf ===")
pdf_path = os.path.join(wipo_dir, "SLIDE TRINH BAY VE CNCL.pdf")
reader = pypdf.PdfReader(pdf_path)
for i, page in enumerate(reader.pages):
    print(f"--- Page {i+1} ---")
    print(page.extract_text())

# Parse Viettel.pptx
print("\n=== AUG 27 AM - VIETTEL.PPTX ===")
ppt_path = os.path.join(wipo_dir, "Aug 27 AM - Viettel.pptx")
prs = pptx.Presentation(ppt_path)
for i, slide in enumerate(prs.slides):
    txts = [s.text.strip() for s in slide.shapes if hasattr(s, "text") and s.text and s.text.strip()]
    full = "\n".join(txts)
    if any(k in full.lower() for k in ['viettel', 'patent', 'analytics', 'portfolio', 'r&d', 'strategy']):
        print(f"--- Slide {i+1} ---")
        print(full[:400])

# Parse IPVN.pptx
print("\n=== AUG 27 PM - IPVN.PPTX ===")
ppt_path = os.path.join(wipo_dir, "Aug 27 PM - IPVN.pptx")
prs = pptx.Presentation(ppt_path)
for i, slide in enumerate(prs.slides):
    txts = [s.text.strip() for s in slide.shapes if hasattr(s, "text") and s.text and s.text.strip()]
    full = "\n".join(txts)
    if any(k in full.lower() for k in ['ip viet nam', 'ipvn', 'recommendation', 'next step', 'training', 'capability', 'policy']):
        print(f"--- Slide {i+1} ---")
        print(full[:400])
