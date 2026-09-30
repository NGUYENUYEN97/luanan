import os
import sys
import docx
import pypdf
import pptx

sys.stdout.reconfigure(encoding='utf-8')

wipo_dir = r"g:\My Drive\Luan an uyen 26\LuanAn_TS_SHTT_KinhTeSo\bài báo sáng chế\WIPO 26-27.8.2026-20260826T025931Z-1-001\WIPO 26-27.8.2026"
out_dir = r"C:\Users\Admin\.gemini\antigravity-ide\brain\2664ec95-b734-4052-ac21-ea79e5aed358\scratch"
os.makedirs(out_dir, exist_ok=True)

out_file = os.path.join(out_dir, "wipo_materials_summary.txt")

with open(out_file, "w", encoding="utf-8") as f:
    f.write("=== SUMMARY OF WIPO WORKSHOP MATERIALS (26-27/08/2026) ===\n\n")

    # 1. Agenda DOCX
    f.write("--- 1. Agenda DOCX ---\n")
    try:
        doc = docx.Document(os.path.join(wipo_dir, "Chương trình làm việc_FN.docx"))
        for p in doc.paragraphs:
            if p.text.strip():
                f.write(p.text + "\n")
        for table in doc.tables:
            for row in table.rows:
                f.write(" | ".join([c.text.replace("\n", " ").strip() for c in row.cells]) + "\n")
    except Exception as e:
        f.write(f"Error reading agenda: {e}\n")
    f.write("\n" + "="*50 + "\n\n")

    # 2. PDF Files
    pdf_files = [
        "Aug 26 AM - Intro_overview of PA.pdf",
        "Aug 26 PM - MST preliminary analysis VN strategic technologies.pdf",
        "SLIDE TRINH BAY VE CNCL.pdf"
    ]
    for pdf_name in pdf_files:
        f.write(f"--- PDF: {pdf_name} ---\n")
        pdf_path = os.path.join(wipo_dir, pdf_name)
        if os.path.exists(pdf_path):
            try:
                reader = pypdf.PdfReader(pdf_path)
                f.write(f"Total pages: {len(reader.pages)}\n")
                for idx, page in enumerate(reader.pages):
                    text = page.extract_text()
                    if text and text.strip():
                        f.write(f"[Page {idx+1}]\n{text.strip()}\n")
            except Exception as e:
                f.write(f"Error reading {pdf_name}: {e}\n")
        else:
            f.write(f"File not found: {pdf_path}\n")
        f.write("\n" + "="*50 + "\n\n")

    # 3. PPTX Files
    pptx_files = [
        "Aug 27 AM - Viettel.pptx",
        "Aug 27 PM - IPVN.pptx",
        "Banner_Ban_do_SC_26_8_2026.pptx"
    ]
    for ppt_name in pptx_files:
        f.write(f"--- PPTX: {ppt_name} ---\n")
        ppt_path = os.path.join(wipo_dir, ppt_name)
        if os.path.exists(ppt_path):
            try:
                prs = pptx.Presentation(ppt_path)
                f.write(f"Total slides: {len(prs.slides)}\n")
                for idx, slide in enumerate(prs.slides):
                    slide_text = []
                    for shape in slide.shapes:
                        if hasattr(shape, "text") and shape.text and shape.text.strip():
                            slide_text.append(shape.text.strip())
                    if slide_text:
                        f.write(f"[Slide {idx+1}]\n" + "\n".join(slide_text) + "\n")
            except Exception as e:
                f.write(f"Error reading {ppt_name}: {e}\n")
        else:
            f.write(f"File not found: {ppt_path}\n")
        f.write("\n" + "="*50 + "\n\n")

    # 4. Scopus article DOCX
    scopus_docx = r"g:\My Drive\Luan an uyen 26\LuanAn_TS_SHTT_KinhTeSo\bài báo sáng chế\Xu hướng sáng chế trong nền kinh tế số Dữ liệu từ scopus.docx"
    f.write("--- Scopus DOCX ---\n")
    if os.path.exists(scopus_docx):
        try:
            doc = docx.Document(scopus_docx)
            for p in doc.paragraphs:
                if p.text.strip():
                    f.write(p.text + "\n")
        except Exception as e:
            f.write(f"Error reading Scopus docx: {e}\n")
    f.write("\n=== END OF SUMMARY ===\n")

print(f"Written summary to {out_file}")
