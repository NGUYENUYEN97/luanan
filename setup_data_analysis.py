import os
import pandas as pd
from docx import Document

# 1. Define paths
BASE_DIR = r"g:\My Drive\Luan an uyen 26\LuanAn_TS_SHTT_KinhTeSo\06_Data_Analysis"
os.makedirs(BASE_DIR, exist_ok=True)

# 2. Markdown Templates
md_files = {
    "DATA_CLEANING_LOG.md": """# DATA CLEANING LOG
Ghi lại toàn bộ các bước làm sạch dữ liệu khảo sát và phỏng vấn.

| Ngày | Dataset | Thao tác | Lý do | Người thực hiện |
|---|---|---|---|---|
| ... | ... | ... | ... | ... |
""",
    "EXPERT_QUOTE_BANK.md": """# EXPERT QUOTE BANK
Kho lưu trữ trích dẫn chuyên gia đã được phân loại (chỉ dùng mã ẩn danh).

## Nhóm Khung pháp lý (LF)
- **CG01**: "[Trích dẫn mẫu chờ nạp]"
"""
}

for filename, content in md_files.items():
    with open(os.path.join(BASE_DIR, filename), "w", encoding="utf-8") as f:
        f.write(content)

# 3. Excel Templates
excel_templates = {
    "DESCRIPTIVE_STATISTICS.xlsx": ["Biến (Variable)", "N", "Mean", "Std.Dev", "Min", "Max", "Skewness", "Kurtosis"],
    "RELIABILITY_ANALYSIS.xlsx": ["Nhóm nhân tố", "Số lượng biến", "Cronbach's Alpha", "Item-Total Correlation min", "Ghi chú"],
    "FACTOR_OR_SCALE_ANALYSIS.xlsx": ["Phương pháp (EFA/CFA)", "KMO", "Bartlett's Test (p-value)", "Phương sai trích (%)", "Factor Loadings", "Điều kiện: N > 150"],
    "HYPOTHESIS_TESTING_REPORT.xlsx": ["Giả thuyết", "Đường dẫn (Path)", "Beta", "T-value", "P-value", "Kết luận (Chấp nhận/Bác bỏ)"],
    "QUALITATIVE_CODING_MATRIX.xlsx": ["Mã chuyên gia", "Nhóm chủ đề", "Node/Theme", "Tần suất xuất hiện", "Trích dẫn tiêu biểu"],
    "MIXED_METHODS_TRIANGULATION_MATRIX.xlsx": ["Câu hỏi nghiên cứu", "Kết quả Định lượng (Khảo sát)", "Kết quả Định tính (Phỏng vấn)", "Đối chiếu (Hỗ trợ/Tương phản)", "Diễn giải"],
    "FIGURE_TABLE_DATA_SOURCE_MAP.xlsx": ["Mã Hình/Bảng", "Tên Hình/Bảng", "File nguồn dữ liệu", "Biến sử dụng", "Vị trí trong luận án (Chương/Mục)"]
}

for filename, columns in excel_templates.items():
    df = pd.DataFrame(columns=columns)
    df.to_excel(os.path.join(BASE_DIR, filename), index=False)

# 4. Word Document Template
def create_analysis_report_docx():
    doc = Document()
    doc.add_heading('BÁO CÁO PHÂN TÍCH DỮ LIỆU LUẬN ÁN', 0)
    doc.add_paragraph('Lưu ý: Tài liệu này là bản ghi nhận kết quả phân tích dữ liệu nhằm hỗ trợ viết luận án.')
    
    sections = [
        "1. Mô tả mẫu",
        "2. Kết quả định lượng (Descriptive, Reliability, EFA, Hypothesis Testing)",
        "3. Kết quả định tính (Coding, Themes)",
        "4. Đối chiếu định lượng – định tính (Triangulation)",
        "5. Kết quả dùng cho Chương 3 (Thực trạng)",
        "6. Kết quả dùng cho Chương 4 (Giải pháp)",
        "7. Kết quả dùng cho bài báo khoa học",
        "8. Kết quả dùng cho đề tài cấp cơ sở",
        "9. Giới hạn dữ liệu"
    ]
    
    for sec in sections:
        doc.add_heading(sec, level=1)
        if "Giới hạn dữ liệu" in sec:
            doc.add_paragraph('Ghi chú: Chỉ chạy EFA/SEM nếu cỡ mẫu N > 5 lần số biến quan sát hoặc N > 150. Nếu nhỏ hơn, sẽ đề xuất dùng PLS-SEM hoặc chỉ dừng ở đánh giá độ tin cậy và thống kê mô tả.')
        else:
            doc.add_paragraph('[Chờ nạp dữ liệu để cập nhật phần này...]')
            
    doc.save(os.path.join(BASE_DIR, "DATA_ANALYSIS_REPORT.docx"))

create_analysis_report_docx()

print("Data Analysis framework and templates created successfully.")
