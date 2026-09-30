# Dựng bảng phân loại tài liệu trong 01_NotebookLM_Inputs cho chuyên đề tổng quan.
# Đầu vào: keep.json (danh sách tệp đã loại trùng, kèm văn bản 3 trang đầu) do bước trích xuất tạo ra.
# Chạy: python build_phan_loai_tong_quan.py <đường dẫn keep.json>
import json
import os
import sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from tong_quan_bo_sung_data import DE_XUAT, CAN_XAC_MINH  # noqa: E402

docs = json.load(open(sys.argv[1], encoding="utf-8"))
PREFIX = "01_NotebookLM_Inputs/"

# ------------------------------------------------------------------ tài liệu đã trích dẫn (đối chiếu thủ công)
DA_TRICH = {
    111: "neves2021", 118: "foster2024", 119: "foster2024", 267: "foster2024", 334: "foster2024", 517: "foster2024",
    159: "fasi2022", 362: "fasi2022", 164: "teece1986", 451: "teece1986", 177: "wipo2026", 182: "cucSHTT2025",
    209: "sinnreich2019", 219: "fisher2001", 198: "nguyenToUyen2023", 269: "nguyenVinhHung2020", 280: "buiXuanChung2025",
    282: "tranVanHai2022", 285: "lamXuanHung2025", 288: "nguyenChiHai2021", 289: "vinhHungNguyen2024", 296: "dangVietDuc2022",
    306: "cameron2019", 481: "cameron2019", 311: "hoThucTai2019", 318: "hoThucTai2019", 312: "hoangVanThuy2024",
    314: "tranMyLinh2025", 335: "luuHoangLong2025b", 336: "nguyenQuocThinh2020", 337: "dangThanhDat2026",
    338: "phamNgocPha2025", 339: "leThuHa2022", 344: "leThuHa2022", 341: "longTranViet2023", 342: "leHa2022",
    345: "nguyenHuuCan2020", 346: "nguyenThanhTung2026", 348: "nguyenHuuXuyen2018", 349: "doanDucLuong2022",
    350: "doanDucLuong2009", 352: "phamHongQuat2021", 353: "leThanhTam2025", 354: "nguyenVanKim2015",
    355: "nguyenMinhHuyenTrang2025", 356: "phamTrungLuan2025", 357: "nguyenHuuCan2017", 358: "nguyenVanPhuc2025",
    359: "nguyenHuuXuyen2024", 360: "hoangLanPhuong2019", 361: "luuHoangLong2025a",
    367: "tranThanhHuyen2021", 368: "tranThanhHuyen2021", 369: "tranThanhHuyen2021", 370: "tranThanhHuyen2021",
    371: "tranThanhHuyen2021", 377: "tranThanhHuyen2021",
    372: "nguyenVanDong2025", 373: "nguyenVanDong2025", 374: "nguyenVanDong2025", 375: "nguyenVanDong2025", 376: "nguyenVanDong2025",
    380: "hoThuyNgoc2025", 381: "hoThuyNgoc2025", 382: "hoThuyNgoc2025", 383: "hoThuyNgoc2025",
    384: "vuPhuongGiang2025", 385: "vuPhuongGiang2025", 386: "vuPhuongGiang2025", 396: "vuPhuongGiang2025", 397: "vuPhuongGiang2025",
    387: "leBichThuy2021", 388: "leBichThuy2021", 389: "leBichThuy2021", 390: "leBichThuy2021", 391: "leBichThuy2021",
    452: "wipo2024", 453: "mazzucato2013", 456: "lundvall1992", 458: "chesbrough2003", 459: "fink2005",
    460: "arora2001", 463: "maskus2000", 474: "unctad2019", 500: "williams2013 (bản NBER WP)",
    512: "acemoglu2012 (bản working paper 2011)", 523: "goldfarb2019 (bản NBER WP)", 531: "goldfarb2019",
    538: "goldfarb2019 (bản thảo 2017)", 528: "bukht2017", 560: "oecd2014",
}

# Bản trùng nội dung (khác tệp) của tài liệu đề xuất
TRUNG_DE_XUAT = {41: 264, 60: 264, 94: 264, 515: 264, 169: 114, 555: 74, 402: 401, 403: 401,
                 398: 401, 399: 401, 400: 401, 378: 379, 405: 404, 406: 404, 407: 404, 408: 404, 409: 404,
                 395: 394, 410: 394, 411: 394, 412: 394, 413: 394, 506: 499, 510: 499}

# ------------------------------------------------------------------ tham khảo mức 2 (đúng chủ đề, giá trị bổ sung thấp hơn)
THAM_KHAO = {
    1: ("1.1.2.2", "Tuyển tập ĐHCN - ĐHQGHN (12/2020) về kinh tế số, AI; tài liệu xám"),
    98: ("1.1.2.1", "Liu & Lin (2025), IREF 102, 104316: trao quyền SHTT, kinh tế số và lực lượng sản xuất mới"),
    100: ("1.1.2.1", "Bản trùng của Liu & Lin (2025)"),
    138: ("1.1.3.1", "Henry (2011), Technology Innovation Management Review: quyền SHTT như một loại tài sản"),
    161: ("1.1.3.1", "De Leon & Fernandez Donoso (2017), Springer: chiến lược thương mại hóa SHTT (chương 3)"),
    163: ("1.1.3.1", "Business Horizons 47(1) 2004: quản trị SHTT theo chuỗi cung ứng tri thức"),
    135: ("1.1.3.1", "Van Norman & Eisenkot (2017): quy trình thương mại hóa trong lĩnh vực y sinh (ngoài trọng tâm)"),
    99: ("1.1.3.1", "Expert Panel on IP, Ontario (2020): SHTT trong hệ sinh thái đổi mới cấp vùng"),
    105: ("1.1.2.1", "He, He, Hou, PLOS ONE: kinh tế số, đổi mới công nghệ và phát triển bền vững"),
    115: ("1.1.1.1", "WIPO (2022), World IP Report 2022: The Direction of Innovation"),
    121: ("1.1.3.1", "Klaila & Hall (2000), Journal of Intellectual Capital: tài sản trí tuệ là chiến lược"),
    122: ("1.1.3.1", "Morricone et al. (2017), Research Policy: chiến lược thương mại hóa và IPO"),
    123: ("1.1.3.1", "Madhusoodanan et al. (2022), J. World Intellectual Property: quyền SHTT và MSME"),
    125: ("1.1.3.1", "Lohrey & Willoughby (2025), World Patent Information: tổng quan SHTT và SME"),
    133: ("1.1.3.1", "Mooi & Wuyts (2021), IJRM: giá trị từ li-xăng công nghệ"),
    160: ("1.1.3.1", "Ziegler et al. (2013), J. Technology Transfer: thương mại hóa SHTT ra bên ngoài"),
    162: ("1.1.3.1", "Wang & Dong (2025), Elsevier: bảo hộ SHTT và vốn hóa dữ liệu"),
    166: ("1.1.3.1", "Holgersson, Granstrand & Bogers (2018), Long Range Planning: chiến lược SHTT trong hệ sinh thái"),
    167: ("1.1.2.1", "Perel, Somech & Elkin-Koren (2026), Big Data & Society: nền tảng số khép kín dữ liệu"),
    168: ("1.1.3.1", "Sun et al. (2024), Information (MDPI): trắc lượng thư mục về SHTT dữ liệu"),
    16: ("1.1.3.1", "IVSC (2025): tài sản vô hình trên thị trường vốn ASEAN"),
    39: ("1.1.3.1", "WIPO (2023): Country Perspectives - China's Journey, tài chính dựa trên SHTT"),
    45: ("Kinh nghiệm", "Huang & Cao (2023), APJIE: cải cách hệ thống quyền SHTT của Trung Quốc"),
    77: ("1.1.1.1", "Roffe & Santa Cruz, ECLAC: quyền SHTT và phát triển bền vững"),
    84: ("1.1.3.1", "Ciuriak (2021), SSRN/CIGI: SHTT và kinh tế số, năm vấn đề về chuẩn mực quốc tế"),
    93: ("1.1.1.1", "WIPO: Methodology for the Development of National IP Strategies (2nd ed.)"),
    66: ("1.1.1.1", "WIPO (2016): Tool 3 - Benchmarking Indicators, chiến lược SHTT quốc gia"),
    144: ("1.1.3.1", "DesForges: The Commercial Exploitation of IPR by Licensing (báo cáo chuyên đề, nguồn tải không chính thức)"),
    139: ("1.1.3.1", "Heim: Intellectual Property Management (sách)"),
    176: ("1.1.1.2", "ICC BASCAP: Thúc đẩy và bảo hộ quyền SHTT tại Việt Nam"),
    200: ("1.1.3.1", ""),
    223: ("1.1.1.1", "Guan (2014), Springer: khái niệm, lịch sử và tranh luận về SHTT (chương 1)"),
    229: ("1.1.1.1", "Breyer (1970): The Uneasy Case for Copyright (tệp không có lớp văn bản)"),
    233: ("1.1.3.1", "Bosworth: The management of IP, chương giới thiệu (sách Elgar)"),
    266: ("1.1.2.1", "Rusche & Scheufen (2018), IW-Report 48/2018: quyền tài sản và khung pháp lý trong kinh tế số"),
    271: ("1.1.1.2", "Phùng Thị Yến và cs. (2026): quyền tác giả đối với tác phẩm do AI tạo ra"),
    272: ("1.1.2.2", "Le Thi Dan Dung, Bui Tien Hanh (2025), JFAR: vốn nhân lực trong kinh tế số"),
    297: ("1.1.2.2", "Nguyễn Phương Lý (2024), Kinh tế và Dự báo: phát triển kinh tế số tại Việt Nam"),
    299: ("1.1.2.2", "Nguyễn Huyền Trang (2025), Giáo dục và Xã hội: phát triển kinh tế số ở Việt Nam"),
    300: ("1.1.2.2", "Trương Quốc Cường, Kinh tế và Dự báo: phát triển kinh tế số và khuyến nghị chính sách"),
    301: ("1.1.2.2", "Bùi Thị Xuân: đổi mới sáng tạo trong nền kinh tế số"),
    308: ("1.1.2.2", "Nguyễn Thị Tuyết Trinh, Nguyễn Thị Hảo (2023): phát triển kinh tế số trong bối cảnh mới"),
    309: ("1.1.1.2", "Lê Hồ Trung Hiếu và cs., Tạp chí Công Thương: quyền tác giả đối với tác phẩm AI"),
    310: ("1.1.2.2", "Lê Thị Thu Hà (2025), ĐH Hải Phòng: thực trạng, giải pháp phát triển kinh tế số"),
    316: ("1.1.2.2", "Võ Xuân Vinh và cs., Tạp chí Công Thương: giải pháp phát triển kinh tế số đến 2030"),
    324: ("1.1.2.2", "Lê Duy Bình, Trần Thị Phương (2020): kinh tế số và chuyển đổi số tại Việt Nam (tài liệu EVFTA)"),
    325: ("1.1.2.2", "Nghiêm Xuân Khoát, Đặng Thị Minh Hiền (2025): đổi mới sáng tạo và công nghệ số"),
    330: ("1.1.2.2", "Lê Thu Hà: thực trạng và giải pháp phát triển kinh tế số"),
    333: ("1.1.2.2", "Bành Thị Hồng Lan, Tạp chí Công Thương: phát triển kinh tế số"),
    347: ("1.1.3.1", "Charumilin (2012), WIPO fellowship: thương mại hóa SHTT, kinh nghiệm Nhật Bản cho Thái Lan"),
    363: ("1.1.3.1", "Cuntz, Fink & Stamm (2024), WIPO ERWP 77: AI và SHTT dưới góc độ kinh tế"),
    422: ("1.1.1.1", "KIIP (2023): The Economic Effects of IPR (tiếng Hàn)"),
    440: ("1.1.1.1", "WIPO-UNU: Impact of the IP System on Economic Growth - Korea"),
    442: ("1.1.1.1", "WIPO: The Economics of IP in the Republic of Korea"),
    455: ("1.1.1.1", "Smarzynska (2002), World Bank PRWP 2786: cơ cấu FDI và bảo hộ quyền SHTT"),
    461: ("1.1.1.1", "Finger & Schuler (eds.) (2004), World Bank: Poor People's Knowledge"),
    464: ("1.1.1.1", "Bravo-Ortega & Lederman (2010), World Bank PRWP 5217: quyền SHTT, vốn nhân lực, R&D"),
    465: ("1.1.3.1", "Arora, Fosfuri & Gambardella (2002), ISSJ: Markets for technology (khác bài 2001 đã trích)"),
    472: ("1.1.2.1", "OECD Digital Economy Outlook 2020"),
    477: ("1.1.2.1", "OECD Digital Economy Outlook 2024, Volume 1"),
    473: ("1.1.2.1", "UNCTAD (2021): Manual for the Production of Statistics on the Digital Economy 2020"),
    475: ("1.1.2.1", "UNCTAD (2021): Digital Economy Report 2021"),
    478: ("1.1.2.1", "UNCTAD (2024): Digital Economy Report 2024, Overview"),
    507: ("1.1.3.1", "EUIPO (2023): IP Infringement and Enforcement Tech Watch Discussion Paper"),
    511: ("1.1.2.1", "Catalini & Gans (2019): Some Simple Economics of the Blockchain"),
    518: ("1.1.2.1", "Oxford Handbook of the Digital Economy (bản xem trước)"),
    519: ("1.1.2.1", "Liu Zhiyi (2022), Springer: Principles of Digital Economics"),
    527: ("1.1.2.1", "Øverby & Audestad (2021), Springer: Introduction to Digital Economics"),
    530: ("1.1.2.1", "Goldfarb, Greenstein & Tucker (eds.) (2015): Economic Analysis of the Digital Economy (chỉ có phần đầu sách)"),
    534: ("1.1.2.1", "Lubacha et al. (eds.) (2024), Routledge: The European Digital Economy"),
    542: ("1.1.2.1", "Śledziewska & Włoch (2021), Routledge: The Economics of Digital Transformation, chương 1"),
    543: ("1.1.2.1", "Brousseau & Curien (eds.) (2007), CUP: Internet and Digital Economics (chỉ có phần đầu sách)"),
    551: ("1.1.2.2", "Ngô Cẩm Tú (2024), LATS HVCTQG: nhân lực cho phát triển kinh tế số ở Thái Nguyên"),
    552: ("1.1.2.2", "Ngô Cẩm Tú (2024), toàn văn luận án"),
    553: ("1.1.2.2", "Ngô Cẩm Tú (2024), tóm tắt tiếng Việt"),
    557: ("1.1.1.1", "WIPO: Economic Aspects of IP in Countries with Economies in Transition"),
    559: ("1.1.3.1", "WIPO (2002): Incentives in Technology Transfer"),
    561: ("1.1.3.1", "WIPO (2019): IP Policy Template for Universities and Research Institutions"),
    562: ("1.1.3.1", "WIPO (2019): IP Policy Writers' Checklist"),
    563: ("1.1.3.1", "WIPO (2019): Guidelines for Customization of the IP Policy Template"),
    558: ("Bối cảnh", "OECD Economic Surveys: Viet Nam 2025"),
    566: ("Bối cảnh", "OECD FDI Qualities Review of Viet Nam"),
}
THAM_KHAO.pop(200)
THAM_KHAO.update({
    415: ("1.1.3.1", "Bradley, Hayter & Link (2013), Foundations and Trends in Entrepreneurship: mô hình chuyển giao công nghệ ở trường đại học"),
    416: ("1.1.3.1", "Etzkowitz (2003), Research Policy: trường đại học khởi nghiệp"),
    126: ("1.1.3.1", "Bản trùng của Lohrey & Willoughby (2025)"),
    95: ("1.1.3.1", "Man et al. (2025), Sustainability (MDPI): công nghệ số và quản trị SHTT"),
    71: ("1.1.1.1", "Khan & Wu (2020), Journal of Politics and Law: tác động của kinh tế số đến luật SHTT"),
    260: ("1.1.1.1", "Buzova & Karelina (2021), Legal Issues in the Digital Age: bảo hộ tư pháp quyền SHTT trong kinh tế số"),
    277: ("1.1.1.1", "Davis, The Digital Dilemma (tham luận hội nghị WIPO)"),
    204: ("1.1.1.1", "Goold, A Critical Introduction to IP Law (CUP), chỉ có phần đầu sách"),
    364: ("1.1.1.1", "Cuntz et al. (2024), WIPO ERWP 78: tiếp cận khoa học ở nước đang phát triển (ngoài trọng tâm)"),
    525: ("1.1.2.1", "OECD: The Digital Economy, Multinational Enterprises and International Investment Policy"),
    544: ("1.1.2.1", "Tapscott, The Digital Economy 20th Anniversary Edition (chương mẫu)"),
    6: ("1.1.1.2", "Thực thi SHTT tại biên giới, Tạp chí Nghiên cứu Tài chính kế toán (2025)"),
    292: ("1.1.1.2", "Nguyễn Trọng Quang (2025), Giáo dục và Xã hội: pháp luật về AI trong kinh tế số"),
    343: ("1.1.1.2", "Diễn đàn KH&CN số 11/2023: hiểu đúng về đổi mới sáng tạo (chưa đọc được tên tác giả)"),
    298: ("1.1.1.2", "Tạp chí Kinh doanh và Công nghệ số 01/2019: quyền SHTT (chưa đọc được tên tác giả)"),
    9: ("1.1.2.2", "Nguyễn Văn Yên, Nguyễn Quang Hưng (2026): hệ sinh thái AI chủ quyền và kinh tế số"),
    12: ("1.1.2.2", "Kinh tế - xã hội kỳ I 02/2022: mô hình kinh doanh mới, kinh tế số Việt Nam"),
    274: ("1.1.2.2", "Tạp chí Nghiên cứu Tài chính kế toán số 283 (2025): phát triển kinh tế số ở Việt Nam"),
    275: ("1.1.2.2", "Tạp chí NCKH Trường ĐH Sao Đỏ số 90 (2025), số đặc biệt về kinh tế số (mục lục)"),
    281: ("1.1.2.2", "Nguyễn Văn Tùng (2024), Tạp chí Tài chính: cải cách thể chế thúc đẩy kinh tế số"),
    284: ("1.1.2.2", "Phan Thị Thanh Tâm, Nguyễn Thị Kim Quyên (2024), Tạp chí Tài chính: hiệu quả quản lý nguồn lực trong kinh tế số"),
    286: ("1.1.2.2", "Đinh Thị Hằng (2025), Tạp chí KH Trường ĐH Mở HN: chính sách pháp luật thúc đẩy KHCN, ĐMST gắn với kinh tế số"),
    305: ("1.1.2.2", "Tạp chí Việt Nam Hội nhập số 383 (2026): kinh tế số và vai trò pháp luật"),
    313: ("1.1.2.2", "Tạp chí Quản lý nhà nước số 348 (1/2025): kinh tế tri thức và chính phủ điện tử"),
    315: ("1.1.2.2", "Tạp chí Việt Nam Hội nhập số 370 (2025): khái niệm kinh tế số"),
    319: ("1.1.2.2", "Tạp chí Nghiên cứu Tài chính kế toán số 262 (2024): kinh tế số"),
    326: ("1.1.2.2", "Số 1/2026: năng lực cạnh tranh quốc tế trong nền kinh tế số"),
    331: ("1.1.2.2", "Tạp chí Nghiên cứu Tài chính kế toán số 296 (2025): tài sản số"),
})
LOAI_THEM = {
    **{i: "Không đọc được lớp văn bản (lỗi phông hoặc bản quét)" for i in (307, 317, 323, 255)},
    **{i: "Bài viết mô tả chung hoặc tạp chí ít uy tín" for i in (7, 67, 251, 257, 268, 270, 278, 279, 294, 340)},
    **{i: "Ngoài trọng tâm đề tài" for i in (332, 454, 565)},
    **{i: "Bài điểm sách, không phải công trình gốc" for i in (457, 462)},
    **{i: "Tệp hành chính của hồ sơ luận án" for i in (549, 550)},
}
KINH_NGHIEM_THEM = {0, 2, 4, 10, 11, 13, 14, 287, 290, 505, 521, 568, 210}

# ------------------------------------------------------------------ không dùng làm nguồn học thuật chính
LOAI = {
    **{i: "Tóm tắt do công cụ AI (LeapSpace) tạo ra, không phải công trình khoa học" for i in (120, 128, 131, 145, 153, 155)},
    **{i: "Bản dịch máy (Machine Translated by Google)" for i in (232, 365, 366, 556)},
    **{i: "Luận văn thạc sĩ hoặc bản xem trước luận văn" for i in (70, 127, 142, 143)},
    **{i: "Tài liệu tập huấn, phổ biến kiến thức hoặc giáo trình nhập môn" for i in (5, 69, 194, 195, 196, 197, 199, 201, 202, 203, 205, 206, 207, 208, 212, 213, 215, 216, 217, 218, 220, 224, 225, 226, 227, 228, 230, 231, 234, 235, 236, 237, 238, 239, 240, 516)},
    **{i: "Tạp chí ít uy tín hoặc bài viết mô tả chung, không đủ chuẩn làm cơ sở lý luận" for i in (43, 68, 72, 73, 78, 79, 80, 81, 82, 83, 85, 86, 87, 88, 90, 91, 101, 102, 104, 106, 108, 109, 112, 116, 117, 134, 137, 140, 154, 158, 252, 253, 254, 256, 258, 259, 261, 262, 263, 265, 304, 496, 498, 503, 504, 514, 524, 547, 564, 567)},
    **{i: "Không liên quan đến đề tài" for i in (89, 141, 414, 417, 497, 501, 502, 508, 509, 529, 539, 541, 548)},
    **{i: "Giáo trình, sách nhập môn, bài giảng hoặc sách mẫu" for i in (92, 96, 97, 107, 165, 222, 526, 532, 535, 536, 540)},
    **{i: "Ngoài trọng tâm đề tài (luật cạnh tranh hoặc chủ đề quản trị chuyên biệt)" for i in (124, 147, 149, 150, 157)},
    **{i: "Tài liệu xám, bản thảo, tài liệu vận động chính sách hoặc báo cáo doanh nghiệp" for i in (75, 110, 136, 148, 156, 513, 522, 533)},
}


def muc_theo_thu_muc(path):
    top = path.split("/")[1]
    if top in ("07_KinhNghiem_TrungQuoc", "08_Kinhnghiem_HanQuoc", "10_KimNghiem_Singapo"):
        return "D", "Kinh nghiệm quốc tế (Chương 2/Chương 4), không đưa vào tổng quan"
    if top in ("09_BaoCao_SoHuuTriTue", "11_Baocao_Kinhteso_VN", "12_Baocao_KTS_QT", "13_Baocao_KTS_ĐNA", "20_sach trang"):
        return "D", "Số liệu, báo cáo bối cảnh (Chương 3, tính cấp thiết)"
    return None, None


ids_de_xuat = {x["id"]: x for x in DE_XUAT}
ids_xac_minh = {x["id"]: x for x in CAN_XAC_MINH}
rows = []
for r in docs:
    i, p = r["id"], r["path"]
    ext = p.rsplit(".", 1)[-1].lower()
    mo_ta = "".join(ch for ch in (r.get("text") or "")[:220] if ch >= " ")
    if i in DA_TRICH:
        cat, muc, note = "A", "", "Đã trích dẫn: " + DA_TRICH[i]
    elif i in ids_de_xuat:
        x = ids_de_xuat[i]; cat, muc, note = "B", x["muc"], f"{x['tac_gia']} ({x['nam']}), {x['ten']}"
    elif i in TRUNG_DE_XUAT:
        x = ids_de_xuat[TRUNG_DE_XUAT[i]]; cat, muc, note = "B*", x["muc"], f"Bản trùng hoặc tệp đi kèm của: {x['tac_gia']} ({x['nam']})"
    elif i in ids_xac_minh:
        x = ids_xac_minh[i]; cat, muc, note = "F", x["muc"], x["mo_ta"] + ". Thiếu: " + x["thieu"]
    elif i in THAM_KHAO:
        cat, (muc, note) = "C", THAM_KHAO[i]
    elif i in LOAI:
        cat, muc, note = "E", "", LOAI[i]
    elif i in LOAI_THEM:
        cat, muc, note = ("G" if "lớp văn bản" in LOAI_THEM[i] or "hành chính" in LOAI_THEM[i] else "E"), "", LOAI_THEM[i]
    elif i in KINH_NGHIEM_THEM:
        cat, muc, note = "D", "", "Kinh nghiệm quốc tế, văn bản pháp luật hoặc số liệu bối cảnh; không đưa vào tổng quan"
    elif ext in ("txt", "md", "js") or "/_files/" in p or len(mo_ta) < 40:
        cat, muc, note = "G", "", "Tệp phụ trợ, tệp rỗng hoặc tệp không có lớp văn bản"
    else:
        c, n = muc_theo_thu_muc(p)
        if c:
            cat, muc, note = c, "", n
        elif p.split("/")[1] in ("19_bosung luan an vn shtt",):
            cat, muc, note = "G", "", "Tệp hành chính của hồ sơ luận án (quyết định, trang thông tin)"
        else:
            cat, muc, note = "C", "", "Chưa xếp mục; đúng chủ đề chung nhưng giá trị bổ sung thấp"
    rows.append(dict(id=i, tep=p[len(PREFIX):], trang=r.get("pages", ""), nhom=cat, muc=muc, ghi_chu=note,
                     trung=" | ".join(x[len(PREFIX):] for x in r.get("dups", [])[1:]), mo_ta=mo_ta))

NHOM = {
    "A": "Đã trích dẫn trong chuyên đề",
    "B": "Đề xuất bổ sung, ưu tiên (đã xác minh trên bản gốc)",
    "B*": "Bản trùng hoặc tệp đi kèm của tài liệu nhóm B",
    "C": "Tham khảo thêm (đúng chủ đề, ưu tiên thấp hơn)",
    "D": "Dùng cho phần khác của luận án (kinh nghiệm quốc tế, số liệu)",
    "E": "Không dùng làm nguồn học thuật chính",
    "F": "Có giá trị nhưng thiếu thông tin thư mục, cần xác minh trước khi dùng",
    "G": "Tệp phụ trợ, tệp rỗng, tệp hành chính",
}

# ------------------------------------------------------------------ Excel
HEAD = Font(bold=True, color="FFFFFF"); FILL = PatternFill("solid", fgColor="1F4E78"); WRAP = Alignment(wrap_text=True, vertical="top")


def sheet(ws, header, data, widths):
    ws.append(header)
    for c in ws[1]:
        c.font, c.fill, c.alignment = HEAD, FILL, WRAP
    for d in data:
        ws.append(d)
    for k, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(k)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = WRAP
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


wb = Workbook()
ws = wb.active; ws.title = "1_DeXuat_BoSung"
sheet(ws, ["STT", "Mục", "Hướng nghiên cứu", "Tác giả", "Năm", "Tên công trình", "Nguồn", "Tập (số), trang, DOI", "Loại",
           "Mức xác minh", "Giá trị đối với luận án", "Ghi chú", "Tệp trong 01_NotebookLM_Inputs"],
      [[k + 1, x["muc"], x["huong"], x["tac_gia"], x["nam"], x["ten"], x["nguon"], x["so_trang"], x["loai"], x["xac_minh"],
        x["dung_cho_LA"], x["ghi_chu"], next(r["tep"] for r in rows if r["id"] == x["id"])] for k, x in enumerate(DE_XUAT)],
      [5, 12, 30, 24, 8, 45, 30, 26, 16, 9, 55, 35, 45])
ws = wb.create_sheet("2_LiteratureReview")
sheet(ws, ["STT (No)", "Tác giả (Authors)", "Năm (year)", "Tiêu đề (Title)", "Tạp chí (Journal) / nơi xuất bản", "Số (trang) (Volume/page)",
           "Vấn đề đặt ra", "Kết quả nghiên cứu (research result)", "Phương pháp nghiên cứu (Research methods)", "Mẫu (Sample)",
           "Phương pháp phân tích (Method of data analysis)", "Nội dung bằng tiếng Anh nguyên gốc dùng cho luận án",
           "Nội dung dùng để phân tích trong luận án dịch sang tiếng Việt", "Số trang trích dẫn nội dung bài báo dùng cho luận án", "Mục đề xuất"],
      [[k + 1, x["tac_gia"], x["nam"], x["ten"], x["nguon"], x["so_trang"], x["van_de"], x["ket_qua"], x["phuong_phap"], x["mau"],
        x["phan_tich"], x["nguyen_van"], x["dung_cho_LA"], x["trang"], x["muc"]] for k, x in enumerate(DE_XUAT)],
      [6, 22, 8, 40, 26, 20, 35, 55, 22, 26, 22, 45, 45, 22, 12])
ws = wb.create_sheet("3_CanXacMinh")
sheet(ws, ["STT", "Mục", "Mô tả đọc được", "Thông tin còn thiếu", "Tệp"],
      [[k + 1, x["muc"], x["mo_ta"], x["thieu"], x["tep"]] for k, x in enumerate(CAN_XAC_MINH)], [5, 16, 70, 40, 60])
ws = wb.create_sheet("4_ToanBo_TaiLieu")
order = {"B": 0, "B*": 1, "F": 2, "C": 3, "A": 4, "D": 5, "E": 6, "G": 7}
rows_sorted = sorted(rows, key=lambda r: (order[r["nhom"]], r["muc"], r["tep"]))
sheet(ws, ["Mã", "Nhóm", "Diễn giải nhóm", "Mục gợi ý", "Ghi chú, lý do", "Tệp", "Số trang PDF", "Bản trùng (cùng nội dung)", "Đoạn đầu văn bản"],
      [[r["id"], r["nhom"], NHOM[r["nhom"]], r["muc"], r["ghi_chu"], r["tep"], r["trang"], r["trung"], r["mo_ta"]] for r in rows_sorted],
      [6, 6, 30, 12, 60, 60, 8, 40, 60])
ws = wb.create_sheet("5_ChuGiai")
sheet(ws, ["Nhóm", "Ý nghĩa", "Số tệp"], [[k, v, sum(1 for r in rows if r["nhom"] == k)] for k, v in NHOM.items()], [8, 70, 10])
ws.append([])
ws.append(["", "Phạm vi rà soát: toàn bộ tệp .pdf, .docx, .txt, .md trong 01_NotebookLM_Inputs, trừ 01_VanBanPhapQuy và Niêm giám; tệp trùng nội dung (cùng mã băm) chỉ tính một lần."])
ws.append(["", "Mức xác minh A: tác giả, năm, tên, nguồn đọc được trên bản gốc. B: còn thiếu một chi tiết, đã ghi ở cột Ghi chú."])
ws.append(["", "Số trang in suy từ số trang của tệp được ghi rõ trong cột Ghi chú; cần đối chiếu bản in trước khi đưa vào danh mục tài liệu tham khảo."])
out = os.path.join(HERE, "PhanLoai_TaiLieu_TongQuan_2026-09-30.xlsx")
wb.save(out)

# ------------------------------------------------------------------ Literature review dạng Markdown
cols = ["STT", "Tác giả", "Năm", "Tiêu đề", "Nơi xuất bản", "Số (trang)", "Vấn đề đặt ra", "Kết quả nghiên cứu", "Phương pháp nghiên cứu",
        "Mẫu", "Phương pháp phân tích", "Nội dung tiếng Anh nguyên gốc", "Nội dung dùng để phân tích trong luận án", "Trang trích dẫn"]
md = ["# Bảng Literature Review: tài liệu đề xuất bổ sung cho chuyên đề tổng quan", "",
      "Rà soát ngày 30/9/2026 từ thư mục `01_NotebookLM_Inputs`. Nội dung mỗi dòng được đọc từ bản gốc; cột cuối ghi trang đã đọc. "
      "Cột \"Nội dung dùng để phân tích\" là gợi ý của người rà soát, không phải trích dẫn.", ""]
for muc in ["1.1.1.1", "1.1.1.2", "1.1.2.1", "1.1.2.2", "1.1.3 (đoạn mở đầu) / 1.1.3.1", "1.1.3.1", "1.1.3.2"]:
    grp = [x for x in DE_XUAT if x["muc"] == muc]
    if not grp:
        continue
    md += [f"## Mục {muc}", "", "| " + " | ".join(cols) + " |", "|" + "---|" * len(cols)]
    for k, x in enumerate(grp, 1):
        vals = [str(k), x["tac_gia"], x["nam"], "*" + x["ten"] + "*", x["nguon"], x["so_trang"], x["van_de"], x["ket_qua"], x["phuong_phap"],
                x["mau"], x["phan_tich"], x["nguyen_van"], x["dung_cho_LA"], x["trang"]]
        md.append("| " + " | ".join(v.replace("|", "/") for v in vals) + " |")
    md.append("")
open(os.path.join(HERE, "Literature_Review_BoSung_TongQuan.md"), "w", encoding="utf-8").write("\n".join(md))

# ------------------------------------------------------------------ BibTeX (chỉ trường đã xác minh)
def key(x):
    import unicodedata
    a = unicodedata.normalize("NFD", x["tac_gia"].split(",")[0].split()[0].lower().replace("đ", "d"))
    return "".join(ch for ch in a if ch.isascii() and ch.isalnum()) + x["nam"][:4] + "_" + str(x["id"])


bib = ["% Tài liệu đề xuất bổ sung (rà soát 30/9/2026). Chỉ gồm trường đã đọc được trên bản gốc; mục có ghi chú cần đối chiếu trước khi dùng."]
for x in DE_XUAT:
    typ = ("phdthesis" if x["loai"].startswith("Luận án") else "book" if x["loai"].startswith("Sách")
           else "article" if x["loai"].startswith("Bài báo") or "Tổng quan" in x["loai"] else "techreport")
    fields = {"author": x["tac_gia"].replace(", ", " and ") if typ == "article" else x["tac_gia"], "year": x["nam"][:4], "title": x["ten"]}
    if typ == "article":
        fields["journal"] = x["nguon"]
    elif typ == "phdthesis":
        fields["school"] = x["nguon"]
    elif typ == "book":
        fields["publisher"] = x["nguon"]
    else:
        fields["institution"] = x["nguon"]
    if x["so_trang"]:
        fields["note"] = x["so_trang"]
    if x["ghi_chu"]:
        fields["annote"] = x["ghi_chu"]
    bib.append(f"@{typ}{{{key(x)},\n" + ",\n".join(f"  {k} = {{{v}}}" for k, v in fields.items()) + "\n}")
open(os.path.join(HERE, "references_bo_sung_tong_quan.bib"), "w", encoding="utf-8").write("\n\n".join(bib) + "\n")

print("Đã tạo:", out)
for k, v in NHOM.items():
    print(k, sum(1 for r in rows if r["nhom"] == k), v)
